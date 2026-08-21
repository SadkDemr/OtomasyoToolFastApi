import io

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from database import Base, get_db
from routers.auth import get_current_user
from routers.testcasegen_router import router as testcasegen_router
from services import ai_service, testcasegen_service as svc, knowledge_service


class _FakeUser:
    id = 1
    full_name = "Test Kullanici"
    username = "tester"
    tcg_area_path = "Mobil\\Hesaplama"
    tcg_assigned_to = "Test Kullanici"
    tcg_state = "Ready"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(svc, "DATA_DIR", tmp_path)
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    TestSession = sessionmaker(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = TestSession()

    def override_get_db():
        yield db

    app = FastAPI()
    app.include_router(testcasegen_router)
    app.dependency_overrides[get_current_user] = lambda: _FakeUser()
    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app)


def _session_with_generated_testcase(client, monkeypatch):
    async def fake_text(prompt, model=None, timeout=60.0):
        return '{"steps": [{"step_action": "X tiklanir", "step_expected": "Y"}]}'

    monkeypatch.setattr(ai_service, "ollama_generate_text", fake_text)
    session = client.post("/api/testcasegen/sessions").json()
    session = client.post(
        f"/api/testcasegen/sessions/{session['id']}/titles",
        json={"screen_indices": [0], "text": "Baslik"},
    ).json()
    title_id = session["titles"][0]["id"]
    gen = client.post(
        f"/api/testcasegen/sessions/{session['id']}/generate-testcases",
        json={"title_ids": [title_id]},
    ).json()
    return session["id"], gen["testcases"][0]["id"]


def test_export_returns_valid_xlsx(client, monkeypatch):
    session_id, tc_id = _session_with_generated_testcase(client, monkeypatch)
    resp = client.post(f"/api/testcasegen/sessions/{session_id}/export", json={"testcase_ids": [tc_id]})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws[1][0].value == "ID"
    assert ws[2][2].value == "Baslik"


def test_export_no_testcase_ids_returns_400(client, monkeypatch):
    session_id, _ = _session_with_generated_testcase(client, monkeypatch)
    resp = client.post(f"/api/testcasegen/sessions/{session_id}/export", json={"testcase_ids": []})
    assert resp.status_code == 400


def test_export_writes_area_path_assigned_to_state_on_testcase_row(client, monkeypatch):
    session_id, tc_id = _session_with_generated_testcase(client, monkeypatch)
    resp = client.post(f"/api/testcasegen/sessions/{session_id}/export", json={"testcase_ids": [tc_id]})
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    testcase_row = ws[2]
    assert testcase_row[6].value == "Mobil\\Hesaplama"
    assert testcase_row[7].value == "Test Kullanici"
    assert testcase_row[8].value == "Ready"


def test_export_adds_scenario_knowledge_entry_per_testcase(client, monkeypatch):
    calls = []

    async def fake_add_entry(db, entry_type, title, content, user_id=None, source_ref=None):
        calls.append((entry_type, title, content))
        return 1

    monkeypatch.setattr(knowledge_service, "add_entry", fake_add_entry)

    session_id, tc_id = _session_with_generated_testcase(client, monkeypatch)
    resp = client.post(
        f"/api/testcasegen/sessions/{session_id}/export",
        json={"testcase_ids": [tc_id]},
    )
    assert resp.status_code == 200
    assert len(calls) == 1
    assert calls[0][0] == "scenario"


def test_export_succeeds_even_if_knowledge_add_entry_fails(client, monkeypatch):
    async def failing_add_entry(*args, **kwargs):
        raise RuntimeError("embedding modeli yok")

    monkeypatch.setattr(knowledge_service, "add_entry", failing_add_entry)

    session_id, tc_id = _session_with_generated_testcase(client, monkeypatch)
    resp = client.post(
        f"/api/testcasegen/sessions/{session_id}/export",
        json={"testcase_ids": [tc_id]},
    )
    assert resp.status_code == 200
