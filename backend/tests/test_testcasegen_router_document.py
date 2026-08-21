import io

import openpyxl
import pytest
from docx import Document as DocxDocument
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from database import Base, get_db
from routers.auth import get_current_user
from routers.testcasegen_router import router as testcasegen_router
from services import ai_service, testcasegen_service as svc, knowledge_service


class _FakeUser:
    id = 1


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(svc, "DATA_DIR", tmp_path)

    async def _default_ollama_embed(text):
        return None

    monkeypatch.setattr(ai_service, "ollama_embed", _default_ollama_embed)

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    TestSession = sessionmaker(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = TestSession()

    def override_get_db():
        yield db

    app = FastAPI()
    app.include_router(testcasegen_router)
    app.dependency_overrides[get_current_user] = lambda: _FakeUser()
    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app)


def _docx_bytes(text):
    doc = DocxDocument()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_document_rejects_unsupported_extension(client):
    session = client.post("/api/testcasegen/sessions").json()
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("analiz.pdf", b"%PDF-1.4", "application/pdf")},
    )
    assert resp.status_code == 400


def test_upload_document_stores_extracted_text(client):
    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("Giris ekrani kullanici adi ve sifre alir.")
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("analiz.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["document"]["filename"] == "analiz.docx"
    assert "Giris ekrani kullanici adi ve sifre alir." in body["document"]["extracted_text"]
    assert body["document"]["truncated"] is False


def test_upload_document_empty_text_returns_400(client):
    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("")
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("bos.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 400


def test_upload_document_corrupt_docx_returns_400(client):
    session = client.post("/api/testcasegen/sessions").json()
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("bad.docx", b"not a docx file", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 400


def test_upload_document_stores_extracted_text_from_xlsx(client):
    session = client.post("/api/testcasegen/sessions").json()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Giris ekrani", "kullanici adi ve sifre alir"])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("analiz.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["document"]["ext"] == ".xlsx"
    assert "Giris ekrani" in body["document"]["extracted_text"]
    assert "kullanici adi ve sifre alir" in body["document"]["extracted_text"]


def test_delete_document_clears_field(client):
    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("Bir seyler yaz.")
    session = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("a.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    resp = client.delete(f"/api/testcasegen/sessions/{session['id']}/document")
    assert resp.status_code == 200
    assert resp.json()["document"] is None


def test_extract_titles_from_document_without_upload_returns_400(client):
    session = client.post("/api/testcasegen/sessions").json()
    resp = client.post(f"/api/testcasegen/sessions/{session['id']}/document/extract-titles")
    assert resp.status_code == 400


def test_extract_titles_from_document_populates_candidates(client, monkeypatch):
    async def fake_text(prompt, model=None, timeout=60.0):
        return '{"titles": [{"text": "Giris basarili senaryo", "excerpt": "Kullanici adi ve sifre girilir"}]}'

    monkeypatch.setattr(ai_service, "ollama_generate_text", fake_text)
    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("Kullanici adi ve sifre girilir, giris yapilir.")
    session = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("a.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()

    resp = client.post(f"/api/testcasegen/sessions/{session['id']}/document/extract-titles")
    assert resp.status_code == 200
    candidates = resp.json()["document"]["candidate_titles"]
    assert candidates == [{"text": "Giris basarili senaryo", "excerpt": "Kullanici adi ve sifre girilir"}]


def test_extract_titles_from_document_ollama_failure_returns_503(client, monkeypatch):
    async def failing_text(prompt, model=None, timeout=60.0):
        raise ConnectionError("ollama kapali")

    monkeypatch.setattr(ai_service, "ollama_generate_text", failing_text)
    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("Bir metin.")
    session = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("a.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()

    resp = client.post(f"/api/testcasegen/sessions/{session['id']}/document/extract-titles")
    assert resp.status_code == 503


def test_add_title_with_doc_excerpt_sets_document_source(client):
    session = client.post("/api/testcasegen/sessions").json()
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/titles",
        json={"text": "Odeme akisi", "doc_excerpt": "Kullanici tutar girer, onaylar."},
    )
    assert resp.status_code == 200
    title = resp.json()["titles"][0]
    assert title["source"] == "document"
    assert title["doc_excerpt"] == "Kullanici tutar girer, onaylar."
    assert title["screen_indices"] == []


def test_add_title_without_screens_or_doc_excerpt_returns_400(client):
    session = client.post("/api/testcasegen/sessions").json()
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/titles",
        json={"text": "Baslik"},
    )
    assert resp.status_code == 400


def test_prioritize_scores_document_sourced_title_without_crash(client, monkeypatch):
    async def fake_text(prompt, model=None, timeout=60.0):
        return "Gerekce metni."

    monkeypatch.setattr(ai_service, "ollama_generate_text", fake_text)
    session = client.post("/api/testcasegen/sessions").json()
    session = client.post(
        f"/api/testcasegen/sessions/{session['id']}/titles",
        json={"text": "Odeme onaylanir", "doc_excerpt": "Kullanici odeme onaylar."},
    ).json()

    resp = client.post(f"/api/testcasegen/sessions/{session['id']}/prioritize")
    assert resp.status_code == 200
    title = resp.json()["titles"][0]
    assert title["level"] in ("high", "medium", "low")
    assert title["score"] is not None


def test_generate_testcases_for_document_title_uses_text_prompt_not_vision(client, monkeypatch):
    captured = {}

    async def fake_text(prompt, model=None, timeout=60.0):
        captured["prompt"] = prompt
        return '{"steps": [{"step_action": "X tıklanır", "step_expected": "Y görülür"}]}'

    async def fake_vision(prompt, images_b64, model=None, timeout=60.0):
        raise AssertionError("Dokuman kaynakli title icin vision cagrilmamali")

    monkeypatch.setattr(ai_service, "ollama_generate_text", fake_text)
    monkeypatch.setattr(ai_service, "ollama_generate_vision", fake_vision)

    session = client.post("/api/testcasegen/sessions").json()
    session = client.post(
        f"/api/testcasegen/sessions/{session['id']}/titles",
        json={"text": "Odeme onaylanir", "doc_excerpt": "Kullanici tutari girer ve onaylar."},
    ).json()
    title_id = session["titles"][0]["id"]

    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/generate-testcases",
        json={"title_ids": [title_id]},
    )
    assert resp.status_code == 200
    testcases = resp.json()["testcases"]
    assert len(testcases) == 1
    assert "Kullanici tutari girer ve onaylar." in captured["prompt"]


def test_upload_document_adds_knowledge_entry(client, monkeypatch):
    calls = []

    async def fake_add_entry(db, entry_type, title, content, user_id=None, source_ref=None):
        calls.append((entry_type, title, content))
        return 1

    monkeypatch.setattr(knowledge_service, "add_entry", fake_add_entry)

    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("Giris ekrani test edilir.")
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("analiz.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200
    assert len(calls) == 1
    assert calls[0][0] == "document"
    assert calls[0][1] == "analiz.docx"
    assert "Giris ekrani test edilir." in calls[0][2]


def test_upload_document_succeeds_even_if_knowledge_add_entry_fails(client, monkeypatch):
    async def failing_add_entry(*args, **kwargs):
        raise RuntimeError("embedding modeli yok")

    monkeypatch.setattr(knowledge_service, "add_entry", failing_add_entry)

    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("Giris ekrani test edilir.")
    resp = client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("analiz.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200


def test_extract_titles_from_document_injects_rule_context(client, monkeypatch):
    captured = {}

    async def fake_text(prompt, model=None, timeout=60.0):
        captured["prompt"] = prompt
        return '{"titles": []}'

    monkeypatch.setattr(ai_service, "ollama_generate_text", fake_text)

    async def fake_search(db, query_text, top_k=5, entry_types=None):
        assert entry_types == ["rule"]
        return ["Sirket standardi: her ekran icin en az 1 negatif senaryo yazilmali"]

    monkeypatch.setattr(knowledge_service, "search", fake_search)

    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("Giris ekrani test edilir.")
    client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("analiz.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    client.post(f"/api/testcasegen/sessions/{session['id']}/document/extract-titles")
    assert "negatif senaryo" in captured["prompt"]


def test_extract_titles_from_document_continues_when_search_fails(client, monkeypatch):
    async def fake_text(prompt, model=None, timeout=60.0):
        return '{"titles": []}'

    monkeypatch.setattr(ai_service, "ollama_generate_text", fake_text)

    async def failing_search(db, query_text, top_k=5, entry_types=None):
        raise RuntimeError("db hatasi")

    monkeypatch.setattr(knowledge_service, "search", failing_search)

    session = client.post("/api/testcasegen/sessions").json()
    data = _docx_bytes("Giris ekrani test edilir.")
    client.post(
        f"/api/testcasegen/sessions/{session['id']}/document/upload",
        files={"file": ("analiz.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resp = client.post(f"/api/testcasegen/sessions/{session['id']}/document/extract-titles")
    assert resp.status_code == 200
