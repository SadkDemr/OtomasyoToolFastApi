import io

import openpyxl
import pytest
from docx import Document as DocxDocument

from services import testcasegen_service as svc


def _make_docx_bytes(paragraphs, table_rows=None):
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx_bytes(sheet_name, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_text_from_docx_reads_paragraphs_and_tables():
    data = _make_docx_bytes(
        ["Giris ekrani kullanici adi ve sifre alir.", "Onayla butonuna basilinca giris yapilir."],
        table_rows=[["Alan", "Zorunlu"], ["Kullanici Adi", "Evet"]],
    )
    text, truncated = svc.extract_text_from_docx(data)
    assert "Giris ekrani kullanici adi ve sifre alir." in text
    assert "Kullanici Adi | Evet" in text
    assert truncated is False


def test_extract_text_from_docx_truncates_long_text():
    data = _make_docx_bytes(["a" * (svc.DOC_TEXT_CHAR_LIMIT + 20000)])
    text, truncated = svc.extract_text_from_docx(data)
    assert len(text) == svc.DOC_TEXT_CHAR_LIMIT
    assert truncated is True


def test_extract_text_from_xlsx_keeps_every_sheet_when_one_sheet_is_huge():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws1 = wb.create_sheet("BuyukSayfa")
    for i in range(400):
        ws1.append([f"satir{i}", "x" * 200])
    ws2 = wb.create_sheet("KucukSayfa")
    ws2.append(["Login", "Kullanici giris yapar"])
    buf = io.BytesIO()
    wb.save(buf)

    text, truncated = svc.extract_text_from_xlsx(buf.getvalue())
    assert "Sayfa: BuyukSayfa" in text
    assert "Sayfa: KucukSayfa" in text
    assert "Login | Kullanici giris yapar" in text
    assert truncated is True


def test_extract_text_from_docx_invalid_bytes_raises_value_error():
    with pytest.raises(ValueError):
        svc.extract_text_from_docx(b"not a docx file")


def test_extract_text_from_xlsx_reads_sheet_rows():
    data = _make_xlsx_bytes("Analiz", [["Ekran", "Aksiyon"], ["Login", "Kullanici giris yapar"]])
    text, truncated = svc.extract_text_from_xlsx(data)
    assert "Sayfa: Analiz" in text
    assert "Login | Kullanici giris yapar" in text
    assert truncated is False


def test_extract_text_from_xlsx_invalid_bytes_raises_value_error():
    with pytest.raises(ValueError):
        svc.extract_text_from_xlsx(b"not an xlsx file")
