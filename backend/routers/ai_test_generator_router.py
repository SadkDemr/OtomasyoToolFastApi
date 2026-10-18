"""
AI Test Generator Router  v3  —  Hibrit Mimari
===============================================
1. Kural Motoru  → Türkçe test adımlarını deterministik Appium/Selenium koduna çevirir
2. LLM Desteği   → Kural motorunun tanımadığı adımlar için Ollama (gemma3:4b)
3. Vision Fallback → Runtime'da element bulunamazsa screenshot → LLM/Gemini → tap(x,y)
"""

import asyncio
import os
import re
import subprocess
import sys
import tempfile
import textwrap
import uuid
from typing import Optional

import httpx
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from pathlib import Path
from database import get_db
from sqlalchemy.orm import Session
from routers.auth import get_current_user
from models.db_models import User, UserWorkspace
from services import git_service
import json as _json
import datetime as _datetime

router = APIRouter(prefix="/api/ai-gen", tags=["ai-generator"])

_running_processes: dict[str, subprocess.Popen] = {}

import time as _time_mod

_active_sessions: dict[tuple[int, str], dict] = {}   # (user_id, platform) -> session info
_SESSION_IDLE_TIMEOUT_SEC = 15 * 60


def _touch_session(user_id: int, platform: str) -> None:
    key = (user_id, platform)
    if key in _active_sessions:
        _active_sessions[key]["last_used_at"] = _time_mod.time()


def _get_session(user_id: int, platform: str) -> dict | None:
    return _active_sessions.get((user_id, platform))


def _remove_session(user_id: int, platform: str) -> None:
    info = _active_sessions.pop((user_id, platform), None)
    if info and info.get("lease_token"):
        leases.release(info.get("device_id", ""), info["lease_token"])


def _quit_session(key: tuple[int, str]) -> None:
    """Oturuma bağlanıp driver.quit() eder ve kaydı (telefon kirasıyla birlikte) siler."""
    info = _active_sessions.get(key)
    if not info:
        return
    try:
        if key[1] == "mobile":
            from services.session_reconnect import attach_appium_session
            d = attach_appium_session(info["command_executor_url"], info["session_id"])
        else:
            from services.session_reconnect import attach_selenium_session
            d = attach_selenium_session(info["command_executor_url"], info["session_id"])
        d.quit()
    except Exception as e:
        print(f"[SESSION] Kapatma hatasi ({key}): {e}")
    _remove_session(*key)


# ── Telefon paylaşımı (birden fazla kullanıcı, birden fazla telefon) ─────────
# Telefonu kullanan her yol önce kiralar: services/device_lease_service.py
from services import device_lease_service as leases


from services import device_node_service as dnodes


def _adb_device_ids() -> list[str]:
    """Tüm cihaz sunucularında (yerel + uzak) 'device' durumundaki telefonlar."""
    return dnodes.adb_devices()


def _resolve_device(device_id: str) -> str:
    """Cihaz seçilmediyse tek bağlı telefon kullanılır; birden fazlaysa seçim zorunlu
    (aksi halde Appium rastgele bir telefonu — belki başkasının telefonunu — alırdı)."""
    device_id = (device_id or "").strip()
    if device_id:
        return device_id
    devs = _adb_device_ids()
    if len(devs) == 1:
        return devs[0]
    if not devs:
        raise HTTPException(400, "Bağlı telefon bulunamadı.")
    raise HTTPException(400, "Birden fazla telefon bağlı — lütfen kullanacağınız telefonu seçin.")


def _busy_http(e: "leases.DeviceBusyError") -> HTTPException:
    return HTTPException(423, leases.busy_http_detail(e))


# ── Tek-adım test havuzu: SQL/API/OKU adımlarının kaydettiği {{değişken}}ler,
# izole alt-süreçlerde çalışan "▶ Test Et" adım testleri arasında hayatta kalsın ──
_step_vars_pool: dict[int, dict[str, str]] = {}   # user_id -> {var_adı: değer}
_VARS_OUT_RE = re.compile(r'\[(?:SQL|API|OKU)→_VARS\]\s*(\w+)\s*=\s*(.*)')


def _get_step_vars(user_id: int) -> dict:
    return _step_vars_pool.setdefault(user_id, {})


def _ingest_step_vars(user_id: int, output: str) -> None:
    if not output:
        return
    pool = _step_vars_pool.setdefault(user_id, {})
    for name, value in _VARS_OUT_RE.findall(output):
        pool[name] = value.strip()


def _reset_step_vars(user_id: int) -> None:
    _step_vars_pool[user_id] = {}


async def _session_idle_reaper():
    """15 dakika boşta kalan oturumları otomatik kapatır (driver.quit())."""
    while True:
        await asyncio.sleep(60)
        now = _time_mod.time()
        for key in list(_active_sessions.keys()):
            info = _active_sessions.get(key)
            if not info:
                continue
            if now - info["last_used_at"] > _SESSION_IDLE_TIMEOUT_SEC:
                # Telefon kirası da düşer → başkaları kullanabilir
                await asyncio.to_thread(_quit_session, key)
                print(f"[SESSION-REAPER] Idle oturum kapatildi: user={key[0]} platform={key[1]}")


_stop_files:        dict[str, str]               = {}   # run_id → stop flag dosyası
_desktop_lock:      asyncio.Lock                 = None # tek seferlik desktop kilidi (startup'ta init edilir)

def _get_desktop_lock() -> asyncio.Lock:
    global _desktop_lock
    if _desktop_lock is None:
        _desktop_lock = asyncio.Lock()
    return _desktop_lock

# ── Stop flag helper (her platforma gömülür) ─────────────────────────────────
STOP_HELPER = '''\
import os as _os_s, sys as _sys_s
_STOP_FILE = _os_s.environ.get("_AI_STOP_FILE", "")
def _check_stop():
    if _STOP_FILE and _os_s.path.exists(_STOP_FILE):
        try: _os_s.unlink(_STOP_FILE)
        except: pass
        raise SystemExit(0)
'''

from services.ai_service import get_ollama_config as _get_ollama_cfg

def _ollama_url():   return _get_ollama_cfg()["url"]
def _ollama_model(): return _get_ollama_cfg()["model"]

OLLAMA_URL   = "http://localhost:11434/api/generate"  # geriye dönük compat
OLLAMA_MODEL = "gemma3:4b"

# Backend kök dizini (ai_config.json burada)
_ROUTER_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ─── Pydantic ─────────────────────────────────────────────────────────────────

class StepItem(BaseModel):
    type:          str = "text"   # "text" | "image" | "sql" | "api" | "comment" | "condition"
    disabled:      bool = False   # True ise adım çalıştırılmaz, "# [DEVRE DIŞI]" yorumuna döner
    content:       str = ""
    image_b64:     str = ""
    action:        str = "click"  # image: click/double_click/right_click/assert | sql: query_print/execute | api: print/execute/store
    # SQL adımı
    db_connection: str = ""       # sqlite:///path.db  veya  mysql+pymysql://user:pass@host/db
    sql_query:     str = ""
    # API adımı
    api_url:       str = ""
    api_method:    str = "GET"
    api_headers:   str = ""       # JSON string
    api_body:      str = ""       # JSON string
    # Değişken sistemi
    var_name:      str = ""       # dolu ise sonuç _VARS["var_name"]'e kaydedilir
    json_field:    str = ""       # API yanıtından çekilecek alan yolu: "id" veya "user.name"
    # Koordinat tıklama (nokta işaretle modu)
    click_x:       int = -1       # >= 0 ise kırpılan/tam görsele göre nokta tıklaması kullanılır
    click_y:       int = -1       # (coklu-olcek eslestirme + normalize nokta, bkz. _find_b64_and_click_point)
    # Fare ile sürükleme (action="drag"): click_x/click_y = başlangıç, drag_x/drag_y = bitiş noktası
    drag_x:        int = -1
    drag_y:        int = -1
    # ── Koşullu dallanma (type == "condition") ────────────────────────────
    condition_type: str = ""      # "" | "image_exists" | "element_exists" | "variable_compare"
    compare_var:    str = ""      # variable_compare: {{degisken}} adı (_VARS anahtarı)
    compare_op:     str = ""      # variable_compare: "eq" | "neq" | "contains"
    compare_value:  str = ""      # variable_compare: karşılaştırılacak sabit değer
    then_steps:     list["StepItem"] = []
    else_steps:     list["StepItem"] = []

StepItem.model_rebuild()   # self-referential model — forward ref ("StepItem") çözümü için gerekli

class GenerateRequest(BaseModel):
    test_case:    str         = ""   # geriye dönük uyumluluk
    steps_list:   list[StepItem] = []   # yeni yapılandırılmış adım listesi
    platform:     str
    browser:      Optional[str] = "chrome"
    target_url:   Optional[str] = ""
    device_id:    Optional[str] = ""
    app_package:  Optional[str] = ""
    app_activity: Optional[str] = ""
    appium_url:   Optional[str] = "http://localhost:4723"
    # Paralel job hücrelerinde her cihaz için ayrı UiAutomator2 systemPort gerekir
    # (aynı Appium sunucusunda eşzamanlı oturumlar varsayılan portta çakışır).
    # Manuel tek-cihaz akışında None kalır → script'e systemPort yazılmaz (eski davranış).
    system_port:  Optional[int] = None
    app_path:       Optional[str]  = ""
    window_title:   Optional[str]  = ""
    reuse_instance: Optional[bool] = True
    single_step_test: Optional[bool] = False   # tek adım "Test Et" çağrısı mı?
    locator_file:   Optional[str]  = ""      # mobil: workspace/locators/<isim>.json (kayıt modu XPath'leri)

class ExecuteRequest(BaseModel):
    script:          str
    platform:        str
    scenario_name:   str       = "Adsız Senaryo"
    scenario_folder: str       = ""
    step_labels:     list[str] = []   # adım etiketleri (log için)


# ══════════════════════════════════════════════════════════════════════════════
# BÖLÜM 1: KURAL MOTORU
# Türkçe test adımı → Appium/Selenium Python satırları
# ══════════════════════════════════════════════════════════════════════════════

def _clean_step(raw: str) -> str:
    """Adım numarası, etiket ve beklenen sonuç suffix'ini temizle."""
    s = raw.strip()
    s = re.sub(r'^(adım|adim|step)\s*\d+\s*[:\-\.\)]\s*', '', s, flags=re.IGNORECASE)
    # "→ Beklenen: ..." suffix'ini kural motoru için kes (LLM bağlamı için orijinal satır kullanılır)
    s = re.sub(r'\s*[→>]+\s*[Bb]eklenen:.*$', '', s)
    return s.strip()


def _extract_quoted(text: str) -> list[str]:
    """Metindeki tırnak içi değerleri döndürür."""
    return re.findall(r'["\u201c\u201d\u2018\u2019]([^"\u201c\u201d\u2018\u2019]+)["\u201c\u201d\u2018\u2019]', text)


def _extract_key_target(text: str) -> str | None:
    """
    Karmaşık adım açıklamalarından asıl hedef elementi çıkarır.
    "Toplam Varlıklarım kartında yer alan favori fonlarım" → "favori fonlarım"
    "Ana Sayfa ekranındaki Hesaplarım butonu" → "Hesaplarım"
    """
    # "X kartında/bölümünde/ekranında/menüsünde [yer alan/bulunan/içindeki] Y"
    m = re.search(
        r'(?:kart[ıi]nda|bölümünde|ekranında|sayfasında|menüsünde|widget[ıi]nda)'
        r'(?:\s+(?:yer\s+alan|bulunan|içindeki|görünen))?\s+(.+)$',
        text, re.IGNORECASE
    )
    if m:
        return m.group(1).strip()
    # "yer alan / bulunan / içindeki" marker
    m2 = re.search(r'(?:yer\s+alan|içindeki|bulunan|görünen)\s+(.+)$', text, re.IGNORECASE)
    if m2:
        return m2.group(1).strip()
    return None


def _get_auto_value(context: str) -> str:
    """Alan bağlamına göre otomatik test değeri üretir."""
    ctx = context.lower()
    if re.search(r'tutar|miktar|amount|fiyat|para|finansman', ctx):
        return '10000'
    if re.search(r'telefon|gsm|phone|cep|numara', ctx):
        return '05551234567'
    if re.search(r'tc\b|tckn|kimlik', ctx):
        return '12345678901'
    if re.search(r'şifre|sifre|password|pin|parola', ctx):
        return '123456'
    if re.search(r'email|eposta|e-posta|mail', ctx):
        return 'test@test.com'
    if re.search(r'isim|ad\b|name|first', ctx):
        return 'Test'
    if re.search(r'soyisim|soyad|last', ctx):
        return 'Kullanici'
    return '121212'


def _parse_azure_devops_csv(decoded: str) -> str | None:
    """
    Azure DevOps CSV formatını tanır ve Step Action sütununu çıkarır.
    Döndürür: "Adım 1: ...\nAdım 2: ..." formatında metin, tanımazsa None.
    """
    import io, csv
    reader = csv.reader(io.StringIO(decoded))
    rows = list(reader)
    if not rows:
        return None

    header = [c.strip() for c in rows[0]]
    header_lower = [h.lower() for h in header]

    # Azure DevOps format tespiti
    if not any('step action' in h for h in header_lower):
        return None

    step_num_idx    = next((i for i, h in enumerate(header_lower) if 'test step' in h), None)
    step_action_idx = next((i for i, h in enumerate(header_lower) if 'step action' in h), None)
    if step_action_idx is None:
        return None

    step_expected_idx = next((i for i, h in enumerate(header_lower) if 'step expected' in h), None)

    steps = []
    for row in rows[1:]:
        if not any(c.strip() for c in row):
            continue
        id_col   = row[0].strip() if len(row) > 0 else ""
        step_num = (row[step_num_idx].strip() if step_num_idx is not None and len(row) > step_num_idx else "")
        if not id_col and step_num and step_num.isdigit() and len(row) > step_action_idx:
            action = row[step_action_idx].strip()
            if not action:
                continue
            # Beklenen sonucu kısa tut (ilk cümle, max 120 karakter)
            expected = ""
            if step_expected_idx is not None and len(row) > step_expected_idx:
                raw_exp = row[step_expected_idx].strip()
                # İlk cümleyi al
                first_sentence = re.split(r'(?<=[.!?])\s', raw_exp)[0]
                expected = first_sentence[:120]
            if expected:
                steps.append(f"Adım {step_num}: {action} → Beklenen: {expected}")
            else:
                steps.append(f"Adım {step_num}: {action}")

    return "\n".join(steps) if steps else None


# ── Locator dosyası (kayıt modu XPath'leri) ───────────────────────────────────
# Seçiliyse _rule_mobile çıktısındaki tıklama/yazma hedefleri dosyada aranır:
# bulunursa önce kayıtlı XPath denenir, eleman yoksa eski akıllı arama (metin,
# kaydırma, vision) aynen çalışır — locator hiçbir zaman tek yol değildir.
from contextvars import ContextVar
import ast as _ast
from services import locator_service

_MOBILE_LOCATORS: ContextVar[dict | None] = ContextVar("ai_gen_mobile_locators", default=None)


def _locator_lookup(user_id: int, db: Session, name: Optional[str]) -> dict | None:
    name = (name or "").strip()
    if not name:
        return None
    ws = _get_user_workspace(user_id, db)
    if not ws:
        return None
    try:
        return locator_service.load_lookup(ws.local_path, name)
    except (FileNotFoundError, ValueError) as e:
        print(f"[LOCATOR] {e}")
        return None


def _call_args(line: str, fn: str):
    """'_safe_click(driver, 'X', 'Y')' → (girinti, [argüman kaynak metinleri], [sabit değerler|None])"""
    m = re.match(rf"^(\s*){fn}\(driver, .*\)\s*$", line)
    if not m:
        return None
    src = line.strip()
    try:
        call = _ast.parse(src, mode="eval").body
    except SyntaxError:
        return None
    args = call.args[1:]
    return (m.group(1), [_ast.get_source_segment(src, a) for a in args],
            [a.value if isinstance(a, _ast.Constant) and isinstance(a.value, str) else None for a in args])


def _apply_locators_mobile(step: str, lines: list[str] | None, lookup: dict | None) -> tuple[list | None, bool]:
    if not lines or not lookup:
        return lines, False
    out, hit = [], False
    for ln in lines:
        c = _call_args(ln, "_safe_click")
        if c:
            ind, srcs, vals = c
            xp = None
            for t in vals:
                if t and not t.startswith(("//", "id/", "@id/")):
                    xp = locator_service.resolve_xpath(lookup, t, "tap")
                    if xp:
                        break
            if xp:
                out += [f"{ind}try:",
                        f"{ind}    driver.find_element(AppiumBy.XPATH, {xp!r}).click()",
                        f"{ind}    print({('[OK] Locator XPath tıklandı: ' + xp)!r})",
                        f"{ind}except Exception:",
                        f"{ind}    _safe_click(driver, {', '.join(srcs)})"]
                hit = True
                continue
        c = _call_args(ln, "_safe_assert")
        if c and c[2] and c[2][0]:
            # Sadece yazısız (ikon) kayıtlar: yazılı hedefte doğrulama yazının kendisini arar
            xp = locator_service.resolve_xpath(lookup, c[2][0], "assert")
            if xp:
                out.append(f"{c[0]}_safe_assert(driver, {xp!r})")
                hit = True
                continue
        c = _call_args(ln, "_safe_input")
        if c and len(c[1]) == 2:
            ind, srcs, _ = c
            q = _extract_quoted(_clean_step(step))
            xp = locator_service.resolve_xpath(lookup, q[0], "type") if len(q) >= 2 else None
            if xp:
                out += [f"{ind}_lx = {xp!r} if driver.find_elements(AppiumBy.XPATH, {xp!r}) else {srcs[0]}",
                        f"{ind}_safe_input(driver, _lx, {srcs[1]})"]
                hit = True
                continue
        out.append(ln)
    return out, hit


def _rule_mobile(step: str) -> list[str] | None:
    """
    Tek bir Türkçe adımı Appium satırlarına çevirir.
    Tanıyamazsa None döner → LLM devreye girer.
    """
    s = _clean_step(step)
    if not s:
        return None

    quoted = _extract_quoted(s)
    sl = s.lower()

    # ── {{varname}} "//xpath" yaz ─────────────────────────────────────────────
    _vm = re.match(r'^\{\{(\w+)\}\}(?:\s+"([^"]+)")?\s*(?:yaz|gir|doldur|type\b)?\.?\s*$', s, re.IGNORECASE)
    if _vm:
        vn, xpath = _vm.group(1), (_vm.group(2) or '//android.widget.EditText[1]')
        return [f'# {{{{ {vn} }}}} yaz', f"_safe_input(driver, {repr(xpath)}, str(_VARS.get({repr(vn)}, '')))"]

    # ── "hedef" oku {{varname}} — elementten değer oku ────────────────────────
    _rm = re.match(r'^"([^"]+)"\s*oku\s*\{\{(\w+)\}\}\.?\s*$', s, re.IGNORECASE)
    if _rm:
        hint, vn = _rm.group(1), _rm.group(2)
        return [
            f'# "{hint}" oku {{{{ {vn} }}}}',
            f"_VARS[{repr(vn)}] = _safe_read(driver, {repr(hint)})",
            f"if not _VARS[{repr(vn)}]: raise ValueError({repr('Deger okunamadi: ' + hint)})",
            f"print(f'[OKU→_VARS] {vn} = {{_VARS[{repr(vn)}]}}')",
        ]

    # ── Ok ikonları: '"X" yanındaki aşağı ok(una) tıkla', 'Aşağı oka tıkla' ──────
    # Tırnaklı metin oka YAKINLIK için kullanılır, tıklanacak hedef değildir. Fiil
    # yazılmasa da ("... aşağı ok butonu") doğrulama fiili yoksa tıklama sayılır.
    # Evren'de oklar "<" / ">" gibi chevron ikonlarıyla eğitildi → adımda sembolün
    # kendisi de yazılabilir: '"<" iconuna tıkla', '"Takımları Gör" yanındaki > ikonu'
    _ARROW_SYMS = {"<": "sol", "‹": "sol", "←": "sol", ">": "sağ", "›": "sağ", "→": "sağ",
                   "˅": "aşağı", "⌄": "aşağı", "▼": "aşağı", "↓": "aşağı",
                   "^": "yukarı", "˄": "yukarı", "⌃": "yukarı", "▲": "yukarı", "↑": "yukarı"}
    _SM = re.search(r'(?:^|\s)["“”\']?([<>‹›←→˅⌄▼↓^˄⌃▲↑])["“”\']?\s+(?:[iı]con|[iı]kon|simge|buton|ok)\w*', s, re.IGNORECASE)
    _AM = re.search(r'\b(sol|sağ|sag|aşağı|asagi|yukarı|yukari)\s+ok(?:a|u|un|una|lar\w*)?\b', sl)
    if (_SM or _AM) and not re.search(r'g[öo]r[üu]l[üu]r|g[öo]r[üu]nmeli|kontrol|do[ğg]rula', sl):
        if _SM:
            _dir = _ARROW_SYMS[_SM.group(1)]
        else:
            _dir = {"sag": "sağ", "asagi": "aşağı", "yukari": "yukarı"}.get(_AM.group(1), _AM.group(1))
        _near = next((q for q in quoted if q.strip() not in _ARROW_SYMS), "")
        if not _near:
            _nm = re.match(r'^(.+?)\s+(?:yan[ıi]ndaki|sa[ğg][ıi]ndaki|solundaki)\b', s, re.IGNORECASE)
            _near = _nm.group(1).strip() if _nm else ""
        _phrase = (_near + " " if _near else "") + _dir + " ok"
        return [f'# {s}', f"_safe_click_arrow(driver, {repr(_phrase)})", 'time.sleep(0.8)']

    # ── Pasif Türkçe: "X [suffix] tıklanır" (tırnaksız) ─────────────────────
    _PC = re.search(
        r'^(.+?)\s+(?:seçene[ğg]ine|componentine|butonu?na|simgesine|ikonuna'
        r'|bağlantısına|sekmesin[ei]|tab[ıi]na|alan[ıi]na|[üu]zerine'
        r'|linkine|ikona|logosuna)\s+t[ıi]klan[ıi]r\.?\s*$',
        s, re.IGNORECASE
    )
    if _PC:
        full_t = _PC.group(1).strip()
        key_t  = _extract_key_target(full_t)
        if key_t:
            return [f'# {s}', f"_safe_click(driver, {repr(full_t)}, {repr(key_t)})", 'time.sleep(0.8)']
        return [f'# {s}', f"_safe_click(driver, {repr(full_t)})", 'time.sleep(0.8)']

    # ── "Uygulama açılır" — uygulama zaten başlatıldı, sadece yorum ─────────
    if re.search(r'^uygulama\s+a[çc][ıi]l[ıi]r\.?\s*$', sl):
        return [f'# {s}']

    # ── "X açılır" - bileşen/ekran açılır (click to open) ────────────────────
    _AC = re.search(r'^(.+?)\s+a[çc][ıi]l[ıi]r\.?\s*$', s, re.IGNORECASE)
    if _AC:
        full_t = _AC.group(1).strip()
        key_t  = _extract_key_target(full_t)
        if key_t:
            return [f'# {s}', f"_safe_click(driver, {repr(full_t)}, {repr(key_t)})", 'time.sleep(1.0)']
        return [f'# {s}', f"_safe_click(driver, {repr(full_t)})", 'time.sleep(1.0)']

    # ── Compound: "seçim yapılır, bir X girilir" — hem seç hem yaz ──────────────
    if re.search(r'se[çc]im\s+yapıl[ıi]r', sl) and re.search(r'bir\s+\w+\s+girilir', sl):
        _m_amt = re.search(r'bir\s+(tutar|miktar|değer|numara|sayı)\s+girilir', sl, re.IGNORECASE)
        value = _get_auto_value(_m_amt.group(1) if _m_amt else 'tutar')
        return [
            f'# {s}',
            '# Dropdown seçimi: ilk mevcut seçeneği tıkla',
            'try:',
            '    _ddl = driver.find_elements(AppiumBy.XPATH, \'//android.widget.ListView//*[@clickable="true"]\')',
            '    if not _ddl:',
            '        _ddl = driver.find_elements(AppiumBy.XPATH, \'//*[@clickable="true" and string-length(@text)>0]\')',
            '    if _ddl: _ddl[0].click(); time.sleep(0.5)',
            'except Exception as _e: print(f"Seçim hatası: {_e}")',
            f"# Tutar girişi — otomatik test verisi: {value}",
            f"_safe_input(driver, '//android.widget.EditText[1]', '{value}')",
            'time.sleep(0.5)',
        ]

    # ── "bir X girilir" — basit adım, değer belirtilmemiş, otomatik doldur ─────
    _GR = re.search(r'\bgirilir\b', sl)
    if _GR and not quoted:
        _GM = re.search(r'(?:bir\s+)?(.+?)\s+(?:alan[ıi]na\s+)?girilir\.?\s*$', s, re.IGNORECASE)
        if _GM:
            ctx = _GM.group(1)
            value = _get_auto_value(ctx)
            is_pw = bool(re.search(r'şifre|sifre|password|pin|parola', ctx, re.IGNORECASE))
            field_xpath = ('//android.widget.EditText[@password="true"]'
                           if is_pw else '//android.widget.EditText[1]')
            return [
                f'# {s}',
                f"# Otomatik test verisi: {value}",
                f"_safe_input(driver, '{field_xpath}', '{value}')",
                'time.sleep(0.5)',
            ]

    # ── Tıklama ───────────────────────────────────────────────────────────────
    CLICK_WORDS = r'(?:t[iı]kla|tikla|bas|click|se[cç]|tap|a[cç])'
    _ASSERT_PRE = r'(?:g[oö]r[uü]l[uü]r|g[oö]r[uü]n[uü]r|g[oö]r[uü]nt[uü]len|do[gğ]rula|kontrol)'
    _sl_noq = re.sub(r'"[^"]*"', '', sl)
    if re.search(CLICK_WORDS, sl) and quoted and not (re.search(_ASSERT_PRE, _sl_noq) and not re.search(CLICK_WORDS, _sl_noq)):
        text = quoted[0]
        return [
            f'# {s}',
            f'_safe_click(driver, {repr(text)})',
            'time.sleep(0.8)',
        ]

    # ── Doğrulama / Görme — ÖNCE kontrol et (gir/giriş çakışmasını önle) ─────
    ASSERT_WORDS = r'(?:g[oö]r[uü]l[uü]r|g[oö]r[uü]n[uü]r|g[oö]r[uü]nt[uü]len|g[oö]r\b|do[gğ]rula|kontrol|assert|verify|check|var m[iı]|bekleniyor)'
    if re.search(ASSERT_WORDS, sl) and quoted:
        text = quoted[0]
        return [
            f'# {s}',
            f'_safe_assert(driver, {text!r})',
        ]

    # ── Metin girişi ──────────────────────────────────────────────────────────
    # "gir" kelimesi sadece bağımsız geçiyorsa eşleşir (giriş/giriyor değil)
    INPUT_WORDS = r'(?:yaz\b|(?<![a-zğışöüçA-ZĞİŞÖÜÇ])gir(?!iş|iyor|di)|doldur|type|enter\b|input\b)'
    if re.search(INPUT_WORDS, sl) and quoted:
        value = quoted[-1]  # Son tırnak değeri = yazılacak metin
        # Şifre alanı mı?
        if re.search(r'(?:şifre|sifre|password|pin|parola)', sl):
            field_xpath = '//android.widget.EditText[@password="true"]'
        elif len(quoted) >= 2 and not quoted[0].startswith("//"):
            # '"Tutar" alanına "100" yaz' → alan ADIYLA bul (çok alanlı formlarda ilk
            # EditText'e yazılmasın): alanın kendi ipucu/metni ya da önündeki etiket
            lbl = quoted[0].replace('"', "'")
            field_xpath = (f'(//android.widget.EditText[@text="{lbl}" or @hint="{lbl}"]'
                           f' | //*[@text="{lbl}" or @content-desc="{lbl}"]/following::android.widget.EditText[1])[1]')
        else:
            field_xpath = '//android.widget.EditText[1]'
        return [
            f'# {s}',
            f'_safe_input(driver, {field_xpath!r}, {value!r})',
            'time.sleep(0.5)',
        ]

    # ── Doğrulama (altta kalır — ASSERT önceden yakalandı) ───────────────────
    if False and quoted:  # devre dışı - yukarıda ele alındı
        text = quoted[0]
        return [
            f'# {s}',
            f'_safe_assert(driver, \'{text}\')',
        ]

    # ── Kaydırma ──────────────────────────────────────────────────────────────
    if re.search(r'(?:a[sş]a[gğ][iı]\s*kaydır|scroll\s*down|swipe\s*down)', sl):
        return [f'# {s}', 'driver.swipe(540, 1400, 540, 400, 800)', 'time.sleep(0.5)']
    if re.search(r'(?:yukar[iı]\s*kaydır|scroll\s*up|swipe\s*up)', sl):
        return [f'# {s}', 'driver.swipe(540, 400, 540, 1400, 800)', 'time.sleep(0.5)']
    if re.search(r'(?:sa[gğ]a\s*kaydır|swipe\s*right|scroll\s*right)', sl):
        return [f'# {s}', 'driver.swipe(900, 800, 150, 800, 800)', 'time.sleep(0.5)']
    if re.search(r'(?:sola\s*kaydır|swipe\s*left|scroll\s*left)', sl):
        return [f'# {s}', 'driver.swipe(150, 800, 900, 800, 800)', 'time.sleep(0.5)']
    if re.search(r'(?:kaydır|scroll|swipe)', sl):
        return [f'# {s}', 'driver.swipe(540, 1400, 540, 400, 800)', 'time.sleep(0.5)']

    # ── Geri ─────────────────────────────────────────────────────────────────
    if re.search(r'(?:geri\s*d[oö]n|geri\s*git|back\b)', sl):
        return [f'# {s}', 'driver.back()', 'time.sleep(0.5)']

    # ── Klavye kapat ──────────────────────────────────────────────────────────
    if re.search(r'(?:klavye\s*kapat|klavyeyi\s*kapat|kapat\s*klavye|hide\s*keyboard|dismiss\s*keyboard)', sl):
        return [
            f'# {s}',
            'try:',
            '    driver.hide_keyboard()',
            'except Exception:',
            '    driver.back()',
            'time.sleep(0.5)',
        ]

    # ── Bekleme ───────────────────────────────────────────────────────────────
    if re.search(r'(?:bekle|wait|sleep)', sl):
        m = re.search(r'(\d+)', s)
        secs = m.group(1) if m else '2'
        return [f'# {s}', f'time.sleep({secs})']

    # ── Tırnaksız doğrulama/durum (görülür, açıldı, kapandı — quoted olmadan) ─
    ASSERT_WORDS = r'(?:g[oö]r[uü]l[uü]r|g[oö]r[uü]n[uü]r|g[oö]r[uü]nt[uü]len|g[oö]r\b|do[gğ]rula|kontrol|assert|verify|check|var m[iı]|bekleniyor)'
    if re.search(ASSERT_WORDS, sl) and not quoted:
        return [f'# {s}  # assertion: eleman metni belirtilmemiş, manuel doğrula']

    return None  # Tanınmadı → LLM


def _rule_web(step: str) -> list[str] | None:
    s  = _clean_step(step)
    if not s:
        return None
    sl     = s.lower()
    quoted = _extract_quoted(s)

    # ── {{varname}} "alanHint" yaz ────────────────────────────────────────────
    _vm = re.match(r'^\{\{(\w+)\}\}(?:\s+"([^"]+)")?\s*(?:yaz|gir|doldur|type\b)?\.?\s*$', s, re.IGNORECASE)
    if _vm:
        vn, hint = _vm.group(1), (_vm.group(2) or "")
        if hint:
            return [f'# {{{{ {vn} }}}} yaz', f"_safe_input_web(driver, {repr(hint)}, str(_VARS.get({repr(vn)}, '')))"]
        return [f'# {{{{ {vn} }}}} yaz', f"import pyperclip as _pc; _pc.copy(str(_VARS.get({repr(vn)}, ''))); import pyautogui as _pa; _pa.hotkey('ctrl','v')"]

    # ── "hedef" oku {{varname}} — elementten değer oku ────────────────────────
    _rm = re.match(r'^"([^"]+)"\s*oku\s*\{\{(\w+)\}\}\.?\s*$', s, re.IGNORECASE)
    if _rm:
        hint, vn = _rm.group(1), _rm.group(2)
        return [
            f'# "{hint}" oku {{{{ {vn} }}}}',
            f"_VARS[{repr(vn)}] = _safe_read_web(driver, {repr(hint)})",
            f"if not _VARS[{repr(vn)}]: raise ValueError({repr('Deger okunamadi: ' + hint)})",
            f"print(f'[OKU→_VARS] {vn} = {{_VARS[{repr(vn)}]}}')",
        ]

    _WEB_KEY_MAP = {
        'ctrl': 'Keys.CONTROL', 'control': 'Keys.CONTROL',
        'alt': 'Keys.ALT', 'shift': 'Keys.SHIFT',
        'delete': 'Keys.DELETE', 'del': 'Keys.DELETE',
        'backspace': 'Keys.BACK_SPACE', 'bs': 'Keys.BACK_SPACE',
        'enter': 'Keys.ENTER', 'return': 'Keys.RETURN',
        'tab': 'Keys.TAB',
        'escape': 'Keys.ESCAPE', 'esc': 'Keys.ESCAPE',
        'home': 'Keys.HOME', 'end': 'Keys.END',
        'pageup': 'Keys.PAGE_UP', 'pagedown': 'Keys.PAGE_DOWN',
        'f1': 'Keys.F1', 'f2': 'Keys.F2', 'f3': 'Keys.F3', 'f4': 'Keys.F4',
        'f5': 'Keys.F5', 'f6': 'Keys.F6', 'f7': 'Keys.F7', 'f8': 'Keys.F8',
        'f9': 'Keys.F9', 'f10': 'Keys.F10', 'f11': 'Keys.F11', 'f12': 'Keys.F12',
        'space': 'Keys.SPACE',
        'up': 'Keys.ARROW_UP', 'down': 'Keys.ARROW_DOWN',
        'left': 'Keys.ARROW_LEFT', 'right': 'Keys.ARROW_RIGHT',
    }
    _WEB_MODS = {'ctrl', 'control', 'alt', 'shift'}

    def _web_key_lines(key_str):
        parts = [p.strip().lower().replace(' ', '') for p in key_str.split('+')]
        mods  = [p for p in parts if p in _WEB_MODS]
        keys  = [p for p in parts if p not in _WEB_MODS]
        ln = [
            'from selenium.webdriver.common.keys import Keys',
            'from selenium.webdriver.common.action_chains import ActionChains',
            '_ac_web = ActionChains(driver)',
        ]
        for m in mods:
            ln.append(f'_ac_web.key_down({_WEB_KEY_MAP[m]})')
        for k in keys:
            mapped = _WEB_KEY_MAP.get(k)
            ln.append(f'_ac_web.send_keys({mapped})' if mapped else f'_ac_web.send_keys({repr(k)})')
        for m in reversed(mods):
            ln.append(f'_ac_web.key_up({_WEB_KEY_MAP[m]})')
        ln.append('_ac_web.perform()')
        return ln

    # ── Kısayol tuşu: "Ctrl+A" bas / "Delete" bas ────────────────────────
    hk = re.match(r'^["\']([^"\']+)["\']\s*(?:bas|kısayol|hotkey|tu[sş])$', s, re.IGNORECASE)
    if hk:
        return [f'# {s}'] + _web_key_lines(hk.group(1).strip()) + ['time.sleep(0.2)']

    # ── Tekil tuş (tırnaksız) ─────────────────────────────────────────────
    if re.search(r'\benter\s*bas\b', sl):
        return [f'# {s}'] + _web_key_lines('enter') + ['time.sleep(0.3)']
    if re.search(r'\btab\s*bas\b', sl):
        return [f'# {s}'] + _web_key_lines('tab')   + ['time.sleep(0.2)']
    if re.search(r'(?:escape|esc)\s*bas\b', sl):
        return [f'# {s}'] + _web_key_lines('escape') + ['time.sleep(0.2)']
    if re.search(r'(?:delete(?:\s*tuşuna?)?|sil(?:\s*tuşuna?)?)\s*bas\b', sl):
        return [f'# {s}'] + _web_key_lines('delete') + ['time.sleep(0.2)']
    if re.search(r'backspace\s*bas\b', sl):
        return [f'# {s}'] + _web_key_lines('backspace') + ['time.sleep(0.2)']
    if re.search(r'geri\s*sil', sl):
        return [f'# {s}'] + _web_key_lines('backspace') + ['time.sleep(0.2)']

    # ── Ok tuşları: "sağ ok (tuşuna) bas" ──────────────────────────────────
    _arrow_m = re.search(r'(sağ|sag|sol|yukarı|yukari|aşağı|asagi)\s*ok(?:\s*tuşuna?)?\s*bas', sl)
    if _arrow_m:
        _arrow_map = {'sağ': 'right', 'sag': 'right', 'sol': 'left',
                      'yukarı': 'up', 'yukari': 'up', 'aşağı': 'down', 'asagi': 'down'}
        return [f'# {s}'] + _web_key_lines(_arrow_map[_arrow_m.group(1)]) + ['time.sleep(0.2)']

    if re.search(r'çift\s*boşluk|cift\s*bosluk', sl):
        return [f'# {s}'] + _web_key_lines('space') + _web_key_lines('space') + ['time.sleep(0.2)']
    if re.search(r'boşluk(?:\s*tuşuna?)?\s*bas|bosluk\s*tusuna?\s*bas', sl):
        return [f'# {s}'] + _web_key_lines('space') + ['time.sleep(0.2)']
    if re.search(r'home\s*tuşuna?\s*bas', sl):
        return [f'# {s}'] + _web_key_lines('home') + ['time.sleep(0.2)']
    if re.search(r'end\s*tuşuna?\s*bas', sl):
        return [f'# {s}'] + _web_key_lines('end') + ['time.sleep(0.2)']
    if re.search(r'page\s*up', sl):
        return [f'# {s}'] + _web_key_lines('pageup') + ['time.sleep(0.2)']
    if re.search(r'page\s*down', sl):
        return [f'# {s}'] + _web_key_lines('pagedown') + ['time.sleep(0.2)']
    if re.search(r'tümünü\s*seç|tumunu\s*sec', sl):
        return [f'# {s}'] + _web_key_lines('ctrl+a') + ['time.sleep(0.2)']
    if re.search(r'\bkopyala\b', sl):
        return [f'# {s}'] + _web_key_lines('ctrl+c') + ['time.sleep(0.2)']
    if re.search(r'yapıştır|yapistir', sl):
        return [f'# {s}'] + _web_key_lines('ctrl+v') + ['time.sleep(0.2)']

    CLICK_WORDS  = r'(?:t[iı]kla|tikla|click|tap)'
    INPUT_WORDS  = r'(?:yaz\b|(?<![a-zğışöüçA-ZĞİŞÖÜÇ])gir(?!iş|iyor)|doldur|type\b)'
    ASSERT_WORDS = r'(?:g[oö]r[uü]l[uü]r|g[oö]r[uü]n[uü]r|g[oö]r[uü]nt[uü]len|g[oö]r\b|do[gğ]rula|kontrol|assert|verify|var m[iı]|bekleniyor|görünmeli)'
    NAV_WORDS    = r'(?:git\b|navigate|a[cç]\b|open\b|ziyaret|url|adres)'

    # ── URL / Navigasyon ─────────────────────────────────────────────────
    if re.search(NAV_WORDS, sl) and quoted:
        url = quoted[0]
        if not url.startswith("http"):
            url = "https://" + url
        return [f'# {s}', f'driver.get("{url}")', 'time.sleep(1.5)']

    # ── Tıklama ──────────────────────────────────────────────────────────
    if re.search(CLICK_WORDS, sl) and quoted:
        text = quoted[0]
        return [f'# {s}', f'_safe_click_web(driver, \'{text}\')', 'time.sleep(0.5)']

    # ── Doğrulama ────────────────────────────────────────────────────────
    if re.search(ASSERT_WORDS, sl) and quoted:
        text = quoted[0]
        return [f'# {s}', f'_safe_assert_web(driver, \'{text}\')']

    # ── Element görünene kadar bekle ──────────────────────────────────────
    if re.search(r'g[oö]r[uü]n[eien]+.*kadar\s*bekle|appear.*wait', sl) and quoted:
        text = quoted[0]
        return [
            f'# {s}',
            f"WebDriverWait(driver, 15).until(lambda d: len(d.find_elements(By.XPATH, '//*[contains(.,\"{text}\")]')) > 0)",
            'time.sleep(0.3)',
        ]

    # ── Metin girişi ─────────────────────────────────────────────────────
    if re.search(INPUT_WORDS, sl) and quoted:
        value  = quoted[-1]
        is_pw  = bool(re.search(r'(?:şifre|sifre|password|pin|parola)', sl))
        # Alan ipucu: tırnak içlerini ve fiili sil, kalan metin alanı tanımlar
        hint_raw = re.sub(
            r'["\u201c\u201d\u2018\u2019][^"\u201c\u201d\u2018\u2019]+["\u201c\u201d\u2018\u2019]',
            '', s
        )
        hint_raw = re.sub(
            r'\b(?:yaz|gir|doldur|type|enter|alan[ıi]na|alan[ıi]|alanına|field|kutusuna|bölümüne)\b',
            '', hint_raw, flags=re.IGNORECASE
        )
        hint = hint_raw.strip().rstrip('.,;:').strip().lower()
        return [
            f'# {s}',
            f'_safe_input_web(driver, \'{hint}\', \'{value}\', is_password={is_pw})',
            'time.sleep(0.3)',
        ]

    # ── Kaydırma ─────────────────────────────────────────────────────────
    if re.search(r'(?:a[sş]a[gğ][iı]|scroll\s*down)', sl):
        return [f'# {s}', 'driver.execute_script("window.scrollBy(0, 400);")', 'time.sleep(0.4)']
    if re.search(r'(?:yukar[iı]|scroll\s*up)', sl):
        return [f'# {s}', 'driver.execute_script("window.scrollBy(0, -400);")', 'time.sleep(0.4)']
    if re.search(r'(?:sa[gğ]a\s*kaydır|scroll\s*right)', sl):
        return [f'# {s}', 'driver.execute_script("window.scrollBy(400, 0);")', 'time.sleep(0.4)']
    if re.search(r'(?:sola\s*kaydır|scroll\s*left)', sl):
        return [f'# {s}', 'driver.execute_script("window.scrollBy(-400, 0);")', 'time.sleep(0.4)']

    # ── Sayfa işlemleri ───────────────────────────────────────────────────
    if re.search(r'(?:yenile|refresh|f5)', sl):
        return [f'# {s}', 'driver.refresh()', 'time.sleep(1.5)']
    if re.search(r'(?:geri|back\b)', sl):
        return [f'# {s}', 'driver.back()', 'time.sleep(0.5)']
    if re.search(r'(?:ileri|forward\b)', sl):
        return [f'# {s}', 'driver.forward()', 'time.sleep(0.5)']
    if re.search(r'(?:ekran\s*g[oö]r[uü]nt[uü]s[uü]|screenshot)', sl):
        return [f'# {s}', 'driver.save_screenshot(f"screenshot_{int(time.time())}.png")', 'time.sleep(0.3)']

    # ── Bekleme ───────────────────────────────────────────────────────────
    if re.search(r'(?:bekle|wait|sleep)', sl):
        m = re.search(r'(\d+)', s)
        secs = m.group(1) if m else '2'
        return [f'# {s}', f'time.sleep({secs})']

    # ── Tırnaksız doğrulama/durum (görülür, açıldı — quoted olmadan) ─────────
    if re.search(ASSERT_WORDS, sl) and not quoted:
        return [f'# {s}  # assertion: eleman metni belirtilmemiş, manuel doğrula']

    return None


def _rule_desktop(step: str) -> list[str] | None:
    """Türkçe / İngilizce desktop adımını pyautogui helper çağrılarına çevirir."""
    s  = _clean_step(step)
    if not s:
        return None
    sl     = s.lower()
    quoted = _extract_quoted(s)

    # ── {{varname}} yaz — değişkenden değer yaz ──────────────────────────────
    _vm = re.match(r'^\{\{(\w+)\}\}\s*(?:yaz|gir|doldur|type\b)?\.?\s*$', s, re.IGNORECASE)
    if _vm:
        vn = _vm.group(1)
        return [f'# {{{{ {vn} }}}} yaz', f"_safe_type_desktop(str(_VARS.get({repr(vn)}, '')))"]

    CLICK_WORDS        = r'(?:t[iı]kla|tikla|bas|click|tap|se[çc]|dokun)'
    DOUBLE_CLICK_WORDS = r'(?:çift\s*t[iı]kla|double[\s_]click)'
    RIGHT_CLICK_WORDS  = r'(?:sağ\s*t[iı]kla|right[\s_]click)'
    INPUT_WORDS        = r'(?:yaz\b|(?<![a-zğışöüçA-ZĞİŞÖÜÇ])gir(?!iş|iyor)|doldur|type\b)'
    ASSERT_WORDS       = r'(?:g[oö]r[uü]l[uü]r|g[oö]r[uü]n[uü]r|g[oö]r[uü]nt[uü]len|g[oö]r\b|do[gğ]rula|kontrol|assert|verify|var m[iı]|görünmeli|bekleniyor)'

    # ── Kısayol tuşu: "Ctrl+S" bas ───────────────────────────────────────────
    hk = re.match(r'^["\']([^"\']+)["\']\s*(?:bas|kısayol|hotkey|tu[sş])$', s, re.IGNORECASE)
    if hk:
        keys = "', '".join(hk.group(1).lower().replace(' ', '').split('+'))
        return [f'# {s}', f"_pag.hotkey('{keys}')", 'time.sleep(0.3)']

    # ── Enter / Tab / Escape ──────────────────────────────────────────────────
    if re.search(r'\benter\b', sl):
        return [f'# {s}', '_pag.press("enter")', 'time.sleep(0.3)']
    if re.search(r'\btab\b', sl):
        return [f'# {s}', '_pag.press("tab")',   'time.sleep(0.2)']
    if re.search(r'\bescape|esc\b', sl):
        return [f'# {s}', '_pag.press("esc")',   'time.sleep(0.2)']

    # ── Ok tuşları: "sağ ok (tuşuna) bas" ──────────────────────────────────
    _arrow_m = re.search(r'(sağ|sag|sol|yukarı|yukari|aşağı|asagi)\s*ok(?:\s*tuşuna?)?\s*bas', sl)
    if _arrow_m:
        _arrow_map = {'sağ': 'right', 'sag': 'right', 'sol': 'left',
                      'yukarı': 'up', 'yukari': 'up', 'aşağı': 'down', 'asagi': 'down'}
        return [f'# {s}', f'_pag.press("{_arrow_map[_arrow_m.group(1)]}")', 'time.sleep(0.2)']

    if re.search(r'çift\s*boşluk|cift\s*bosluk', sl):
        return [f'# {s}', '_pag.press("space")', 'time.sleep(0.05)', '_pag.press("space")', 'time.sleep(0.2)']
    if re.search(r'boşluk(?:\s*tuşuna?)?\s*bas|bosluk\s*tusuna?\s*bas', sl):
        return [f'# {s}', '_pag.press("space")', 'time.sleep(0.2)']
    if re.search(r'delete(?:\s*tuşuna?)?\s*bas|\bsil\s*tuşuna?\s*bas', sl):
        return [f'# {s}', '_pag.press("delete")', 'time.sleep(0.2)']
    if re.search(r'backspace|geri\s*sil', sl):
        return [f'# {s}', '_pag.press("backspace")', 'time.sleep(0.2)']
    if re.search(r'home\s*tuşuna?\s*bas', sl):
        return [f'# {s}', '_pag.press("home")', 'time.sleep(0.2)']
    if re.search(r'end\s*tuşuna?\s*bas', sl):
        return [f'# {s}', '_pag.press("end")', 'time.sleep(0.2)']
    if re.search(r'page\s*up', sl):
        return [f'# {s}', '_pag.press("pageup")', 'time.sleep(0.2)']
    if re.search(r'page\s*down', sl):
        return [f'# {s}', '_pag.press("pagedown")', 'time.sleep(0.2)']
    if re.search(r'tümünü\s*seç|tumunu\s*sec', sl):
        return [f'# {s}', '_pag.hotkey("ctrl", "a")', 'time.sleep(0.2)']

    # ── Kopyala / Yapıştır ────────────────────────────────────────────────────
    if re.search(r'(?:kopyala|copy)', sl) and quoted:
        text = quoted[0]
        return [f'# {s}', '_refocus_window()', f"_safe_click_desktop('{text}', _WINDOW_TITLE)", '_pag.hotkey("ctrl", "c")', 'time.sleep(0.3)']
    if re.search(r'(?:kopyala|copy)', sl):
        return [f'# {s}', '_pag.hotkey("ctrl", "c")', 'time.sleep(0.2)']
    if re.search(r'(?:yapıştır|paste)', sl):
        return [f'# {s}', '_refocus_window()', '_pag.hotkey("ctrl", "v")', 'time.sleep(0.3)']

    # ── Çift tıklama (CLICK_WORDS'ten önce kontrol edilmeli) ──────────────────
    if re.search(DOUBLE_CLICK_WORDS, sl) and quoted:
        text = quoted[0]
        return [f'# {s}', '_refocus_window()', f"_safe_double_click_desktop('{text}', _WINDOW_TITLE)", 'time.sleep(0.5)']

    # ── Sağ tıklama ───────────────────────────────────────────────────────────
    if re.search(RIGHT_CLICK_WORDS, sl) and quoted:
        text = quoted[0]
        return [f'# {s}', '_refocus_window()', f"_safe_right_click_desktop('{text}', _WINDOW_TITLE)", 'time.sleep(0.5)']

    # ── Menü navigasyonu: "Dosya > Yeni" tıkla ───────────────────────────────
    if re.search(CLICK_WORDS, sl) and quoted and '>' in quoted[0]:
        items = [i.strip() for i in quoted[0].split('>')]
        return [f'# {s}', '_refocus_window()', f"_safe_menu_click_desktop({items!r}, _WINDOW_TITLE)", 'time.sleep(0.5)']

    # ── Tıklama ───────────────────────────────────────────────────────────────
    if re.search(CLICK_WORDS, sl) and quoted:
        text = quoted[0]
        return [f'# {s}', '_refocus_window()', f"_safe_click_desktop('{text}', _WINDOW_TITLE)", 'time.sleep(0.5)']

    # ── Doğrulama ─────────────────────────────────────────────────────────────
    if re.search(ASSERT_WORDS, sl) and quoted:
        text = quoted[0]
        return [f'# {s}', '_refocus_window()', f"_safe_assert_desktop('{text}')"]

    # ── Metin girişi ──────────────────────────────────────────────────────────
    if re.search(INPUT_WORDS, sl) and quoted:
        value = quoted[-1]
        return [f'# {s}', '_refocus_window()', f"_safe_type_desktop('{value}')", 'time.sleep(0.3)']

    # ── Element görünene kadar bekle (genel bekle'den önce) ───────────────────
    if re.search(r'g[oö]r[uü]n[eu]ne\s*kadar\s*bekle|wait.*(?:appear|visible)', sl) and quoted:
        text = quoted[0]
        m = re.search(r'(\d+)', s)
        timeout = m.group(1) if m else '10'
        return [f'# {s}', f"_safe_wait_for_desktop('{text}', timeout={timeout})"]

    # ── Bekleme ───────────────────────────────────────────────────────────────
    if re.search(r'(?:bekle|wait|sleep)', sl):
        m = re.search(r'(\d+)', s)
        return [f'# {s}', f'time.sleep({m.group(1) if m else "1"})']

    # ── Kaydırma ─────────────────────────────────────────────────────────────
    if re.search(r'(?:a[sş]a[gğ]|scroll\s*down|swipe\s*down)', sl):
        return [f'# {s}', '_pag.scroll(-5)', 'time.sleep(0.3)']
    if re.search(r'(?:yukar[iı]|scroll\s*up|swipe\s*up)', sl):
        return [f'# {s}', '_pag.scroll(5)',  'time.sleep(0.3)']
    if re.search(r'(?:sa[gğ]a\s*kaydır|scroll\s*right|swipe\s*right)', sl):
        return [f'# {s}', "_pag.keyDown('shift')", '_pag.scroll(-5)', "_pag.keyUp('shift')", 'time.sleep(0.3)']
    if re.search(r'(?:sola\s*kaydır|scroll\s*left|swipe\s*left)', sl):
        return [f'# {s}', "_pag.keyDown('shift')", '_pag.scroll(5)', "_pag.keyUp('shift')", 'time.sleep(0.3)']

    # ── Ekran görüntüsü ───────────────────────────────────────────────────────
    if re.search(r'ekran\s*g[oö]r[uü]nt[uü]s[uü]|screenshot', sl):
        return [f'# {s}', f'_pag.screenshot(f"screenshot_{{int(time.time())}}.png")', 'print("[OK] Ekran görüntüsü alındı")']

    # ── Pencere odakla ────────────────────────────────────────────────────────
    if re.search(r'(?:pencere|window|odakla|focus|aktif)', sl) and quoted:
        title = quoted[0]
        return [
            f'# {s}',
            f'if _PW_OK:',
            f'    try:',
            f'        _PwApp(backend="uia").connect(title_re=".*{title}.*", timeout=5).top_window().set_focus()',
            f'        print("[OK] Pencere odaklandı: {title}")',
            f'    except Exception as _fe: print(f"[WARN] Pencere bulunamadı: {{_fe}}")',
        ]

    # ── Tırnaksız doğrulama/durum (görülür, açıldı — quoted olmadan) ─────────
    if re.search(ASSERT_WORDS, sl) and not quoted:
        return [f'# {s}  # assertion: eleman metni belirtilmemiş, manuel doğrula']

    return None


def _llm_fallback_steps_prompt(unrecognized: list[str], platform: str) -> str:
    """Kural motorunun tanıyamadığı adımlar için LLM prompt'u."""
    if platform == "mobile":
        api = "AppiumBy.XPATH, '//*[@text=\"X\"]'"
    elif platform == "desktop":
        api = "_safe_click_desktop('X') | _safe_type_desktop('X') | _safe_assert_desktop('X') | _pag.hotkey('ctrl','x')"
    else:
        api = "By.XPATH, '//*[contains(text(),\"X\")]'"

    steps_text = "\n".join(f"- {s}" for s in unrecognized)
    return f"""Write Python Appium/Selenium lines for these test steps. No imports, no explanations.
Use only: driver.find_element({api}).click() | .send_keys("x") | .is_displayed() | time.sleep(n)
Output only code lines + # comments.

Steps:
{steps_text}

Code:"""


# ══════════════════════════════════════════════════════════════════════════════
# BÖLÜM 2: VISION FALLBACK FONKSİYONU (script içine gömülür)
# ══════════════════════════════════════════════════════════════════════════════

VISION_HELPER = '''
import base64 as _b64, requests as _req, re as _re

# ── Vision fallback: OCR (metni bulur) + Evren (nesne tespiti, sınıflı) ───────
# Koşum sırasında LLM (Gemini/Ollama) KULLANILMAZ — yavaştı (15-60 sn timeout).
# OCR: rapidocr_onnxruntime (lokal). Evren: ai_config.json'da key + version varsa
# paralel çağrılır. Sınıflar: Buton, Metin, İcon, Logo, Input, sol/sağ/aşağı/yukarı ok.
#  - Metin bir Buton/Input kutusundaysa kutunun merkezine dokunulur
#  - Metni olmayan oklar (geri, ileri, açılır liste) doğrudan Evren'den bulunur
#  - EditText bulunamayan input'larda etikete en yakın Input kutusuna dokunulur
_OCR_ENGINE = None
_TR_MAP = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")

import threading as _vis_threading
_OCR_LOCK = _vis_threading.Lock()


def _ocr_engine():
    global _OCR_ENGINE
    with _OCR_LOCK:
        if _OCR_ENGINE is None:
            from rapidocr_onnxruntime import RapidOCR
            _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _ocr_preload():
    """Her koşum yeni süreç → model ilk vision çağrısında yüklenince ~5 sn kaybediliyordu.
    Script açılırken arka planda yükle + ısıt; ilk adımlar koşarken hazır olur."""
    try:
        from PIL import Image
        import io
        _b = io.BytesIO()
        Image.new("RGB", (64, 32), "white").save(_b, "PNG")
        _ocr_engine()(_b.getvalue(), use_cls=False)
    except Exception:
        pass


_vis_threading.Thread(target=_ocr_preload, daemon=True).start()


def _vis_norm(s):
    """Türkçe harfleri sadeleştirip karşılaştırmaya hazırlar. OCR 'ç'yi sık
    sık 'q' okuduğundan (Seçimi → Seqimi) q da c'ye çevrilir; 'ı'yı da 'l'
    okuyabildiğinden (Takımları → Takimlarl) l → i. İki tarafa birden
    uygulandığı için gerçekten l içeren kelimeler (Ayarlar) etkilenmez."""
    s = str(s).translate(_TR_MAP).lower().replace("q", "c").replace("l", "i")
    return "".join(ch for ch in s if ch.isalnum())


# Evren sınıf adları da _vis_norm'dan geçer → "sol ok" / "Sol_Ok" / "solok" aynı
_CLS_BUTON = _vis_norm("Buton")
_CLS_INPUT = _vis_norm("Input")
_ARROW_PHRASES = {
    _vis_norm("sol ok"):    ["sol ok", "geri", "back"],
    _vis_norm("sağ ok"):    ["sağ ok", "ileri", "next"],
    _vis_norm("aşağı ok"):  ["aşağı ok", "açılır", "dropdown"],
    _vis_norm("yukarı ok"): ["yukarı ok"],
}


def _vis_cfg():
    try:
        import json, pathlib
        _cp = pathlib.Path(_AI_CONFIG_PATH) if "_AI_CONFIG_PATH" in globals() else None
        if _cp and _cp.exists():
            return json.loads(_cp.read_text(encoding="utf-8"))
    except Exception:
        pass
    return {}


def _evren_boxes(png_bytes, img_w, img_h, out):
    """Evren predict → [(x1, y1, x2, y2, sınıf, güven), ...] piksel cinsinden (out'a yazar)."""
    cfg = _vis_cfg()
    key = (cfg.get("evren_api_key") or "").strip()
    ver = (cfg.get("evren_model_version_id") or "").strip()
    if not (key and ver):
        return
    url = cfg.get("evren_api_url") or "https://api.ssyz.org.tr/api/v1/inference/predict"
    try:
        for _h in ({"Authorization": "Bearer " + key}, {"X-API-Key": key}):
            r = _req.post(url, headers=_h, files={"file": ("screen.png", png_bytes, "image/png")},
                          data={"model_version_id": ver}, timeout=10)
            if r.status_code != 401:
                break
        if not r.ok:
            print(f"[VISION] Evren HTTP {r.status_code}: {r.text[:120]}")
            return
        for p in r.json().get("predictions") or []:
            bb = p.get("bbox") or []
            conf = p.get("confidence") or 0
            if len(bb) == 4 and conf >= 0.3:
                out.append((bb[0] * img_w, bb[1] * img_h, bb[2] * img_w, bb[3] * img_h,
                            _vis_norm(p.get("class_name") or ""), float(conf)))
    except Exception as _e:
        print(f"[VISION] Evren hatası: {_e}")


def _vis_score(target, line):
    """Hedef metin ile OCR satırı arasındaki benzerlik (0-1)."""
    import difflib
    if not target or not line:
        return 0.0
    if target == line:
        return 1.0
    if len(target) == 1:
        return 0.0  # tek karakter (PIN rakamı) sadece birebir eşleşir
    if target in line:
        # Hedef satırın ne kadarını kaplıyor → "Takım", "Takım Yönetimi ..." grubundan
        # önce tek başına "Takım" yazan satırı seçsin
        return (0.85 + 0.14 * len(target) / len(line)) if len(target) >= 3 else 0.0
    r = difflib.SequenceMatcher(None, target, line).ratio()
    # Uzunluklar çok farklıysa ceza: "Takvim Ayarları Kaydet" aranırken ekrandaki
    # "Takvim Ayarları" 0.82 alıp YANLIŞ elemana dokunulmasın (iki satıra bölünen
    # yazılar zaten bitişik satır grupları olarak tam uzunlukta değerlendiriliyor)
    lr = min(len(target), len(line)) / max(len(target), len(line))
    return r * lr if lr < 0.8 else r


def _vis_capture(driver):
    import io
    from PIL import Image
    png = _b64.b64decode(driver.get_screenshot_as_base64())
    w, h = Image.open(io.BytesIO(png)).size
    return png, w, h


def _vis_analyze(png, img_w, img_h):
    """OCR (bu thread) + Evren (paralel thread). Döner: (items, boxes, evren_thread)
    items: [(metin, norm, x1, y1, x2, y2, güven)] — yukarıdan aşağı sıralı.
    boxes listesi Evren thread'i bitince dolar; gerekirse çağıran join eder."""
    import threading
    boxes = []
    th = threading.Thread(target=_evren_boxes, args=(png, img_w, img_h, boxes), daemon=True)
    th.start()
    # use_cls=False: ekran metni hep düz, açı sınıflandırma gereksiz (~%20 hızlı)
    ocr_res, _ = _ocr_engine()(png, use_cls=False)
    items = []
    for box, txt, conf in ocr_res or []:
        xs = [p[0] for p in box]; ys = [p[1] for p in box]
        items.append((txt, _vis_norm(txt), min(xs), min(ys), max(xs), max(ys), float(conf)))
    items.sort(key=lambda it: (it[3], it[2]))
    return items, boxes, th


def _vis_adjacent(a, b):
    """b, a'nın hemen altında (aynı sütun) ya da hemen sağında (aynı satır) mı?
    İki satıra bölünmüş buton yazıları ("Takımları" / "Gör") tek aday olsun diye."""
    h = max(a[5] - a[3], b[5] - b[3], 1)
    x_overlap = min(a[4], b[4]) - max(a[2], b[2])
    v_gap = b[3] - a[5]
    if x_overlap > 0 and -0.3 * h <= v_gap <= 0.8 * h:
        return True
    same_row = abs((a[3] + a[5]) / 2 - (b[3] + b[5]) / 2) < 0.5 * h
    h_gap = b[2] - a[4]
    return same_row and -0.3 * h <= h_gap <= 1.5 * h


def _vis_text_cands(items, search_text):
    """Hedef metne benzeyen OCR satırları + bitişik 2-3 satır grupları.
    Döner: (adaylar[(skor, x1, y1, x2, y2, güven, metin)], tüm_skorlar)"""
    target = _vis_norm(search_text)
    groups = [[it] for it in items]
    for i, a in enumerate(items):
        for j in range(i + 1, len(items)):
            b = items[j]
            if not _vis_adjacent(a, b):
                continue
            groups.append([a, b])
            for k in range(j + 1, len(items)):
                if _vis_adjacent(b, items[k]):
                    groups.append([a, b, items[k]])
    cands, scores = [], []
    for g in groups:
        txt = " ".join(it[0] for it in g)
        sc = _vis_score(target, "".join(it[1] for it in g))
        scores.append((sc, txt))
        if sc >= 0.8:
            cands.append((sc, min(it[2] for it in g), min(it[3] for it in g),
                          max(it[4] for it in g), max(it[5] for it in g),
                          min(it[6] for it in g), txt))
    return cands, scores


def _vis_container(boxes, cx, cy, classes):
    """(cx, cy) noktasını içeren, verilen sınıflardaki EN KÜÇÜK Evren kutusu."""
    inside = [b for b in boxes if b[4] in classes and b[0] <= cx <= b[2] and b[1] <= cy <= b[3]]
    return min(inside, key=lambda b: (b[2] - b[0]) * (b[3] - b[1])) if inside else None


def _vis_tap(driver, x, y, img_w, img_h):
    """Screenshot pikseli → dokunma koordinatı (iOS'ta point/piksel oranı farklı)."""
    try:
        ws = driver.get_window_size()
        sx, sy = ws["width"] / img_w, ws["height"] / img_h
    except Exception:
        sx = sy = 1.0
    x, y = int(x * sx), int(y * sy)
    driver.tap([(x, y)])
    return x, y


def _vis_save_fail(png):
    try:
        import tempfile, os, time as _t
        _shot = os.path.join(tempfile.gettempdir(), f"vision_fail_{int(_t.time())}.png")
        open(_shot, "wb").write(png)
        print(f"[VISION] Ekran görüntüsü: {_shot}")
    except Exception:
        pass


def _vis_arrow_intent(search_text):
    """'Geri', 'sol ok', 'aşağı oka' ... → (ok sınıfı, kalan yakın metin, açık_ok_mu)
    ya da (None, None, False). açık_ok: ifade doğrudan "... ok" (önce Evren oku aranır);
    geri/ileri/açılır gibi kelimelerde önce metin aranır (butonda "Geri" yazabilir)."""
    # "İ".lower() birleşik noktalı i üretir ("i̇leri") → önce düz i'ye çevir
    low = str(search_text).replace("İ", "i").replace("I", "ı").lower()
    toks = [_vis_norm(w) for w in low.split()]
    toks = [t for t in toks if t]

    def _has(ph):
        # Kelime bazlı: "Değerleri" içindeki "geri" eşleşmesin; son kelimede kısa
        # ek serbest (ok → oka/okuna, geri → geriye)
        pt = [_vis_norm(w) for w in ph.split()]
        n = len(pt)
        for i in range(len(toks) - n + 1):
            seg = toks[i:i + n]
            if seg[:-1] == pt[:-1] and seg[-1].startswith(pt[-1]) and len(seg[-1]) <= len(pt[-1]) + 3:
                return True
        return False

    for cls, phrases in _ARROW_PHRASES.items():
        for ph in phrases:
            if _has(ph):
                near = low.replace(ph, " ")
                for _w in ("butonuna", "butonu", "ikonuna", "ikonu", "oka", "simgesine",
                           "simgesi", "yanındaki", "yanında", "tıkla", "tıklanır"):
                    near = near.replace(_w, " ")
                return cls, " ".join(near.split()), ph.endswith(" ok")
    return None, None, False


def _vis_pick_arrow(boxes, cls, items, near_text):
    """Ok sınıfındaki kutulardan birini seç: yakın metin varsa ona en yakını,
    geri (sol ok) için ekranın en üst-solundaki, diğerlerinde en yüksek güvenli."""
    arrows = [b for b in boxes if b[4] == cls]
    if not arrows:
        return None, ""
    if near_text and len(_vis_norm(near_text)) >= 3:
        cands, _ = _vis_text_cands(items, near_text)
        if cands:
            c = max(cands, key=lambda c: (round(c[0], 2), c[5]))
            tx, ty = (c[1] + c[3]) / 2, (c[2] + c[4]) / 2
            b = min(arrows, key=lambda b: ((b[0] + b[2]) / 2 - tx) ** 2 + ((b[1] + b[3]) / 2 - ty) ** 2)
            return b, f"'{c[6]}' metnine en yakın"
    if cls == _vis_norm("sol ok"):
        return min(arrows, key=lambda b: (round(b[1] / 50), b[0])), "en üst-sol"
    return max(arrows, key=lambda b: b[5]), "en yüksek güven"


def _vision_find(driver, search_text):
    """Element bulunamazsa: screenshot → OCR ile metni bul (Buton kutusuna hizala)
    → yoksa ok ifadesiyse (geri/ileri/açılır) Evren'den ok bul → tap."""
    import time as _t
    try:
        print(f"[VISION] '{search_text}' aranıyor (OCR + Evren)...")
        _t0 = _t.time()
        png, img_w, img_h = _vis_capture(driver)
        try:
            items, boxes, th = _vis_analyze(png, img_w, img_h)
        except ImportError as _ie:
            import sys as _sys
            print(f"[VISION] OCR yüklenemedi ({_ie}) — Python: {_sys.executable} "
                  f"→ bu Python'a: pip install rapidocr_onnxruntime")
            return False
        _t_ocr = _t.time() - _t0
        arrow_cls, near_text, explicit_arrow = _vis_arrow_intent(search_text)
        # Ok aranıyorsa Evren şart → tam bekle; değilse Evren sadece hizalama/tercih
        # için, OCR bittikten sonra en fazla 1 sn (yavaş/502 dönen Evren yavaşlatmasın)
        th.join(timeout=10 if arrow_cls else 1.0)
        print(f"[VISION] OCR {len(items)} metin ({_t_ocr:.1f} sn), "
              f"Evren {len(boxes)} nesne ({_t.time() - _t0:.1f} sn)")

        def _tap_arrow():
            b, why = _vis_pick_arrow(boxes, arrow_cls, items, near_text)
            if not b:
                print(f"[VISION] Evren ekranda '{arrow_cls}' sınıfında ok bulamadı")
                return False
            x, y = _vis_tap(driver, (b[0] + b[2]) / 2, (b[1] + b[3]) / 2, img_w, img_h)
            print(f"[VISION] Evren ok bulundu ({why}, güven {b[5]:.2f}) → tap ({x}, {y}) "
                  f"[{_t.time() - _t0:.1f} sn]")
            return True

        # "X aşağı ok" gibi açık ok ifadesi: önce ok; yoksa X etiketine dokunulur
        # (dropdown'larda etikete dokunmak da listeyi açar)
        if explicit_arrow:
            if _tap_arrow():
                return True
            search_text = near_text or search_text

        cands, scores = _vis_text_cands(items, search_text)
        if cands:
            def _key(c):
                cx, cy = (c[1] + c[3]) / 2, (c[2] + c[4]) / 2
                in_btn = _vis_container(boxes, cx, cy, (_CLS_BUTON, _CLS_INPUT)) is not None
                return (round(c[0], 2), in_btn, c[5])
            if len(_vis_norm(search_text)) == 1:
                # PIN klavyesi: aynı rakam ekranın başka yerinde de olabilir → en alttaki tuş
                best = max(cands, key=lambda c: c[2])
            else:
                best = max(cands, key=_key)
            sc, x1, y1, x2, y2, conf, txt = best
            cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
            btn = _vis_container(boxes, cx, cy, (_CLS_BUTON, _CLS_INPUT))
            where = "metin merkezi"
            if btn:
                cx, cy = (btn[0] + btn[2]) / 2, (btn[1] + btn[3]) / 2
                where = f"Evren {('Buton' if btn[4] == _CLS_BUTON else 'Input')} merkezi"
            x, y = _vis_tap(driver, cx, cy, img_w, img_h)
            print(f"[VISION] '{txt}' bulundu (benzerlik {sc:.2f}) → {where} "
                  f"tap ({x}, {y}) [{_t.time() - _t0:.1f} sn]")
            return True

        if arrow_cls and not explicit_arrow and _tap_arrow():
            return True

        print(f"[VISION] '{search_text}' ekranda bulunamadı ({_t.time() - _t0:.1f} sn)")
        _near = sorted(scores, reverse=True)[:5]
        if _near:
            print("[VISION] En yakın metinler: " + ", ".join(f"'{t}' ({s:.2f})" for s, t in _near))
        _vis_save_fail(png)
        return False
    except Exception as e:
        print(f"[VISION] Hata: {e}")
        return False


def _vision_focus_input(driver, label):
    """EditText bulunamadığında: OCR ile alan etiketini bul → Evren'in etikete en
    yakın Input kutusuna (içinde ya da altında) dokun. Yazmayı çağıran yapar."""
    import time as _t
    if not label or len(_vis_norm(label)) < 2:
        return False
    try:
        print(f"[VISION] '{label}' alanı aranıyor (OCR + Evren Input)...")
        _t0 = _t.time()
        png, img_w, img_h = _vis_capture(driver)
        items, boxes, th = _vis_analyze(png, img_w, img_h)
        th.join(timeout=10)
        inputs = [b for b in boxes if b[4] == _CLS_INPUT]
        cands, _ = _vis_text_cands(items, label)
        if not cands:
            print(f"[VISION] '{label}' etiketi ekranda bulunamadı")
            _vis_save_fail(png)
            return False
        c = max(cands, key=lambda c: (round(c[0], 2), c[5]))
        tx, ty = (c[1] + c[3]) / 2, (c[2] + c[4]) / 2
        if inputs:
            def _dist(b):
                bx, by = (b[0] + b[2]) / 2, (b[1] + b[3]) / 2
                if b[0] <= tx <= b[2] and b[1] <= ty <= b[3]:
                    return 0  # etiket kutunun içinde (placeholder)
                below = by >= ty  # input genelde etiketin altında
                return ((bx - tx) ** 2 + (by - ty) ** 2) * (1 if below else 4)
            b = min(inputs, key=_dist)
            x, y = _vis_tap(driver, (b[0] + b[2]) / 2, (b[1] + b[3]) / 2, img_w, img_h)
            print(f"[VISION] '{c[6]}' etiketine en yakın Evren Input → tap ({x}, {y}) "
                  f"[{_t.time() - _t0:.1f} sn]")
        else:
            x, y = _vis_tap(driver, tx, ty, img_w, img_h)
            print(f"[VISION] Evren Input bulamadı → '{c[6]}' metnine tap ({x}, {y})")
        _t.sleep(0.4)
        return True
    except Exception as e:
        print(f"[VISION] Input arama hatası: {e}")
        return False


def _safe_click_arrow(driver, phrase):
    """Ok ikonları ('Takımları Gör aşağı ok', 'sol ok') — yazısı olmadığı için
    XPath/scroll denenmez, doğrudan Evren ok tespitine gidilir."""
    print(f"[OK-IKON] '{phrase}' aranıyor...")
    if not _vision_find(driver, phrase):
        raise Exception(f"Ok bulunamadı: {phrase}")


def _safe_click(driver, text, key_text=None):
    """
    Akıllı click stratejisi:
    1. Tam metin ile dene
    2. Kısa anahtar kelime ile dene (karmaşık açıklamalar için)
    3. Aşağı kaydır → tekrar dene
    4. Yukarı kaydır → tekrar dene
    5. Vision fallback (ekran görüntüsü → AI → koordinat)
    key_text: "Toplam Varlıklarım kartında yer alan favori fonlarım" → "favori fonlarım"
    """
    # Ham XPath/resource-id ise direkt kullan, scroll/vision'a gerek yok
    if text.startswith("//") or text.startswith("//*"):
        try:
            driver.find_element(AppiumBy.XPATH, text).click()
            print(f"[OK] XPath tıklandı: {text}")
            return
        except Exception as _xe:
            raise Exception(f"XPath bulunamadı: {text} → {_xe}")
    if text.startswith("id/") or text.startswith("@id/"):
        _rid = text.lstrip("@")
        try:
            driver.find_element(AppiumBy.ID, _rid).click()
            print(f"[OK] resource-id tıklandı: {_rid}")
            return
        except Exception as _xe:
            raise Exception(f"resource-id bulunamadı: {_rid} → {_xe}")

    _all_texts = [t for t in [text, key_text] if t]

    def _try(t):
        for _xp in [
            f\'//*[@text="{t}"]\',
            f\'//android.widget.Button[@text="{t}"]\',
            f\'//android.widget.TextView[@text="{t}"]\',
            f\'//*[contains(@text,"{t}")]\',
            f\'//*[@content-desc="{t}"]\',
            f\'//*[contains(@content-desc,"{t}")]\',
        ]:
            try:
                driver.find_element(AppiumBy.XPATH, _xp).click()
                print(f"[OK] Tıklandı: {t}")
                return True
            except Exception:
                continue
        return False

    # 1 & 2: direkt dene
    for _t in _all_texts:
        if _try(_t): return

    # 3: aşağı kaydır + tekrar
    print(f"[SCROLL↓] '{text}' bulunamadı, aşağı kaydırılıyor...")
    driver.swipe(540, 1400, 540, 500, 800); time.sleep(0.7)
    for _t in _all_texts:
        if _try(_t): return

    # 4: yukarı kaydır + tekrar (orijinal konuma + biraz üstüne)
    print(f"[SCROLL↑] Yukarı kaydırılıyor...")
    driver.swipe(540, 500, 540, 1400, 800); time.sleep(0.7)
    for _t in _all_texts:
        if _try(_t): return

    # 5: Vision fallback — kısa metin varsa onu gönder (AI için daha net)
    _vision_target = key_text if key_text else text
    print(f"[VISION] '{_vision_target}' için ekran görüntüsü analiz ediliyor...")
    if not _vision_find(driver, _vision_target):
        raise Exception(f"Element bulunamadı (tüm stratejiler denendi): {text}")


def _safe_input(driver, xpath, value):
    """Metin girişi — önce ekrana bak, PIN klavyesi varsa tıkla, yoksa send_keys/ADB"""
    import subprocess as _sp, shutil as _sh, os as _os, re as _re2
    str_value = str(value)
    print(f"[INPUT] '{str_value}' yazılıyor... (xpath={xpath})")

    # ── iOS: send_keys yeterli, ADB/mobile:type desteklenmiyor ───────────────
    try:
        _platform = driver.capabilities.get("platformName", "").lower()
    except Exception:
        _platform = ""
    if _platform == "ios":
        try:
            _ios_el = driver.find_element(AppiumBy.XPATH, xpath) if xpath else None
        except Exception:
            _ios_el = None
        try:
            if _ios_el:
                _ios_el.click()
                import time as _t; _t.sleep(0.3)
                _ios_el.send_keys(str_value)
            else:
                driver.switch_to.active_element.send_keys(str_value)
            print(f"[INPUT] ✅ iOS send_keys: {str_value}")
        except Exception as _e:
            print(f"[INPUT] iOS send_keys hatası: {_e}")
            try:
                driver.switch_to.active_element.send_keys(str_value)
                print(f"[INPUT] ✅ iOS active_element send_keys: {str_value}")
            except Exception as _e2:
                print(f"[INPUT] iOS yazma başarısız: {_e2}")
        return

    # ── ADB yolu ──────────────────────────────────────────────────────────────
    def _adb_exec():
        _p = _sh.which("adb")
        if _p: return _p
        for _base in [_os.environ.get("LOCALAPPDATA",""), _os.environ.get("USERPROFILE","")]:
            _c = _os.path.join(_base, "AppData","Local","Android","Sdk","platform-tools","adb.exe")
            if _os.path.exists(_c): return _c
        return "adb"
    _ADB = _adb_exec()
    _adb = [_ADB, "-s", _DEVICE_ID] if _DEVICE_ID else [_ADB]

    # ── Sayfa kaynağını bir kez al ────────────────────────────────────────────
    try:
        _src = driver.page_source
    except Exception:
        _src = ""

    # ── PIN klavyesi tespiti: 6+ rakam butonu görünüyorsa ─────────────────────
    def _detect_pin_keyboard():
        _digit_count = sum(
            1 for _d in "0123456789"
            if f\'text="{_d}"\' in _src or f\'content-desc="{_d}"\' in _src
        )
        print(f"[INPUT] Rakam butonu sayısı: {_digit_count}")
        return _digit_count >= 6

    # ── Tek rakam butonuna tıkla ───────────────────────────────────────────────
    def _click_digit(d):
        # Ekrandaki tüm eleman bilgilerini çekip koordinat bazlı sırala
        _patterns = [
            f\'//android.widget.Button[@text="{d}"]\',
            f\'//*[@text="{d}" and @clickable="true"]\',
            f\'//*[@content-desc="{d}" and @clickable="true"]\',
            f\'//*[contains(@resource-id,"key") and @text="{d}"]\',
            f\'//*[contains(@resource-id,"btn") and @text="{d}"]\',
            f\'//*[contains(@resource-id,"digit") and @text="{d}"]\',
            f\'//*[@text="{d}"]\',
            f\'//*[@content-desc="{d}"]\',
        ]
        for _xp in _patterns:
            try:
                _els = driver.find_elements(AppiumBy.XPATH, _xp)
                if not _els:
                    continue
                # En alt y-koordinatlı elemanı tercih et (klavye genelde altta)
                _best = max(_els, key=lambda e: e.location.get("y", 0))
                _best.click()
                print(f"[PIN] '{d}' tıklandı ({_xp[:50]})")
                time.sleep(0.25)
                return True
            except Exception:
                continue
        # Vision fallback
        print(f"[PIN] '{d}' butonla bulunamadı, vision deneniyor...")
        return _vision_find(driver, d)

    # ── PIN klavyesi açıksa → direkt butonlara tıkla ─────────────────────────
    if _detect_pin_keyboard():
        print(f"[INPUT] PIN klavyesi tespit edildi, buton tıklama modu")
        _ok = 0
        for _ch in str_value:
            if _click_digit(_ch):
                _ok += 1
            else:
                print(f"[PIN] '{_ch}' yazılamadı")
        print(f"[INPUT] PIN sonuç: {_ok}/{len(str_value)} karakter")
        if _ok >= len(str_value):
            return
        # Kısmen yazdıysa yine de devam et
        if _ok > 0:
            return

    # ── Standart EditText modu ────────────────────────────────────────────────
    def _find_el():
        for _xp in [xpath,
                    '//android.widget.EditText[@password="true"]',
                    '//android.widget.EditText[1]',
                    '//android.widget.EditText']:
            try:
                _e = driver.find_element(AppiumBy.XPATH, _xp)
                _e.click(); time.sleep(0.4)
                return _e
            except Exception:
                continue
        return None

    def _field_has_value():
        for _xp in [xpath, '//android.widget.EditText[@password="true"]',
                    '//android.widget.EditText[1]']:
            try:
                _e = driver.find_element(AppiumBy.XPATH, _xp)
                _v = _e.text or _e.get_attribute("text") or _e.get_attribute("value") or ""
                if _v.strip():
                    return True
            except Exception:
                continue
        return False

    _el = _find_el()
    print(f"[INPUT] EditText {'bulundu' if _el else 'bulunamadı'}")
    if not _el:
        # EditText yok (WebView/özel bileşen) → OCR ile etiketi bul, Evren Input'una dokun;
        # yazmayı aşağıdaki ADB/mobile:type stratejileri odaklanan alana yapar
        _lbl = None
        for _attr in ("@text", "@hint", "@content-desc"):
            _ix = (xpath or "").find(_attr)
            if _ix >= 0:
                _rest = xpath[_ix:]
                _qs = [k for k, ch in enumerate(_rest) if ch in (chr(39), chr(34))]
                if len(_qs) >= 2:
                    _lbl = _rest[_qs[0] + 1:_qs[1]]
                    break
        if _lbl and _vision_focus_input(driver, _lbl):
            _el = _find_el()

    # Alan tıklandıktan sonra PIN klavyesi açıldıysa buton moduna geç
    time.sleep(0.4)
    try:
        _src_r = driver.page_source
        _pin_r = sum(1 for _d in "0123456789"
                     if f\'text="{_d}"\' in _src_r or f\'content-desc="{_d}"\' in _src_r)
        if _pin_r >= 6:
            print(f"[INPUT] Alan tıklanınca PIN klavyesi açıldı ({_pin_r} rakam)")
            _ok = 0
            for _ch in str_value:
                if _click_digit(_ch):
                    _ok += 1
            if _ok > 0:
                print(f"[INPUT] PIN sonuç: {_ok}/{len(str_value)}")
                return
    except Exception as _e:
        print(f"[INPUT] PIN re-check hata: {_e}")

    # Placeholder metnini kaydet: typing başarısını doğrulamak için
    # Bazı uygulamalar placeholder\'ı actual text olarak set eder (ör: "Şifre Giriniz")
    # send_keys başarısızsa uygulama placeholder\'ı geri koyar, değiştiyse başarılı
    _placeholder = ""
    try:
        if _el:
            _placeholder = _el.text or _el.get_attribute("text") or ""
            print(f"[INPUT] Placeholder: \'{_placeholder[:40]}\'")
    except:
        pass

    def _typing_succeeded(el):
        # Placeholder kontrolü ÖNCE gelir — _field_has_value placeholder\'ı dolu sayar
        if _placeholder:
            try:
                _curr = el.text or el.get_attribute("text") or "" if el else ""
            except:
                _curr = ""
            # Placeholder hala görünüyorsa typing kesinlikle başarısız
            if _curr == _placeholder:
                return False
            # Placeholder değiştiyse başarılı
            return True
        # Placeholder yoksa field_has_value\'ya bak
        return _field_has_value()

    # Strateji A: send_keys
    if _el:
        try:
            _el.clear(); time.sleep(0.2)
            _el.send_keys(str_value); time.sleep(0.6)
            if _typing_succeeded(_el):
                print(f"[OK] send_keys başarılı: {value}"); return
            print("[INPUT] send_keys başarısız (placeholder değişmedi)")
        except Exception as _e:
            print(f"[INPUT] send_keys hata: {_e}")

    # Strateji B: ADB input text
    try:
        _sp.run(_adb + ["shell", "input", "text", str_value],
                capture_output=True, timeout=5)
        time.sleep(0.6)
        if _typing_succeeded(_el):
            print(f"[OK] ADB input text başarılı: {value}"); return
        print("[INPUT] ADB input text başarısız (placeholder değişmedi)")
    except Exception as _e:
        print(f"[INPUT] ADB hata: {_e}")

    # Strateji B2: ADB keyevent — rakam dizileri için (IME\'den bağımsız, doğrudan tuş eventi)
    if str_value.isdigit():
        try:
            _km = {\'0\':7,\'1\':8,\'2\':9,\'3\':10,\'4\':11,\'5\':12,\'6\':13,\'7\':14,\'8\':15,\'9\':16}
            for _ch in str_value:
                _sp.run(_adb + ["shell", "input", "keyevent", str(_km[_ch])],
                        capture_output=True, timeout=3)
                time.sleep(0.12)
            time.sleep(0.5)
            if _typing_succeeded(_el):
                print(f"[OK] ADB keyevent başarılı: {value}"); return
            print("[INPUT] ADB keyevent başarısız (placeholder değişmedi)")
        except Exception as _e:
            print(f"[INPUT] ADB keyevent hata: {_e}")

    # Strateji C: mobile:type
    try:
        if _el or _find_el():
            driver.execute_script("mobile: type", {"text": str_value})
            time.sleep(0.4)
            if _typing_succeeded(_el):
                print(f"[OK] mobile:type başarılı: {value}"); return
        print("[INPUT] mobile:type başarısız")
    except Exception as _e:
        print(f"[INPUT] mobile:type hata: {_e}")

    # Strateji D: Clipboard paste
    try:
        driver.set_clipboard_text(str_value); time.sleep(0.2)
        _sp.run(_adb + ["shell", "input", "keyevent", "279"], capture_output=True, timeout=5)
        time.sleep(0.4)
        if _typing_succeeded(_el):
            print(f"[OK] Clipboard başarılı: {value}"); return
        print("[INPUT] Clipboard başarısız")
    except Exception as _e:
        print(f"[INPUT] Clipboard hata: {_e}")

    # Strateji E: Alan tekrar bul → PIN klavyesi açıldıysa tıkla
    time.sleep(0.5)
    try:
        _src2 = driver.page_source
        _digit_count2 = sum(
            1 for _d in "0123456789"
            if f\'text="{_d}"\' in _src2 or f\'content-desc="{_d}"\' in _src2
        )
        if _digit_count2 >= 4:
            print(f"[INPUT] Klavye sonradan açıldı ({_digit_count2} rakam), buton moduna geçiliyor")
            _ok2 = 0
            for _ch in str_value:
                if _click_digit(_ch):
                    _ok2 += 1
            print(f"[INPUT] Geç PIN sonuç: {_ok2}/{len(str_value)}")
            if _ok2 > 0:
                return
    except Exception as _e:
        print(f"[INPUT] Geç klavye tespiti hata: {_e}")

    # Strateji F: Vision per-karakter (son çare)
    print(f"[INPUT] Vision per-karakter başlatılıyor: {value}")
    _ok_v = 0
    for _ch in str_value:
        if _vision_find(driver, _ch):
            _ok_v += 1; time.sleep(0.25)
        else:
            print(f"[INPUT] Vision '{_ch}' bulunamadı")
    print(f"[INPUT] Vision sonuç: {_ok_v}/{len(str_value)}")


def _safe_assert(driver, text, key_text=None):
    """
    Element doğrulama — scroll + vision destekli.
    1. Direkt bul
    2. Aşağı kaydır → tekrar bul
    3. Yukarı kaydır → tekrar bul
    4. Vision ile koordinat al → görsel doğrulama
    Bulunamazsa WARNING verir ama testi durdurmaz (beklenen ekran gecikmeli gelebilir).
    """
    _all_texts = [t for t in [text, key_text] if t]
    _is_xp = text.startswith("//") or text.startswith("(//")

    def _try_find():
        for _t in _all_texts:
            for _xp in ([_t] if _is_xp else [
                f\'//*[@text="{_t}"]\',
                f\'//*[contains(@text,"{_t}")]\',
                f\'//*[@content-desc="{_t}"]\',
                f\'//*[contains(@content-desc,"{_t}")]\',
            ]):
                try:
                    el = driver.find_element(AppiumBy.XPATH, _xp)
                    if el.is_displayed():
                        print(f"[OK] Doğrulandı: {_t}")
                        return True
                except Exception:
                    continue
        return False

    if _try_find(): return True

    # Aşağı kaydır + tekrar
    print(f"[SCROLL↓] '{text}' aranıyor...")
    driver.swipe(540, 1400, 540, 500, 800); time.sleep(0.7)
    if _try_find(): return True

    # Yukarı kaydır + tekrar
    print(f"[SCROLL↑] '{text}' aranıyor...")
    driver.swipe(540, 500, 540, 1400, 800); time.sleep(0.7)
    if _try_find(): return True

    # Vision fallback (XPath hedefte görüntüde aranacak yazı yok)
    if _is_xp:
        print(f"[WARN] XPath ile eleman bulunamadı: {text} — adım geçiliyor")
        return False
    print(f"[VISION] '{text}' için ekran görüntüsü analiz ediliyor...")
    if _vision_find(driver, key_text or text):
        print(f"[OK] Vision ile doğrulandı: {text}")
        return True

    print(f"[WARN] '{text}' ekranda bulunamadı — adım geçiliyor")
    return False


def _safe_read(driver, text, key_text=None):
    """Hedef metni/içeriği içeren elementin text/content-desc değerini okur (scroll ile arar)."""
    _all_texts = [t for t in [text, key_text] if t]

    def _try_find():
        for _t in _all_texts:
            for _xp in [
                f'//*[contains(@text,"{_t}")]',
                f'//*[contains(@content-desc,"{_t}")]',
            ]:
                try:
                    el = driver.find_element(AppiumBy.XPATH, _xp)
                    val = (el.get_attribute("text") or el.get_attribute("content-desc") or "").strip()
                    if val:
                        return val
                except Exception:
                    continue
        return None

    val = _try_find()
    if val is not None:
        print(f"[OK] Okundu: {text}")
        return val

    print(f"[SCROLL↓] '{text}' aranıyor...")
    driver.swipe(540, 1400, 540, 500, 800); time.sleep(0.7)
    val = _try_find()
    if val is not None:
        return val

    print(f"[SCROLL↑] '{text}' aranıyor...")
    driver.swipe(540, 500, 540, 1400, 800); time.sleep(0.7)
    val = _try_find()
    if val is not None:
        return val

    print(f"[WARN] '{text}' ekranda bulunamadı — değer okunamadı")
    return ""
'''


VISION_HELPER_WEB = '''
from selenium.webdriver.support.ui import WebDriverWait as _WDWEB
from selenium.webdriver.support import expected_conditions as _EC_WEB

def _wdwait(driver, timeout=10):
    return _WDWEB(driver, timeout)


def _safe_click_web(driver, text):
    """Web elementi tıkla — WebDriverWait + scroll + çoklu strateji"""
    from selenium.webdriver.common.by import By
    import time
    strategies = [
        (By.XPATH,         f'//*[normalize-space(text())="{text}"]'),
        (By.XPATH,         f'//button[contains(normalize-space(.),"{text}")]'),
        (By.XPATH,         f'//a[contains(normalize-space(.),"{text}")]'),
        (By.XPATH,         f'//input[@value="{text}"]'),
        (By.XPATH,         f'//*[@aria-label="{text}"]'),
        (By.XPATH,         f'//*[contains(text(),"{text}")]'),
        (By.LINK_TEXT,     text),
        (By.PARTIAL_LINK_TEXT, text),
    ]
    for scroll_pass in range(3):
        for by, val in strategies:
            try:
                el = _wdwait(driver, 5).until(_EC_WEB.element_to_be_clickable((by, val)))
                driver.execute_script("arguments[0].scrollIntoView({block:'center',inline:'nearest'});", el)
                time.sleep(0.15)
                el.click()
                print(f"[OK] Tıklandı: {text}")
                return
            except Exception:
                continue
        driver.execute_script(f"window.scrollBy(0, {350 * (scroll_pass + 1)});")
        time.sleep(0.4)
    raise Exception(f"Web elementi bulunamadı: {text}")


def _safe_input_web(driver, field_hint, value, is_password=False):
    """Input alanini placeholder/label/name/id ile bul ve deger gir"""
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.action_chains import ActionChains
    import time
    hint = (field_hint or "").strip().lower()
    val  = str(value)

    def _type_into(el):
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
        time.sleep(0.15)
        el.clear()
        el.send_keys(val)

    # 1) Her durumda once aktif (odaktaki) elemana yaz
    #    Gorsel tiklamadan sonra oraya gelir, hint olmasa da calisir
    try:
        _ae = driver.switch_to.active_element
        if _ae and _ae.tag_name.lower() in ("input", "textarea"):
            _type_into(_ae)
            print(f"[OK] Aktif elemana girildi: '{val}'")
            return
    except Exception:
        pass

    # 2) Password
    if is_password:
        for _by, _sel in [
            (By.XPATH,        '//input[@type="password"]'),
            (By.CSS_SELECTOR, 'input[type="password"]'),
        ]:
            try:
                _e = _wdwait(driver, 4).until(_EC_WEB.visibility_of_element_located((_by, _sel)))
                _type_into(_e)
                print(f"[OK] Sifre girildi")
                return
            except Exception:
                continue

    # 3) Hint ile ara
    if hint:
        _UC = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
        _LC = "abcdefghijklmnopqrstuvwxyz"
        for _by, _sel in [
            (By.XPATH, f'//input[contains(translate(@placeholder,"{_UC}","{_LC}"),"{hint}")]'),
            (By.XPATH, f'//textarea[contains(translate(@placeholder,"{_UC}","{_LC}"),"{hint}")]'),
            (By.XPATH, f'//input[contains(translate(@name,"{_UC}","{_LC}"),"{hint}")]'),
            (By.XPATH, f'//input[contains(translate(@id,"{_UC}","{_LC}"),"{hint}")]'),
            (By.XPATH, f'//input[contains(translate(@aria-label,"{_UC}","{_LC}"),"{hint}")]'),
            (By.XPATH, f'//*[contains(translate(.,"{_UC}","{_LC}"),"{hint}")]//following-sibling::input[1]'),
            (By.XPATH, f'//*[contains(translate(.,"{_UC}","{_LC}"),"{hint}")]/..//input[1]'),
        ]:
            try:
                _e = _wdwait(driver, 3).until(_EC_WEB.visibility_of_element_located((_by, _sel)))
                _type_into(_e)
                print(f"[OK] Girildi: '{val}' (hint: {hint})")
                return
            except Exception:
                continue

    # 4) Sayfadaki ilk gorunur input/textarea
    for _by, _sel in [
        (By.CSS_SELECTOR, 'input:not([type=hidden]):not([type=submit]):not([type=button]):not([type=checkbox]):not([type=radio]):not([type=file])'),
        (By.TAG_NAME,     'textarea'),
    ]:
        try:
            _e = _wdwait(driver, 3).until(_EC_WEB.visibility_of_element_located((_by, _sel)))
            _type_into(_e)
            print(f"[OK] Girildi: '{val}' (ilk input)")
            return
        except Exception:
            continue

    # 5) Son care: ActionChains ile mevcut odaga gonder
    try:
        ActionChains(driver).send_keys(val).perform()
        print(f"[OK] ActionChains ile girildi: '{val}'")
    except Exception as _ex:
        raise Exception(f"Input girilemiyor: {_ex}")


def _safe_assert_web(driver, text, raise_on_fail=True):
    """Metni sayfada doğrula — kaynak + element + scroll"""
    from selenium.webdriver.common.by import By
    import time
    if text in driver.page_source:
        print(f"[OK] Doğrulandı: {text}"); return True
    for xp in [
        f'//*[contains(normalize-space(.),"{text}")]',
        f'//*[contains(@value,"{text}")]',
        f'//*[contains(@placeholder,"{text}")]',
    ]:
        try:
            el = driver.find_element(By.XPATH, xp)
            if el.is_displayed():
                print(f"[OK] Doğrulandı: {text}"); return True
        except Exception:
            continue
    driver.execute_script("window.scrollTo(0,0);"); time.sleep(0.5)
    _found = text in driver.page_source
    if _found:
        return True
    if raise_on_fail:
        assert _found, f"Beklenen metin yok: {text}"
    return False


def _safe_read_web(driver, hint):
    """Hedef metni içeren elementin görünen değerini okur (contains eşleşmesi)."""
    from selenium.webdriver.common.by import By
    for xp in [
        f'//*[contains(normalize-space(.),"{hint}")][not(.//*[contains(normalize-space(.),"{hint}")])]',
        f'//*[contains(@value,"{hint}")]',
        f'//*[contains(@placeholder,"{hint}")]',
    ]:
        try:
            el = driver.find_element(By.XPATH, xp)
            val = (el.text or el.get_attribute("value") or "").strip()
            if val:
                return val
        except Exception:
            continue
    print(f"[WARN] Okunamadı: {hint}")
    return ""
'''


IMAGE_HELPER = '''
import base64 as _ib64, tempfile as _itmp, os as _ios, time as _itime

def _get_window_region():
    """_WINDOW_TITLE penceresinin ekran bölgesini (x,y,w,h) döndürür. Bulamazsa None."""
    try:
        _wt = _WINDOW_TITLE  # generated script globals'ta tanımlı
    except NameError:
        return None
    if not _wt:
        return None
    # pywinauto ile pencere dikdörtgeni al
    try:
        import pywinauto as _pwr
        _app = _pwr.Application(backend="uia").connect(title_re=f".*{_wt}.*", timeout=2)
        _rect = _app.top_window().rectangle()
        _x, _y = max(0, _rect.left), max(0, _rect.top)
        _w, _h = _rect.width(), _rect.height()
        if _w > 10 and _h > 10:
            return (_x, _y, _w, _h)
    except Exception:
        pass
    # win32gui fallback
    try:
        import win32gui as _wg
        def _cb(h, acc):
            if _wt.lower() in _wg.GetWindowText(h).lower():
                acc.append(h)
        _hs = []; _wg.EnumWindows(_cb, _hs)
        if _hs:
            _x1,_y1,_x2,_y2 = _wg.GetWindowRect(_hs[0])
            _w, _h = _x2-_x1, _y2-_y1
            if _w > 10 and _h > 10:
                return (max(0,_x1), max(0,_y1), _w, _h)
    except Exception:
        pass
    return None

def _find_b64_on_screen(b64_data, threshold=0.80, timeout=10):
    """Base64 gorsel arar — _WINDOW_TITLE varsa yalnizca o pencere icinde."""
    _tmp_path = _ios.path.join(_itmp.gettempdir(), f"_img_find_{id(b64_data)}.png")
    try:
        with open(_tmp_path, "wb") as _f:
            _f.write(_ib64.b64decode(b64_data))
        # Pencere bolgesi — _WINDOW_TITLE varsa (Desktop) tam ekran aramaz. Web script'lerinde
        # _WINDOW_TITLE hic tanimli olmadigi icin globals() ile guvenli kontrol edilir.
        _wt_set = '_WINDOW_TITLE' in globals() and bool(_WINDOW_TITLE)
        _region = _get_window_region()
        if _wt_set and _region is None:
            _wt_end = _itime.time() + 3
            while _itime.time() < _wt_end:
                _itime.sleep(0.3)
                _region = _get_window_region()
                if _region: break
            if _region is None:
                print(f"[HATA] '{_WINDOW_TITLE}' penceresi bulunamadi — gorsel aranamaz.")
                return None
        _start = _itime.time()
        while _itime.time() - _start < timeout:
            if _wt_set:
                _upd = _get_window_region()
                if _upd: _region = _upd
            _elapsed = _itime.time() - _start
            _conf = max(0.50, threshold - (_elapsed / timeout) * 0.30)
            _coords = _find_image_on_screen(_tmp_path, threshold=_conf, region=_region)
            if _coords:
                return _coords
            _itime.sleep(0.3)
        print(f"[DEBUG] Gorsel {timeout}s icinde bulunamadi. Farkli bir bolgeyi kroplayip tekrar deneyin.")
        return None
    finally:
        try: _ios.unlink(_tmp_path)
        except Exception: pass

def _find_image_on_screen(target_path, threshold=0.80, region=None, return_bbox=False):
    """OpenCV multi-scale template matching. region=(x,y,w,h) ile kısıtlı arama yapar.
    return_bbox=True ise merkez nokta yerine eslesen alanin (x,y,w,h) dikdortgenini dondurur."""
    try:
        import cv2 as _cv2, numpy as _np
        import pyautogui as _pag_sc
        # Windows DPI ölçeğini tespit et — template farklı DPI'da alınmış olabilir
        _dpi_scale = 1.0
        try:
            import ctypes as _ct
            _dpi_scale = _ct.windll.user32.GetDpiForSystem() / 96.0
        except Exception:
            pass
        _screenshot = _pag_sc.screenshot(region=region) if region else _pag_sc.screenshot()
        _screen      = _np.array(_screenshot)
        _ox, _oy     = (region[0], region[1]) if region else (0, 0)
        _screen_bgr  = _cv2.cvtColor(_screen, _cv2.COLOR_RGB2BGR)
        _screen_gray = _cv2.cvtColor(_screen_bgr, _cv2.COLOR_BGR2GRAY)
        _screen_gray = _cv2.GaussianBlur(_screen_gray, (3,3), 0)
        _tmpl_bgr  = _cv2.imread(target_path)
        if _tmpl_bgr is None: return None
        _tmpl_gray = _cv2.cvtColor(_tmpl_bgr, _cv2.COLOR_BGR2GRAY)
        _tmpl_gray = _cv2.GaussianBlur(_tmpl_gray, (3,3), 0)
        _h, _w = _tmpl_gray.shape[:2]
        _best = None; _best_val = -1
        # DPI farklılıklarını (125%/150%/200%) karşılamak için geniş scale aralığı
        # Merkez: 1/dpi_scale (fiziksel px template → mantıksal px ekran)
        _center = 1.0 / _dpi_scale
        _lo = max(0.3, _center * 0.6)
        _hi = min(2.5, _center * 1.7)
        for _sc in _np.linspace(_lo, _hi, 25):
            _rw, _rh = int(_w*_sc), int(_h*_sc)
            if _rw < 8 or _rh < 8: continue
            if _rw > _screen_gray.shape[1] or _rh > _screen_gray.shape[0]: continue
            _res = _cv2.matchTemplate(_screen_gray, _cv2.resize(_tmpl_gray,(_rw,_rh)), _cv2.TM_CCOEFF_NORMED)
            _, _val, _, _loc = _cv2.minMaxLoc(_res)
            if _val > _best_val: _best_val = _val; _best = (_loc, _rw, _rh)
        if _best_val >= threshold and _best:
            _loc, _rw, _rh = _best
            if return_bbox:
                return (_loc[0] + _ox, _loc[1] + _oy, _rw, _rh)
            return (_loc[0] + _rw//2 + _ox, _loc[1] + _rh//2 + _oy)
    except ImportError:
        try:
            import pyautogui as _pag2
            _loc2 = _pag2.locateOnScreen(target_path, grayscale=True,
                                          confidence=min(0.8, threshold), region=region)
            if _loc2:
                if return_bbox:
                    return (_loc2.left, _loc2.top, _loc2.width, _loc2.height)
                return (_pag2.center(_loc2).x, _pag2.center(_loc2).y)
        except Exception: pass
    except Exception: pass
    return None

def _do_robust_click(x, y, button="left"):
    import pyautogui as _pag3
    _pag3.moveTo(x, y, duration=0.2)
    _itime.sleep(0.3)
    _pag3.mouseDown(button=button)
    _itime.sleep(0.05)
    _pag3.mouseUp(button=button)

def _find_b64_on_screen_web(b64_data, threshold=0.80, timeout=10):
    """Web için: görseli ekranda ara (pencere kısıtlaması olmadan)."""
    return _find_b64_on_screen(b64_data, threshold=threshold, timeout=timeout)

def _bbox_click_point(bbox, click_x, click_y, orig_w, orig_h):
    """bbox=(x,y,w,h) eslesen ekran alani; click_x/click_y kirpilan/tam gorseldeki orijinal
    piksel nokta; orig_w/orig_h o gorselin boyutu. Mutlak ekran (x,y) dondurur."""
    _norm_x = click_x / orig_w
    _norm_y = click_y / orig_h
    _bx, _by, _bw, _bh = bbox
    return (int(_bx + _norm_x * _bw), int(_by + _norm_y * _bh))

def _find_b64_and_click_point(b64_data, click_x, click_y, action="click", threshold=0.80, timeout=10):
    """Base64 gorseli (kirpilmis alan veya eski tam sayfa) ekranda coklu-olcek eslestirmeyle
    bulur, gorsel icindeki (click_x, click_y) orijinal piksel noktasini eslesen alana gore
    mutlak ekran koordinatina cevirip tiklar."""
    import io as _iio3
    try:
        from PIL import Image as _PILY
        _pil = _PILY.open(_iio3.BytesIO(_ib64.b64decode(b64_data)))
        _orig_w, _orig_h = _pil.size
    except Exception as _pe:
        print(f"[WARN] Gorsel boyutu alinamiyor: {_pe}")
        return False

    _tmp_path = _ios.path.join(_itmp.gettempdir(), f"_img_click_{id(b64_data)}.png")
    try:
        with open(_tmp_path, "wb") as _f:
            _f.write(_ib64.b64decode(b64_data))

        _wt_set = '_WINDOW_TITLE' in globals() and bool(_WINDOW_TITLE)
        _region = _get_window_region() if _wt_set else None

        _start = _itime.time()
        while _itime.time() - _start < timeout:
            _elapsed = _itime.time() - _start
            _conf = max(0.50, threshold - (_elapsed / timeout) * 0.30)
            _bbox = _find_image_on_screen(_tmp_path, threshold=_conf, region=_region, return_bbox=True)
            if _bbox:
                _abs_x, _abs_y = _bbox_click_point(_bbox, click_x, click_y, _orig_w, _orig_h)
                import pyautogui as _pag4
                if action == "double_click":
                    _pag4.doubleClick(_abs_x, _abs_y)
                elif action == "right_click":
                    _pag4.rightClick(_abs_x, _abs_y)
                else:
                    _pag4.click(_abs_x, _abs_y)
                print(f"[OK] Nokta tiklandi (bolge esleme): ({_abs_x}, {_abs_y})")
                return True
            _itime.sleep(0.3)
        print(f"[DEBUG] Gorsel {timeout}s icinde bulunamadi — nokta tiklanamadi.")
        return False
    finally:
        try: _ios.unlink(_tmp_path)
        except Exception: pass

def _find_b64_and_drag(b64_data, x1, y1, x2, y2, drag_duration=0.4, threshold=0.80, timeout=10):
    """Kirpilmis referans gorseli ekranda coklu-olcek eslestirmeyle bulur, gorsel icindeki
    iki noktayi ((x1,y1) -> (x2,y2)) eslesen alana gore mutlak ekran koordinatina cevirip
    aralarinda fare ile surukler (mouseDown -> moveTo -> mouseUp)."""
    import io as _iio4
    try:
        from PIL import Image as _PILY2
        _pil2 = _PILY2.open(_iio4.BytesIO(_ib64.b64decode(b64_data)))
        _orig_w, _orig_h = _pil2.size
    except Exception as _pe2:
        print(f"[WARN] Gorsel boyutu alinamiyor: {_pe2}")
        return False

    _tmp_path = _ios.path.join(_itmp.gettempdir(), f"_img_drag_{id(b64_data)}.png")
    try:
        with open(_tmp_path, "wb") as _f:
            _f.write(_ib64.b64decode(b64_data))

        _wt_set = '_WINDOW_TITLE' in globals() and bool(_WINDOW_TITLE)
        _region = _get_window_region() if _wt_set else None

        _start = _itime.time()
        while _itime.time() - _start < timeout:
            _elapsed = _itime.time() - _start
            _conf = max(0.50, threshold - (_elapsed / timeout) * 0.30)
            _bbox = _find_image_on_screen(_tmp_path, threshold=_conf, region=_region, return_bbox=True)
            if _bbox:
                _ax1, _ay1 = _bbox_click_point(_bbox, x1, y1, _orig_w, _orig_h)
                _ax2, _ay2 = _bbox_click_point(_bbox, x2, y2, _orig_w, _orig_h)
                import pyautogui as _pag5
                _pag5.moveTo(_ax1, _ay1, duration=0.2)
                _pag5.mouseDown()
                _itime.sleep(0.1)
                _pag5.moveTo(_ax2, _ay2, duration=drag_duration)
                _itime.sleep(0.05)
                _pag5.mouseUp()
                print(f"[OK] Suruklendi (bolge esleme): ({_ax1},{_ay1}) -> ({_ax2},{_ay2})")
                return True
            _itime.sleep(0.3)
        print(f"[DEBUG] Gorsel {timeout}s icinde bulunamadi — surukleme yapilamadi.")
        return False
    finally:
        try: _ios.unlink(_tmp_path)
        except Exception: pass
'''

IMAGE_HELPER_WEB_CDP = '''
def _find_image_in_page(driver, target_path, threshold=0.80, return_bbox=False):
    """Chrome/Edge (CDP) icin: OS ekrani degil, sayfanin KENDI screenshot'ini kullanir.
    Headless'ta da calisir, paralel tarayicilar birbirini etkilemez (her biri sadece
    kendi render buffer'ini gorur)."""
    try:
        import cv2 as _cv2w, numpy as _npw, io as _iiow
        from PIL import Image as _PILW
        _png = driver.get_screenshot_as_png()
        _img = _PILW.open(_iiow.BytesIO(_png)).convert("RGB")
        _screen = _npw.array(_img)
        _screen_bgr = _cv2w.cvtColor(_screen, _cv2w.COLOR_RGB2BGR)
        _screen_gray = _cv2w.cvtColor(_screen_bgr, _cv2w.COLOR_BGR2GRAY)
        _screen_gray = _cv2w.GaussianBlur(_screen_gray, (3, 3), 0)
        _tmpl_bgr = _cv2w.imread(target_path)
        if _tmpl_bgr is None:
            return None
        _tmpl_gray = _cv2w.cvtColor(_tmpl_bgr, _cv2w.COLOR_BGR2GRAY)
        _tmpl_gray = _cv2w.GaussianBlur(_tmpl_gray, (3, 3), 0)
        _h, _w = _tmpl_gray.shape[:2]
        _best = None
        _best_val = -1
        for _sc in _npw.linspace(0.6, 1.7, 25):
            _rw, _rh = int(_w * _sc), int(_h * _sc)
            if _rw < 8 or _rh < 8:
                continue
            if _rw > _screen_gray.shape[1] or _rh > _screen_gray.shape[0]:
                continue
            _res = _cv2w.matchTemplate(_screen_gray, _cv2w.resize(_tmpl_gray, (_rw, _rh)), _cv2w.TM_CCOEFF_NORMED)
            _, _val, _, _loc = _cv2w.minMaxLoc(_res)
            if _val > _best_val:
                _best_val = _val
                _best = (_loc, _rw, _rh)
        if _best_val >= threshold and _best:
            _loc, _rw, _rh = _best
            if return_bbox:
                return (_loc[0], _loc[1], _rw, _rh)
            return (_loc[0] + _rw // 2, _loc[1] + _rh // 2)
    except Exception as _e:
        print(f"[WARN] Sayfa goruntusunde gorsel arama hatasi: {_e}")
    return None


def _cdp_click(driver, x, y, button="left", click_count=1):
    """Gercek OS faresine dokunmadan, tarayicinin icine dogrudan mouse event basar."""
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {
        "type": "mouseMoved", "x": x, "y": y,
    })
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {
        "type": "mousePressed", "x": x, "y": y, "button": button, "clickCount": click_count,
    })
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {
        "type": "mouseReleased", "x": x, "y": y, "button": button, "clickCount": click_count,
    })


def _cdp_drag(driver, x1, y1, x2, y2, steps=10):
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {"type": "mouseMoved", "x": x1, "y": y1})
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {
        "type": "mousePressed", "x": x1, "y": y1, "button": "left", "clickCount": 1,
    })
    for _i in range(1, steps + 1):
        _ix = x1 + (x2 - x1) * _i / steps
        _iy = y1 + (y2 - y1) * _i / steps
        driver.execute_cdp_cmd("Input.dispatchMouseEvent", {
            "type": "mouseMoved", "x": _ix, "y": _iy, "buttons": 1,
        })
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {
        "type": "mouseReleased", "x": x2, "y": y2, "button": "left", "clickCount": 1,
    })


def _find_b64_in_page_web(driver, b64_data, threshold=0.80, timeout=10):
    """Base64 gorseli sayfa screenshot'unda arar (Chrome/Edge, CDP tabanli)."""
    _tmp_path = _ios.path.join(_itmp.gettempdir(), f"_img_find_web_{id(b64_data)}.png")
    try:
        with open(_tmp_path, "wb") as _f:
            _f.write(_ib64.b64decode(b64_data))
        _start = _itime.time()
        while _itime.time() - _start < timeout:
            _elapsed = _itime.time() - _start
            _conf = max(0.50, threshold - (_elapsed / timeout) * 0.30)
            _coords = _find_image_in_page(driver, _tmp_path, threshold=_conf)
            if _coords:
                return _coords
            _itime.sleep(0.3)
        print(f"[DEBUG] Gorsel {timeout}s icinde bulunamadi (sayfa goruntusu).")
        return None
    finally:
        try:
            _ios.unlink(_tmp_path)
        except Exception:
            pass


def _find_b64_and_click_point_web(driver, b64_data, click_x, click_y, action="click", threshold=0.80, timeout=10):
    """Kirpilmis referans gorseli sayfa screenshot'unda bulur, gorsel icindeki (click_x, click_y)
    noktasini eslesen alana gore CDP ile tiklar."""
    import io as _iiowc
    try:
        from PIL import Image as _PILWC
        _pil = _PILWC.open(_iiowc.BytesIO(_ib64.b64decode(b64_data)))
        _orig_w, _orig_h = _pil.size
    except Exception as _pe:
        print(f"[WARN] Gorsel boyutu alinamiyor: {_pe}")
        return False

    _tmp_path = _ios.path.join(_itmp.gettempdir(), f"_img_click_web_{id(b64_data)}.png")
    try:
        with open(_tmp_path, "wb") as _f:
            _f.write(_ib64.b64decode(b64_data))
        _start = _itime.time()
        while _itime.time() - _start < timeout:
            _elapsed = _itime.time() - _start
            _conf = max(0.50, threshold - (_elapsed / timeout) * 0.30)
            _bbox = _find_image_in_page(driver, _tmp_path, threshold=_conf, return_bbox=True)
            if _bbox:
                _abs_x, _abs_y = _bbox_click_point(_bbox, click_x, click_y, _orig_w, _orig_h)
                _btn = "right" if action == "right_click" else "left"
                _cnt = 2 if action == "double_click" else 1
                _cdp_click(driver, _abs_x, _abs_y, button=_btn, click_count=_cnt)
                print(f"[OK] Nokta tiklandi (Selenium CDP): ({_abs_x}, {_abs_y})")
                return True
            _itime.sleep(0.3)
        print(f"[DEBUG] Gorsel {timeout}s icinde bulunamadi (sayfa goruntusu) — nokta tiklanamadi.")
        return False
    finally:
        try:
            _ios.unlink(_tmp_path)
        except Exception:
            pass


def _find_b64_and_drag_web(driver, b64_data, x1, y1, x2, y2, drag_duration=0.4, threshold=0.80, timeout=10):
    import io as _iiowd
    try:
        from PIL import Image as _PILWD
        _pil2 = _PILWD.open(_iiowd.BytesIO(_ib64.b64decode(b64_data)))
        _orig_w, _orig_h = _pil2.size
    except Exception as _pe2:
        print(f"[WARN] Gorsel boyutu alinamiyor: {_pe2}")
        return False

    _tmp_path = _ios.path.join(_itmp.gettempdir(), f"_img_drag_web_{id(b64_data)}.png")
    try:
        with open(_tmp_path, "wb") as _f:
            _f.write(_ib64.b64decode(b64_data))
        _start = _itime.time()
        while _itime.time() - _start < timeout:
            _elapsed = _itime.time() - _start
            _conf = max(0.50, threshold - (_elapsed / timeout) * 0.30)
            _bbox = _find_image_in_page(driver, _tmp_path, threshold=_conf, return_bbox=True)
            if _bbox:
                _ax1, _ay1 = _bbox_click_point(_bbox, x1, y1, _orig_w, _orig_h)
                _ax2, _ay2 = _bbox_click_point(_bbox, x2, y2, _orig_w, _orig_h)
                _cdp_drag(driver, _ax1, _ay1, _ax2, _ay2)
                print(f"[OK] Suruklendi (sayfa goruntusu): ({_ax1},{_ay1}) -> ({_ax2},{_ay2})")
                return True
            _itime.sleep(0.3)
        print(f"[DEBUG] Gorsel {timeout}s icinde bulunamadi (sayfa goruntusu) — surukleme yapilamadi.")
        return False
    finally:
        try:
            _ios.unlink(_tmp_path)
        except Exception:
            pass
'''


DESKTOP_HELPER = '''
import pyautogui as _pag
import subprocess as _dsproc
import time, base64 as _b64, requests as _req, re as _re3, io as _io

_pag.FAILSAFE = True
_pag.PAUSE    = 0.25

# ── Opsiyonel bağımlılıklar ───────────────────────────────────────────────────
try:
    import pytesseract as _tess
    from PIL import Image as _PILImg
    _TESS_OK = True
except Exception:
    _TESS_OK = False
    print("[DESKTOP] pytesseract kullanılamıyor — OCR devre dışı")

try:
    from pywinauto import Application as _PwApp
    from pywinauto.findwindows import ElementNotFoundError as _PwNotFound
    _PW_OK = True
except Exception:
    _PW_OK = False
    print("[DESKTOP] pywinauto kullanılamıyor — pencere kontrolü devre dışı")


def _screenshot_b64(region=None):
    """Ekran (veya belirtilen bölge) fotoğrafı → base64 PNG."""
    shot = _pag.screenshot(region=region) if region else _pag.screenshot()
    buf  = _io.BytesIO()
    shot.save(buf, format="PNG")
    return _b64.b64encode(buf.getvalue()).decode()


def _vision_find_desktop(search_text):
    """Gemini veya Ollama ile pencere icinde koordinat bul. (x, y) veya (None, None) doner.
    _WINDOW_TITLE varsa yalnizca o pencerenin kroplamasini AI'ye gonderir; donus koordinatlari
    pencere orijinine gore offsetlenir — pencere disina tiklamaz."""
    try:
        _region = _get_window_region()  # (wx, wy, ww, wh) veya None
        _wx, _wy = (_region[0], _region[1]) if _region else (0, 0)
        shot_b64 = _screenshot_b64(region=_region)  # sadece pencere alanini cek

        try:
            import json as _js, pathlib as _pl
            cfg = {}
            _cp = _pl.Path(_AI_CONFIG_PATH) if "_AI_CONFIG_PATH" in dir() else None
            if _cp and _cp.exists():
                cfg = _js.loads(_cp.read_text(encoding="utf-8"))
            gemini_key = cfg.get("gemini_api_key", "")
        except Exception:
            gemini_key = ""

        _prompt = (f"Bu ekran goruntusunde '{search_text}' yazan veya gosteren "
                   f"tiklanabilir elementi bul. Yaniti sadece: x=NNN y=NNN "
                   f"(koordinatlar bu goruntunun sol ust kosesindendir)")

        if gemini_key:
            resp = _req.post(
                f"https://generativelanguage.googleapis.com/v1beta/models/gemini-1.5-flash:generateContent?key={gemini_key}",
                json={"contents": [{"parts": [
                    {"text": _prompt},
                    {"inline_data": {"mime_type": "image/png", "data": shot_b64}}
                ]}]},
                timeout=15
            )
            answer = resp.json()["candidates"][0]["content"]["parts"][0]["text"]
        else:
            resp = _req.post(_ollama_url(), json={
                "model": _ollama_model(), "prompt": _prompt,
                "images": [shot_b64], "stream": False
            }, timeout=60)
            answer = resp.json().get("response", "")

        mx = _re3.search(r"x[=:\s]*(\d+)", answer, _re3.IGNORECASE)
        my = _re3.search(r"y[=:\s]*(\d+)", answer, _re3.IGNORECASE)
        if mx and my:
            # Krop koordinatlarindan ekran koordinatina donustur
            x, y = int(mx.group(1)) + _wx, int(my.group(1)) + _wy
            print(f"[VISION] Koordinat (pencere ofseti {_wx},{_wy}): ({x}, {y})")
            return x, y
        print(f"[VISION] Koordinat bulunamadi — yanit: {answer[:80]}")
    except Exception as _ve:
        print(f"[VISION] Hata: {_ve}")
    return None, None


def _in_window(x, y):
    """Koordinatlar pencere siniri icinde mi? _WINDOW_TITLE yoksa her zaman True."""
    _r = _get_window_region()
    if not _r:
        return True
    _wx, _wy, _ww, _wh = _r
    return _wx <= x <= _wx + _ww and _wy <= y <= _wy + _wh


def _get_active_window():
    """Aktif pencereyi pywinauto ile döndür."""
    if not _PW_OK:
        return None
    try:
        import ctypes
        hwnd = ctypes.windll.user32.GetForegroundWindow()
        return _PwApp(backend="uia").connect(handle=hwnd).top_window()
    except Exception:
        return None


def _safe_click_desktop(target, window_title=""):
    """
    Desktop tıklama — 4 strateji:
    1. pywinauto uia backend
    2. pywinauto win32 backend (Java/Delphi/özel UI için)
    3. pytesseract OCR (kuruluysa)
    4. Vision AI / locateOnScreen
    """
    print(f"[CLICK] '{target}' aranıyor...")

    def _pw_try(backend):
        try:
            if window_title:
                _win = _PwApp(backend=backend).connect(title_re=f".*{window_title}.*", timeout=3).top_window()
            else:
                _win = _get_active_window()
            if not _win:
                return False
            for _ctrl in [
                lambda: _win.child_window(title=target, found_index=0),
                lambda: _win.child_window(title_re=f".*{target}.*", found_index=0),
            ]:
                try:
                    _el = _ctrl()
                    if _el.exists(timeout=1):
                        _el.set_focus()
                        _el.click_input()
                        print(f"[OK] pywinauto({backend}): '{target}'")
                        return True
                except Exception:
                    continue
        except Exception as _e:
            print(f"[CLICK] pywinauto({backend}) başarısız: {_e}")
        return False

    # 1. pywinauto uia ─────────────────────────────────────────────────────────
    if _PW_OK and _pw_try("uia"):
        return True

    # 2. pywinauto win32 ───────────────────────────────────────────────────────
    if _PW_OK and _pw_try("win32"):
        return True

    # 3. pytesseract OCR — yalnizca pencere bolgesi ───────────────────────────
    if _TESS_OK:
        try:
            _ocr_region = _get_window_region()
            _ocr_ox, _ocr_oy = (_ocr_region[0], _ocr_region[1]) if _ocr_region else (0, 0)
            shot = _pag.screenshot(region=_ocr_region) if _ocr_region else _pag.screenshot()
            data = _tess.image_to_data(shot, lang="tur+eng", output_type=_tess.Output.DICT)
            _tl  = target.lower()
            for _i, _w in enumerate(data["text"]):
                if _w and _tl in _w.lower() and data["conf"][_i] > 30:
                    _cx = data["left"][_i] + data["width"][_i] // 2 + _ocr_ox
                    _cy = data["top"][_i]  + data["height"][_i] // 2 + _ocr_oy
                    if _in_window(_cx, _cy):
                        _pag.click(_cx, _cy)
                        print(f"[OK] OCR: '{_w}' @ ({_cx},{_cy})")
                        return True
        except Exception as _e:
            print(f"[CLICK] OCR basarisiz: {_e}")

    # 4. locateOnScreen / Vision AI — pencere ile sinirli ────────────────────
    import os as _os2
    if _os2.path.isfile(target):
        try:
            _ls_region = _get_window_region()
            loc = _pag.locateCenterOnScreen(target, confidence=0.8, region=_ls_region)
            if loc and _in_window(loc.x, loc.y):
                _pag.click(loc)
                print(f"[OK] locateOnScreen: {target}")
                return True
        except Exception: pass

    _vx, _vy = _vision_find_desktop(target)
    if _vx and _vy and _in_window(_vx, _vy):
        _pag.click(_vx, _vy)
        print(f"[OK] Vision AI: '{target}' @ ({_vx},{_vy})")
        return True
    elif _vx and _vy:
        print(f"[HATA] Vision AI koordinati ({_vx},{_vy}) pencere disinda — tiklanmadi.")

    raise Exception(f"Element bulunamadi: '{target}'")


def _safe_type_desktop(text):
    """Aktif alana metin yaz — clipboard ile (Türkçe karakter desteği), fallback pyautogui."""
    print(f"[TYPE] \'{text}\'")
    try:
        import pyperclip as _clip
        _clip.copy(str(text))
        _pag.hotkey("ctrl", "v")
        print(f"[OK] Clipboard paste: \'{text}\'")
    except Exception:
        _pag.write(str(text), interval=0.04)
    time.sleep(0.2)


def _safe_double_click_desktop(target, window_title=""):
    """Desktop çift tıklama — uia → win32 → OCR → vision."""
    print(f"[DBLCLICK] '{target}' aranıyor...")

    def _pw_dbl(backend):
        try:
            _win = _PwApp(backend=backend).connect(title_re=f".*{window_title}.*", timeout=3).top_window() if window_title else _get_active_window()
            if not _win: return False
            for _ctrl in [lambda: _win.child_window(title=target, found_index=0), lambda: _win.child_window(title_re=f".*{target}.*", found_index=0)]:
                try:
                    _el = _ctrl()
                    if _el.exists(timeout=1):
                        _el.set_focus(); _el.double_click_input()
                        print(f"[OK] pywinauto({backend}) çift tıkla: '{target}'"); return True
                except Exception: continue
        except Exception as _e:
            print(f"[DBLCLICK] pywinauto({backend}) başarısız: {_e}")
        return False

    if _PW_OK and _pw_dbl("uia"): return True
    if _PW_OK and _pw_dbl("win32"): return True

    if _TESS_OK:
        try:
            _ocr_region = _get_window_region()
            _ocr_ox, _ocr_oy = (_ocr_region[0], _ocr_region[1]) if _ocr_region else (0, 0)
            shot = _pag.screenshot(region=_ocr_region) if _ocr_region else _pag.screenshot()
            data = _tess.image_to_data(shot, lang="tur+eng", output_type=_tess.Output.DICT)
            for _i, _w in enumerate(data["text"]):
                if _w and target.lower() in _w.lower() and data["conf"][_i] > 30:
                    _cx = data["left"][_i] + data["width"][_i] // 2 + _ocr_ox
                    _cy = data["top"][_i]  + data["height"][_i] // 2 + _ocr_oy
                    if _in_window(_cx, _cy):
                        _pag.doubleClick(_cx, _cy)
                        print(f"[OK] OCR cift tikla: '{_w}' @ ({_cx},{_cy})"); return True
        except Exception as _e:
            print(f"[DBLCLICK] OCR basarisiz: {_e}")

    _vx, _vy = _vision_find_desktop(target)
    if _vx and _vy and _in_window(_vx, _vy):
        _pag.doubleClick(_vx, _vy); print(f"[OK] Vision cift tikla: '{target}' @ ({_vx},{_vy})"); return True
    raise Exception(f"Element bulunamadi (cift tikla): '{target}'")


def _safe_right_click_desktop(target, window_title=""):
    """Desktop sağ tıklama — uia → win32 → OCR → vision."""
    print(f"[RCLICK] '{target}' aranıyor...")

    def _pw_right(backend):
        try:
            _win = _PwApp(backend=backend).connect(title_re=f".*{window_title}.*", timeout=3).top_window() if window_title else _get_active_window()
            if not _win: return False
            for _ctrl in [lambda: _win.child_window(title=target, found_index=0), lambda: _win.child_window(title_re=f".*{target}.*", found_index=0)]:
                try:
                    _el = _ctrl()
                    if _el.exists(timeout=1):
                        _el.set_focus(); _el.right_click_input()
                        print(f"[OK] pywinauto({backend}) sağ tıkla: '{target}'"); return True
                except Exception: continue
        except Exception as _e:
            print(f"[RCLICK] pywinauto({backend}) başarısız: {_e}")
        return False

    if _PW_OK and _pw_right("uia"): return True
    if _PW_OK and _pw_right("win32"): return True

    if _TESS_OK:
        try:
            _ocr_region = _get_window_region()
            _ocr_ox, _ocr_oy = (_ocr_region[0], _ocr_region[1]) if _ocr_region else (0, 0)
            shot = _pag.screenshot(region=_ocr_region) if _ocr_region else _pag.screenshot()
            data = _tess.image_to_data(shot, lang="tur+eng", output_type=_tess.Output.DICT)
            for _i, _w in enumerate(data["text"]):
                if _w and target.lower() in _w.lower() and data["conf"][_i] > 30:
                    _cx = data["left"][_i] + data["width"][_i] // 2 + _ocr_ox
                    _cy = data["top"][_i]  + data["height"][_i] // 2 + _ocr_oy
                    if _in_window(_cx, _cy):
                        _pag.rightClick(_cx, _cy)
                        print(f"[OK] OCR sag tikla: '{_w}' @ ({_cx},{_cy})"); return True
        except Exception as _e:
            print(f"[RCLICK] OCR basarisiz: {_e}")

    _vx, _vy = _vision_find_desktop(target)
    if _vx and _vy and _in_window(_vx, _vy):
        _pag.rightClick(_vx, _vy); print(f"[OK] Vision sag tikla: '{target}' @ ({_vx},{_vy})"); return True
    raise Exception(f"Element bulunamadi (sag tikla): '{target}'")


def _safe_menu_click_desktop(items, window_title=""):
    """Menü zinciri tıkla: [\'Dosya\', \'Yeni\'] sırayla tıklar."""
    print(f"[MENU] {' > '.join(items)}")
    for _item in items:
        _safe_click_desktop(_item, window_title)
        time.sleep(0.4)
    print(f"[OK] Menü tamamlandı: {' > '.join(items)}")


def _safe_wait_for_desktop(target, timeout=10):
    """Element ekranda görünene kadar bekle (polling)."""
    import time as _tw
    print(f"[WAIT] \'{target}\' bekleniyor (max {timeout}s)...")
    _start = _tw.time()
    while _tw.time() - _start < timeout:
        if _PW_OK:
            try:
                _win = _get_active_window()
                if _win:
                    _el = _win.child_window(title_re=f".*{target}.*", found_index=0)
                    if _el.exists(timeout=1):
                        print(f"[OK] pywinauto görüldü: \'{target}\'")
                        return True
            except Exception:
                pass
        if _TESS_OK:
            try:
                _wait_region = _get_window_region()
                _wtxt = _tess.image_to_string(
                    _pag.screenshot(region=_wait_region) if _wait_region else _pag.screenshot(),
                    lang="tur+eng"
                )
                if target.lower() in _wtxt.lower():
                    print(f"[OK] OCR goruldu: \'{target}\'")
                    return True
            except Exception:
                pass
        _tw.sleep(1)
    print(f"[WARN] \'{target}\' {timeout}s icinde gorunmedi")
    return False


def _safe_assert_desktop(target, raise_on_fail=True):
    """Ekranda element / metin doğrula. raise_on_fail=True (varsayılan) ise bulamazsa
    AssertionError fırlatır, False ise sessizce bool döner."""
    print(f"[ASSERT] \'{target}\' kontrol ediliyor...")

    # pywinauto
    if _PW_OK:
        try:
            _win = _get_active_window()
            if _win:
                _el = _win.child_window(title_re=f".*{target}.*", found_index=0)
                if _el.exists(timeout=1):
                    print(f"[OK] pywinauto: \'{target}\'"); return True
        except Exception:
            pass

    # OCR — sadece uygulama penceresiyle sınırlı
    if _TESS_OK:
        try:
            _region = _get_window_region()
            shot = _pag.screenshot(region=_region) if _region else _pag.screenshot()
            _txt = _tess.image_to_string(shot, lang="tur+eng")
            if target.lower() in _txt.lower():
                print(f"[OK] OCR: \'{target}\'"); return True
        except Exception:
            pass

    # Vision
    _vx, _vy = _vision_find_desktop(target)
    if _vx and _vy:
        print(f"[OK] Vision: \'{target}\'"); return True

    if raise_on_fail:
        raise AssertionError(f"[FAIL] \'{target}\' ekranda bulunamadi")
    return False

def _refocus_window():
    """Her adım öncesi pencereyi öne getirir — başlık veya süreç adıyla bulur."""
    import os as _os_rf
    _exe_name = _os_rf.path.basename(_APP_PATH).lower() if "_APP_PATH" in globals() and _APP_PATH else ""

    def _bring_hwnd(hwnd):
        try:
            import win32gui as _wg_r, win32con as _wc_r
            _wg_r.ShowWindow(hwnd, _wc_r.SW_RESTORE)
            _wg_r.SetForegroundWindow(hwnd)
            return True
        except Exception:
            return False

    # 1. Pencere başlığıyla bul ──────────────────────────────────────────────
    if _WINDOW_TITLE:
        try:
            import win32gui as _wg, win32con as _wc
            _found_hwnd = None
            def _cb(hwnd, _):
                nonlocal _found_hwnd
                t = _wg.GetWindowText(hwnd)
                if _WINDOW_TITLE.lower() in t.lower() and _wg.IsWindowVisible(hwnd):
                    _found_hwnd = hwnd
                    return False
                return True
            _wg.EnumWindows(_cb, None)
            if _found_hwnd:
                _wg.ShowWindow(_found_hwnd, _wc.SW_RESTORE)
                _wg.SetForegroundWindow(_found_hwnd)
                return
        except Exception:
            pass
        if _PW_OK:
            try:
                _PwApp(backend="uia").connect(title_re=".*" + _WINDOW_TITLE + ".*", timeout=3).top_window().set_focus()
                return
            except Exception:
                pass

    # 2. Süreç adıyla pencere bul (window_title boşsa) ──────────────────────
    if _exe_name:
        try:
            import win32gui as _wg2, win32process as _wp2, win32con as _wc2
            import psutil as _psu_rf
            _pids = {p.pid for p in _psu_rf.process_iter(["name"]) if p.info["name"].lower() == _exe_name}
            _wins = []
            def _cb2(hwnd, _):
                if _wg2.IsWindowVisible(hwnd) and _wg2.GetWindowText(hwnd):
                    _, _pid = _wp2.GetWindowThreadProcessId(hwnd)
                    if _pid in _pids:
                        _wins.append(hwnd)
                return True
            _wg2.EnumWindows(_cb2, None)
            if _wins:
                _wg2.ShowWindow(_wins[0], _wc2.SW_RESTORE)
                _wg2.SetForegroundWindow(_wins[0])
                return
        except Exception:
            pass
'''


# ══════════════════════════════════════════════════════════════════════════════
# BÖLÜM 3: ŞABLONLAR
# ══════════════════════════════════════════════════════════════════════════════

MAP_TRACE_HELPER = '''
# ── Uygulama haritası izi (backend _AI_MAP_TRACE verirse) ─────────────────
# İlk adımdan önce ve her başarılı adımdan sonra ekran görüntüsü alınır;
# koşum bitince backend bunları OCR'dan geçirip haritaya işler.
import os as _mt_os, json as _mt_json, builtins as _mt_bi
_MAP_TRACE_DIR = _mt_os.environ.get("_AI_MAP_TRACE", "")
_mt_state = {"k": 0, "started": False}


def _map_snap(after_step):
    _drv = globals().get("driver")
    if not _MAP_TRACE_DIR or _drv is None:
        return
    try:
        if _mt_state["k"] == 0:
            _app = globals().get("_TARGET_PKG") or getattr(_drv, "current_package", "") or ""
            _ver = ""
            try:
                import subprocess as _mt_sp
                _out = _mt_sp.run(["adb", "-s", globals().get("_DEVICE_ID", ""), "shell", "dumpsys", "package", _app],
                                  capture_output=True, text=True, timeout=8).stdout
                for _ln in _out.splitlines():
                    if "versionName=" in _ln:
                        _ver = _ln.split("versionName=", 1)[1].strip(); break
            except Exception:
                pass
            with open(_mt_os.path.join(_MAP_TRACE_DIR, "meta.json"), "w", encoding="utf-8") as _mf:
                _mt_json.dump({"app": _app, "version": _ver}, _mf)
        _f = "snap_" + str(_mt_state["k"]) + ".png"
        _drv.get_screenshot_as_file(_mt_os.path.join(_MAP_TRACE_DIR, _f))
        # Sayfa yapısı (resource-id'ler) → ekranı veriden bağımsız tanımak için
        try:
            with open(_mt_os.path.join(_MAP_TRACE_DIR, _f.replace(".png", ".xml")), "w", encoding="utf-8") as _xf:
                _xf.write(_drv.page_source)
        except Exception:
            pass
        with open(_mt_os.path.join(_MAP_TRACE_DIR, "events.jsonl"), "a", encoding="utf-8") as _ef:
            _ef.write(_mt_json.dumps({"file": _f, "after_step": after_step}) + chr(10))
        _mt_state["k"] += 1
    except Exception:
        pass


def print(*_a, **_k):
    _mt_bi.print(*_a, **_k)
    if _MAP_TRACE_DIR and _a and isinstance(_a[0], str):
        if _a[0].startswith("STEP_START:") and not _mt_state["started"]:
            _mt_state["started"] = True
            _map_snap(-1)
        elif _a[0].startswith("STEP_DONE:"):
            try:
                _map_snap(int(_a[0].split(":", 1)[1]))
            except ValueError:
                pass
'''


MAP_RECOVERY_HELPER = '''
# ── Kendini toparlama (uygulama haritası) ──────────────────────────────────
# Bir adım başarısız olunca: ekranı haritada tanı → bağlayıcı olmayan popup'ı
# kapat → yanlış ekrandaysa haritadaki yolla beklenen ekrana dön → adım 1 kez
# daha denenir. Backend _AI_MAP_RECOVER=0 verirse kapalıdır.
_MR_ENABLED = _mt_os.environ.get("_AI_MAP_RECOVER", "1") == "1"
# Sadece işlemi iptal/onay ETMEYEN butonlar — "Tamam/Onayla/Vazgeç/İptal" bilerek yok
_MR_DISMISS = ["kapat", "anladım", "daha sonra", "şimdi değil", "hayır teşekkürler", "atla", "geç"]
_MR_MAX_PATH = 6


def _mr_load():
    import sys as _mr_sys
    _bdir = _mt_os.path.dirname(globals().get("_AI_CONFIG_PATH", ""))
    if _bdir and _bdir not in _mr_sys.path:
        _mr_sys.path.insert(0, _bdir)
    from services import app_map_service as _ams
    return _ams


def _mr_identify(_ams, data):
    _drv = globals()["driver"]
    png = _drv.get_screenshot_as_png()
    try:
        src = _drv.page_source
    except Exception:
        src = None
    lines, w, h = _ams.ocr_lines(png)
    cur, score = _ams.match_screen(data, _ams.observation_tokens(lines, src))
    return cur, score, lines, w, h


def _mr_dismiss_popup(_ams, lines, w, h, step_target):
    """Ekranda bağlayıcı olmayan bir kapatma butonu varsa ona dokun."""
    _tn = _ams._norm(step_target or "")
    for word in _MR_DISMISS:
        _wn = _ams._norm(word)
        for l in lines:
            if l["norm"] == _wn and l["norm"] != _tn:
                _cx = (l["box"][0] + l["box"][2]) / 2 * w
                _cy = (l["box"][1] + l["box"][3]) / 2 * h
                _vis_tap(globals()["driver"], _cx, _cy, w, h)
                print(f"[TOPARLAMA] Popup kapatıldı: '{l['text']}'")
                time.sleep(1.0)
                return True
    return False


def _map_recover(idx, label):
    """True → adım tekrar denenebilir (ekran düzeltildi), False → toparlanamadı."""
    if not _MR_ENABLED or "driver" not in globals():
        return False
    try:
        _ams = _mr_load()
        app = globals().get("_TARGET_PKG") or ""
        data = _ams.load_map(app)
    except Exception as _e:
        print(f"[TOPARLAMA] Harita yüklenemedi: {_e}")
        return False
    if not data.get("screens"):
        print("[TOPARLAMA] Bu uygulama için harita henüz yok — toparlama atlandı")
        return False
    target = _ams.action_from_step(label).get("target") or ""
    expected = _ams.expected_screens(data, target)
    _last = str(globals().get("_LAST_STEP_ERROR") or "")
    if "instrumentation process is not running" in _last:
        print("[TOPARLAMA] Telefondaki UiAutomator2 çöktü — toparlama yapılamaz, oturumu yeniden başlatın.")
        return False
    print(f"[TOPARLAMA] Adım {idx} başarısız ('{label}') — ekran haritada aranıyor...")
    fixed = False
    for attempt in range(3):
        try:
            cur, score, lines, w, h = _mr_identify(_ams, data)
        except Exception as _e:
            print(f"[TOPARLAMA] Ekran tanınamadı: {_e}")
            return False
        cur_name = data["screens"].get(cur, {}).get("name") if cur else None
        print(f"[TOPARLAMA] Şu anki ekran: {cur_name or 'bilinmiyor'} (skor {score:.2f})")
        if _mr_dismiss_popup(_ams, lines, w, h, target):
            fixed = True
            continue
        if cur and (not expected or cur in expected):
            # Doğru ekrandayız; popup kapatıldıysa tekrar denemeye değer
            if not fixed:
                if expected:
                    print("[TOPARLAMA] Zaten beklenen ekrandayız — ekran sorunu değil")
                else:
                    print(f"[TOPARLAMA] '{target or label}' haritada hiçbir ekranda bilinmiyor — toparlama atlandı")
            return fixed
        if not cur:
            if attempt == 0:
                print("[TOPARLAMA] Bilinmeyen ekran → bir kez geri")
                globals()["driver"].back()
                time.sleep(1.0)
                fixed = True
                continue
            return fixed
        best = None
        for e in expected:
            p = _ams.find_path_in(data, cur, e)
            if p is not None and (best is None or len(p) < len(best)):
                best = p
        if best is None or len(best) > _MR_MAX_PATH:
            print(f"[TOPARLAMA] Haritada '{cur_name}' → beklenen ekran yolu yok")
            return fixed
        names = [data["screens"].get(t["to"], {}).get("name", "?") for t in best]
        print(f"[TOPARLAMA] Yol bulundu ({len(best)} adım): {' → '.join(names)}")
        try:
            for t in best:
                a = t.get("action", {})
                if a.get("type") == "back":
                    globals()["driver"].back()
                elif a.get("target"):
                    _safe_click(globals()["driver"], a["target"])
                else:
                    raise Exception("yoldaki geçişin hedefi yok")
                time.sleep(0.8)
        except Exception as _e:
            print(f"[TOPARLAMA] Yol izlenemedi: {_e}")
            return False
        if _MAP_TRACE_DIR:
            _map_snap(-2)   # harita izi: toparlama dokunuşları adım geçişi sayılmasın
        print("[TOPARLAMA] Beklenen ekrana dönüldü → adım tekrar deneniyor")
        return True
    return fixed
'''


# ── iOS uyum katmanı (iPhone, Mac düğümündeki Appium/XCUITest) ────────────────
# Kural motoru ve yardımcılar Android XPath'i üretir (@text, @content-desc,
# android.widget.EditText ...). iOS'ta driver.find_element(s) çağrıları burada iOS
# karşılıklarına çevrilir (@label/@value/@name, XCUIElementType...); kaydırma
# koordinatları (1080x2400 referanslı) iPhone ekranına ölçeklenir. Android'de no-op.
IOS_SHIM_HELPER = r'''
import re as _re_ios
_IOS_CLASSES = {
    "android.widget.EditText": "XCUIElementTypeTextField",
    "android.widget.Button": "XCUIElementTypeButton",
    "android.widget.ImageButton": "XCUIElementTypeButton",
    "android.widget.TextView": "XCUIElementTypeStaticText",
    "android.widget.ImageView": "XCUIElementTypeImage",
    "android.widget.CheckBox": "XCUIElementTypeSwitch",
    "android.widget.Switch": "XCUIElementTypeSwitch",
    "android.widget.ListView": "XCUIElementTypeTable",
    "android.widget.ScrollView": "XCUIElementTypeScrollView",
    "androidx.recyclerview.widget.RecyclerView": "XCUIElementTypeCollectionView",
}
_IOS_LIT = r'("[^"]*"|\'[^\']*\')'


def _ios_xpath(xp):
    for _a, _i in _IOS_CLASSES.items():
        xp = xp.replace(_a, _i)
    xp = xp.replace('XCUIElementTypeTextField[@password="true"]', "XCUIElementTypeSecureTextField")
    xp = xp.replace("XCUIElementTypeTextField[@password='true']", "XCUIElementTypeSecureTextField")
    xp = _re_ios.sub(r"contains\(\s*@text\s*,\s*" + _IOS_LIT + r"\s*\)",
                     r"(contains(@label,\1) or contains(@value,\1) or contains(@name,\1))", xp)
    xp = _re_ios.sub(r"@text\s*=\s*" + _IOS_LIT, r"(@label=\1 or @value=\1 or @name=\1)", xp)
    xp = xp.replace("@text", "@label").replace("@content-desc", "@name")
    xp = xp.replace("@resource-id", "@name").replace("@hint", "@placeholderValue")
    xp = xp.replace("@clickable", "@enabled").replace("@password", "@enabled")
    return xp


def _ios_install(drv):
    """iOS'ta find_element(s)/swipe'ı çevirir; Android'de dokunmaz."""
    if not _IS_IOS:
        return
    try:   # bağlanılan (attach) oturumda yetenekler boş gelir → yardımcılar iOS olduğunu bilsin
        drv.caps = {**(getattr(drv, "caps", None) or {}), "platformName": "iOS"}
    except Exception:
        pass
    _fe, _fes, _sw = drv.find_element, drv.find_elements, drv.swipe

    def _find_element(by="id", value=None):
        if by == AppiumBy.XPATH and value:
            value = _ios_xpath(value)
        return _fe(by, value)

    def _find_elements(by="id", value=None):
        if by == AppiumBy.XPATH and value:
            value = _ios_xpath(value)
        return _fes(by, value)

    _ws = {}

    def _swipe(x1, y1, x2, y2, duration=0):
        if not _ws:
            try:
                _ws.update(drv.get_window_size())
            except Exception:
                _ws.update({"width": 390, "height": 844})
        fx, fy = _ws["width"] / 1080.0, _ws["height"] / 2400.0
        return _sw(int(x1 * fx), int(y1 * fy), int(x2 * fx), int(y2 * fy), duration)

    drv.find_element, drv.find_elements, drv.swipe = _find_element, _find_elements, _swipe
    print("[iOS] XPath/kaydırma uyum katmanı etkin")
'''


def _build_mobile_template(
    req: GenerateRequest,
    steps_code: str,
    session_mode: str = "fresh",
    session_id: str = "",
    command_executor_url: str = "",
    vars_seed: dict | None = None,
    mjpeg_port: int = 0,
) -> str:
    device  = req.device_id    or "emulator-5554"
    pkg     = req.app_package  or ""
    act     = req.app_activity or ""
    appium  = req.appium_url   or "http://localhost:4723"
    # Telefon başka bir makinede (cihaz sunucusu): oturum o makinenin Appium'unda açılır,
    # script'in kendi adb çağrıları (monkey, dumpsys, ADB ile yazma) o makinenin adb sunucusuna gider
    adb_env: dict = {}
    is_ios = False
    if req.device_id:
        appium = dnodes.appium_url_for(req.device_id, appium)
        adb_env = dnodes.adb_env_for(req.device_id)
        is_ios = dnodes.platform_for(req.device_id) == "ios"   # pkg = Bundle ID
    env_lines = ([f"import os as _os_env; _os_env.environ.update({adb_env!r})  # telefon uzak cihaz sunucusunda"]
                 if adb_env else [])

    if session_mode in ("attach", "end"):
        lines = [
            "# -*- coding: utf-8 -*-",
            *env_lines,
            "import sys",
            f'sys.path.insert(0, r"{_ROUTER_BASE_DIR}")',
            "from services.session_reconnect import attach_appium_session",
            "from appium.webdriver.common.appiumby import AppiumBy",
            "from selenium.webdriver.support.ui import WebDriverWait",
            "from selenium.webdriver.support import expected_conditions as EC",
            "import time",
            f"_VARS = {(vars_seed or {})!r}  # adımlar arası değişken havuzu",
            "def _sub_vars(_s):",
            "    for _k, _v in _VARS.items():",
            "        _s = _s.replace('{{' + _k + '}}', str(_v))",
            "    return _s",
            "",
            f"driver = attach_appium_session({command_executor_url!r}, {session_id!r})",
            f"_IS_IOS = {is_ios}",
            IOS_SHIM_HELPER,
            "_ios_install(driver)",
            "wait = WebDriverWait(driver, 15)",
            f'_AI_CONFIG_PATH = r"{os.path.join(_ROUTER_BASE_DIR, "ai_config.json")}"',
            f"_TARGET_PKG = {(req.app_package or '')!r}",
            "",
            VISION_HELPER,
            MAP_TRACE_HELPER,
            MAP_RECOVERY_HELPER,
            "",
            "try:",
        ]
        for line in steps_code.splitlines():
            lines.append("    " + line if line.strip() else "")
        lines.append("    pass  # adım kodu boşsa (örn. oturum başlatma) try bloğu boş kalmasın")
        if session_mode == "end":
            lines += [
                "",
                "except Exception as e:",
                '    print(f"HATA: {e}")',
                "finally:",
                "    driver.quit()",
                '    print("SESSION_ENDED")',
            ]
        else:  # attach - quit YOK
            lines += [
                "",
                "except Exception as e:",
                '    print(f"HATA: {e}")',
                "    raise",
            ]
        return "\n".join(lines)

    # session_mode in ("fresh", "start") - asagidaki mevcut kod aynen devam eder

    if is_ios:
        # iPhone: XCUITest (WebDriverAgent imzası Mac'teki Appium'un varsayılan yeteneklerinde)
        caps = (f'options.bundle_id              = "{pkg}"\n' if pkg else "") + '''\
options.no_reset               = True
options.set_capability("appium:wdaLaunchTimeout", 180000)
options.set_capability("appium:newCommandTimeout", 900)'''
        pre = ""
    elif act:
        caps = f'''\
options.app_package            = "{pkg}"
options.app_activity           = "{act}"
options.no_reset               = True
options.auto_grant_permissions = True'''
        pre = ""
    else:
        caps = f'''\
options.app_package            = "{pkg}"
options.no_reset               = True
options.auto_grant_permissions = True'''
        pre = f'''\
import subprocess
print("Uygulama ADB ile baslatiliyor: {pkg}")
subprocess.run(
    ["adb", "-s", "{device}", "shell", "monkey",
     "-p", "{pkg}", "-c", "android.intent.category.LAUNCHER", "1"],
    capture_output=True
)
time.sleep(3)
'''

    lines = [
        "# -*- coding: utf-8 -*-",
        *env_lines,
        "from appium import webdriver",
        "from appium.options.android import UiAutomator2Options",
        "from appium.webdriver.common.appiumby import AppiumBy",
        "from selenium.webdriver.support.ui import WebDriverWait",
        "from selenium.webdriver.support import expected_conditions as EC",
        "import time, socket",
        "_VARS = {}  # adımlar arası değişken havuzu",
        "def _sub_vars(_s):",
        "    for _k, _v in _VARS.items():",
        "        _s = _s.replace('{{' + _k + '}}', str(_v))",
        "    return _s",
        "",
        STOP_HELPER,
        f'_DEVICE_ID = "{device}"  # ADB komutları için cihaz ID',
        f'_AI_CONFIG_PATH = r"{os.path.join(_ROUTER_BASE_DIR, "ai_config.json")}"',
        "",
        "# ── Appium bağlantı kontrolü ─────────────────────────────────",
        f'_appium_host = "{appium.split("//")[-1].split("/")[0]}"',
        f'_appium_port = {appium.rstrip("/").split(":")[-1] if ":" in appium.split("//")[-1] else 4723}',
        "try:",
        "    _s = socket.create_connection((_appium_host.split(':')[0], _appium_port), timeout=3)",
        "    _s.close()",
        "except OSError:",
        f'    print(f"[HATA] Appium sunucusu {appium} adresinde çalışmıyor!")',
        '    print("[HATA] Lütfen terminalde \\"appium\\" komutunu çalıştırın.")',
        "    raise SystemExit(1)",
        "",
        VISION_HELPER,
        MAP_TRACE_HELPER,
        MAP_RECOVERY_HELPER,
        "",
    ]

    if pre:
        lines += pre.splitlines() + [""]

    # Cihaz seçimi: Appium 2 / UiAutomator2'de deviceName güvenilir bir seçici DEĞİLDİR,
    # asıl seçici udid'dir. device_id verilmişse udid de yazılır (ek satır — device_name
    # davranışı aynen korunur). system_port yalnızca job motoru gönderdiğinde eklenir.
    _opt_lines = [
        "from appium.options.ios import XCUITestOptions",
        "options = XCUITestOptions()",
        'options.platform_name = "iOS"',
        f'options.device_name   = "{device}"',
    ] if is_ios else [
        "options = UiAutomator2Options()",
        f'options.platform_name = "Android"',
        f'options.device_name   = "{device}"',
    ]
    if req.device_id:
        _opt_lines.append(f'options.udid          = "{device}"')
    if is_ios:
        # Aynı Mac'teki her iPhone'un WebDriverAgent'ı ayrı yerel portta (varsayılan 8100
        # paylaşılırsa iki iPhone'da eşzamanlı oturum çakışır); telefon başına sabit port
        from services.ai_job_service import _acquire_system_port
        _wda = int(getattr(req, "system_port", None) or _acquire_system_port(device))
        _opt_lines.append(f'options.set_capability("appium:wdaLocalPort", {_wda})')
    elif getattr(req, "system_port", None):
        _opt_lines.append(f'options.system_port   = {int(req.system_port)}')
    if mjpeg_port:   # kalıcı oturum: telefon ekranı canlı MJPEG akışı (Telefon paneli / kayıt modu)
        _opt_lines.append(f'options.set_capability("appium:mjpegServerPort", {int(mjpeg_port)})')

    lines += _opt_lines + caps.splitlines() + [
        "",
        f'print("Appium baglaniyor: {appium}")',
        f'driver = webdriver.Remote("{appium}", options=options)',
        f"_IS_IOS = {is_ios}",
        IOS_SHIM_HELPER,
        "_ios_install(driver)",
        "wait   = WebDriverWait(driver, 15)",
        "time.sleep(2)",
        "",
        f'# Uygulama açık değilse başlat',
        f'_TARGET_PKG = "{pkg}"',
        "if _TARGET_PKG and _IS_IOS:",
        "    try:",
        "        driver.activate_app(_TARGET_PKG)   # bundle_id ile zaten açıldıysa öne getirir",
        "    except Exception as _launch_err:",
        f'        print(f"activate_app başarısız: {{_launch_err}}")',
        "elif _TARGET_PKG:",
        "    try:",
        "        _curr_pkg = driver.current_package",
        "        if _curr_pkg != _TARGET_PKG:",
        f'            print(f"Uygulama başlatılıyor: {{_TARGET_PKG}} (şu an: {{_curr_pkg}})")',
        "            driver.activate_app(_TARGET_PKG)",
        "            time.sleep(2.5)",
        "        else:",
        f'            print(f"Uygulama zaten açık: {{_TARGET_PKG}}")',
        "    except Exception as _launch_err:",
        f'        print(f"activate_app başarısız: {{_launch_err}}, ADB ile deneniyor...")',
        "        import subprocess as _sub",
        f'        _sub.run(["adb", "-s", _DEVICE_ID, "shell", "monkey",',
        f'                  "-p", _TARGET_PKG, "-c", "android.intent.category.LAUNCHER", "1"],',
        "                 capture_output=True, timeout=5)",
        "        time.sleep(3)",
        "",
        "try:",
    ]

    for line in steps_code.splitlines():
        lines.append("    " + line if line.strip() else "")
    lines.append("    pass  # adım kodu boşsa (örn. oturum başlatma) try bloğu boş kalmasın")

    lines += [
        "",
        "except Exception as e:",
        '    print(f"HATA: {e}")',
        "    raise",
        "finally:",
        ("    pass" if session_mode == "start" else "    driver.quit()"),
    ] + (
        [
            "",
            'print("SESSION_ID:", driver.session_id)',
            'print("EXECUTOR_URL:", driver.command_executor._url)',
            "import sys as _sys_st, os as _os_st",
            "_sys_st.stdout.flush()",
            # Service.__del__ normal script sonunda chromedriver/appium'u kapatir - os._exit ile atlanir
            "_os_st._exit(0)",
        ]
        if session_mode == "start" else []
    )

    return "\n".join(lines)


def _build_web_template(
    req: GenerateRequest,
    steps_code: str,
    session_mode: str = "fresh",
    session_id: str = "",
    command_executor_url: str = "",
    vars_seed: dict | None = None,
) -> str:
    browser = (req.browser or "chrome").lower()
    url     = (req.target_url or "").strip()

    if session_mode in ("attach", "end"):
        lines = [
            "# -*- coding: utf-8 -*-",
            "import sys",
            f'sys.path.insert(0, r"{_ROUTER_BASE_DIR}")',
            "from services.session_reconnect import attach_selenium_session",
            "from selenium.webdriver.common.by import By",
            "from selenium.webdriver.support.ui import WebDriverWait",
            "from selenium.webdriver.support import expected_conditions as EC",
            "import time",
            f"_VARS = {(vars_seed or {})!r}  # adımlar arası değişken havuzu",
            "def _sub_vars(_s):",
            "    for _k, _v in _VARS.items():",
            "        _s = _s.replace('{{' + _k + '}}', str(_v))",
            "    return _s",
            "",
            f"driver = attach_selenium_session({command_executor_url!r}, {session_id!r})",
            "wait = WebDriverWait(driver, 10)",
            "",
            STOP_HELPER,
            IMAGE_HELPER,
            (IMAGE_HELPER_WEB_CDP if browser in ("chrome", "edge") else ""),
            VISION_HELPER_WEB,
            "",
            "try:",
        ]
        for line in steps_code.splitlines():
            lines.append("    " + line if line.strip() else "")
        lines.append("    pass  # adım kodu boşsa (örn. oturum başlatma) try bloğu boş kalmasın")
        if session_mode == "end":
            lines += [
                "",
                "except Exception as e:",
                '    print(f"HATA: {e}")',
                "finally:",
                "    driver.quit()",
                '    print("SESSION_ENDED")',
            ]
        else:  # attach - quit YOK
            lines += [
                "",
                "except Exception as e:",
                '    print(f"HATA: {e}")',
                "    raise",
            ]
        return "\n".join(lines)

    # session_mode in ("fresh", "start") - asagidaki mevcut kod aynen devam eder

    if browser == "firefox":
        _fx_headless = "" if session_mode == "start" else "_fx_opts.add_argument('--headless')\n"
        driver_setup = (
            "from selenium.webdriver.firefox.options import Options as _FxOpts\n"
            "from selenium.webdriver.firefox.service import Service as _FxSvc\n"
            "_fx_opts = _FxOpts()\n"
            f"{_fx_headless}"
            "_fx_opts.add_argument('--width=1920')\n"
            "_fx_opts.add_argument('--height=1080')\n"
            "try:\n"
            "    from webdriver_manager.firefox import GeckoDriverManager\n"
            "    driver = webdriver.Firefox(service=_FxSvc(GeckoDriverManager().install()), options=_fx_opts)\n"
            "except Exception:\n"
            "    driver = webdriver.Firefox(options=_fx_opts)"
        )
    elif browser == "edge":
        _edge_headless = "" if session_mode == "start" else "_edge_opts.add_argument('--headless')\n"
        driver_setup = (
            "from selenium.webdriver.edge.options import Options as _EdgeOpts\n"
            "_edge_opts = _EdgeOpts()\n"
            f"{_edge_headless}"
            "_edge_opts.add_argument('--window-size=1920,1080')\n"
            "_edge_opts.add_argument('--disable-gpu')\n"
            "try:\n"
            "    from webdriver_manager.microsoft import EdgeChromiumDriverManager\n"
            "    from selenium.webdriver.edge.service import Service as _EdgeSvc\n"
            "    driver = webdriver.Edge(service=_EdgeSvc(EdgeChromiumDriverManager().install()), options=_edge_opts)\n"
            "except Exception:\n"
            "    driver = webdriver.Edge(options=_edge_opts)"
        )
    else:  # chrome (default)
        _cr_headless = "" if session_mode == "start" else "_cr_opts.add_argument('--headless=new')\n"
        driver_setup = (
            "from selenium.webdriver.chrome.options import Options as _CrOpts\n"
            "_cr_opts = _CrOpts()\n"
            f"{_cr_headless}"
            "_cr_opts.add_argument('--window-size=1920,1080')\n"
            "_cr_opts.add_argument('--disable-gpu')\n"
            "_cr_opts.add_argument('--no-sandbox')\n"
            "_cr_opts.add_argument('--disable-dev-shm-usage')\n"
            "try:\n"
            "    from webdriver_manager.chrome import ChromeDriverManager\n"
            "    from selenium.webdriver.chrome.service import Service as _CrSvc\n"
            "    driver = webdriver.Chrome(service=_CrSvc(ChromeDriverManager().install()), options=_cr_opts)\n"
            "except Exception:\n"
            "    driver = webdriver.Chrome(options=_cr_opts)"
        )

    lines = [
        "# -*- coding: utf-8 -*-",
        "from selenium import webdriver",
        "from selenium.webdriver.common.by import By",
        "from selenium.webdriver.support.ui import WebDriverWait",
        "from selenium.webdriver.support import expected_conditions as EC",
        "import time",
        "_VARS = {}  # adımlar arası değişken havuzu",
        "def _sub_vars(_s):",
        "    for _k, _v in _VARS.items():",
        "        _s = _s.replace('{{' + _k + '}}', str(_v))",
        "    return _s",
        "",
        driver_setup,
        "",
        "# ── Tarayıcı başlatılamadıysa net hata ver ───────────────────",
        "if 'driver' not in dir():",
        '    print("[HATA] Tarayıcı sürücüsü başlatılamadı. Chrome/Firefox kurulu mu?")',
        "    raise SystemExit(1)",
        "",
        STOP_HELPER,
        IMAGE_HELPER,
        (IMAGE_HELPER_WEB_CDP if browser in ("chrome", "edge") else ""),
        VISION_HELPER_WEB,
        "",
        "driver.maximize_window()",
        "wait = WebDriverWait(driver, 10)",
        "",
        "try:",
    ]

    if url:
        lines += [
            f'    driver.get("{url}")',
            "    time.sleep(1.5)",
        ]

    for line in steps_code.splitlines():
        lines.append("    " + line if line.strip() else "")
    lines.append("    pass  # adım kodu boşsa (örn. oturum başlatma) try bloğu boş kalmasın")

    lines += [
        "",
        "except Exception as e:",
        '    print(f"HATA: {e}")',
        "    raise",
        "finally:",
        ("    pass" if session_mode == "start" else "    driver.quit()"),
    ] + (
        [
            "",
            'print("SESSION_ID:", driver.session_id)',
            'print("EXECUTOR_URL:", driver.command_executor._url)',
            "import sys as _sys_st, os as _os_st",
            "_sys_st.stdout.flush()",
            # Service.__del__ normal script sonunda chromedriver/appium'u kapatir - os._exit ile atlanir
            "_os_st._exit(0)",
        ]
        if session_mode == "start" else []
    )

    return "\n".join(lines)


def _build_desktop_template(req: GenerateRequest, steps_code: str, vars_seed: dict | None = None) -> str:
    app_path      = (req.app_path     or "").strip()
    window_title  = (req.window_title or "").strip()
    reuse_instance = getattr(req, "reuse_instance", True)

    # Escape backslashes so Windows paths don't cause unicode escape SyntaxError
    app_path_esc     = app_path.replace('\\', '\\\\')
    window_title_esc = window_title.replace('\\', '\\\\')

    # _POPEN_KW: uygulama başlatırken stdout/stderr/stdin devralmasın, pipe kapanmasın
    _POPEN_DEVNULL = (
        "_dsproc.Popen(r\"{}\", shell=True, "
        "stdout=_dsproc.DEVNULL, stderr=_dsproc.DEVNULL, stdin=_dsproc.DEVNULL)"
    )

    lines = [
        "# -*- coding: utf-8 -*-",
        "import pyautogui as _pag",
        "import subprocess as _dsproc",
        "import sys",
        "import time",
        "",
        f'_AI_CONFIG_PATH = r"{os.path.join(_ROUTER_BASE_DIR, "ai_config.json")}"',
        f'_WINDOW_TITLE   = "{window_title_esc}"',
        f'_APP_PATH       = r"{app_path_esc}"',
        f"_VARS = {(vars_seed or {})!r}  # adımlar arası değişken havuzu",
        "def _sub_vars(_s):",
        "    for _k, _v in _VARS.items():",
        "        _s = _s.replace('{{' + _k + '}}', str(_v))",
        "    return _s",
        "",
        STOP_HELPER,
        IMAGE_HELPER,
        DESKTOP_HELPER,
        "",
    ]

    # Uygulama başlat
    if app_path:
        popen_line = _POPEN_DEVNULL.format(app_path_esc)
        if reuse_instance:
            lines += [
                '# ── Uygulama başlat (zaten açıksa geç + öne getir) ──────────',
                'def _is_app_running():',
                '    """Uygulama çalışıyorsa True döner ve penceresini öne getirir."""',
                '    # 1) Pencere başlığıyla kontrol (win32gui) ─────────────────',
                f'    _wtitle = "{window_title_esc}"',
                '    if _wtitle:',
                '        try:',
                '            import win32gui as _wg2',
                '            _fnd = []',
                '            def _cb2(hwnd, _):',
                '                t = _wg2.GetWindowText(hwnd)',
                '                if _wtitle.lower() in t.lower() and _wg2.IsWindowVisible(hwnd):',
                '                    _fnd.append(hwnd)',
                '                return True',
                '            _wg2.EnumWindows(_cb2, None)',
                '            if _fnd:',
                '                try: _wg2.SetForegroundWindow(_fnd[0])',
                '                except Exception: pass',
                '                return True',
                '        except Exception: pass',
                '        # pywinauto fallback ─────────────────────────────────',
                '        if _PW_OK:',
                '            try:',
                f'                _PwApp(backend="uia").connect(title_re=".*{window_title_esc}.*", timeout=2).top_window().set_focus()',
                '                return True',
                '            except Exception: pass',
                '    # 2) Surec adiyla kontrol + PID bazli pencere odakla ───────',
                '    try:',
                '        import os as _os2, subprocess as _sp2',
                f'        _exe2 = _os2.path.basename(r"{app_path_esc}")',
                '        if not _exe2: return False',
                '        _proc_found = False',
                '        _boa_pids: set = set()',
                '        # psutil varsa kullan (daha güvenilir) ─────────────────',
                '        try:',
                '            import psutil as _psu2',
                '            for _p2 in _psu2.process_iter(["name", "pid"]):',
                '                if _exe2.lower() == _p2.info["name"].lower():',
                '                    _boa_pids.add(_p2.info["pid"]); _proc_found = True',
                '        except ImportError:',
                '            # tasklist fallback ──────────────────────────────',
                '            _r2 = _sp2.run(["tasklist", "/FI", f"IMAGENAME eq {_exe2}", "/NH"],',
                '                           capture_output=True, text=True, timeout=3)',
                '            _proc_found = _exe2.lower() in _r2.stdout.lower()',
                '        if not _proc_found: return False',
                '        # Sürece ait pencereyi bul ve öne getir ─────────────────',
                '        try:',
                '            import win32gui as _wg3, win32process as _wp3',
                '            _boa_wins = []',
                '            def _cb3(hwnd, _):',
                '                if _wg3.IsWindowVisible(hwnd) and _wg3.GetWindowText(hwnd):',
                '                    _tid, _pid = _wp3.GetWindowThreadProcessId(hwnd)',
                '                    if not _boa_pids or _pid in _boa_pids:',
                '                        _boa_wins.append(hwnd)',
                '                return True',
                '            _wg3.EnumWindows(_cb3, None)',
                '            if _boa_wins: _wg3.SetForegroundWindow(_boa_wins[0])',
                '        except Exception: pass',
                '        return True',
                '    except Exception: pass',
                '    return False',
                'if _is_app_running():',
                '    print("[OK] Uygulama zaten açık — mevcut pencere öne getirildi")',
                'else:',
                f'    print("Başlatılıyor: {app_path_esc}")',
                f'    {popen_line}',
                '    # BOA penceresini bekle (max 30 saniye)',
                '    _wait_start = time.time()',
                '    _win_found = False',
                f'    _wait_exe = "{os.path.basename(app_path_esc).lower()}"',
                f'    _wait_title = "{window_title_esc}"',
                '    while time.time() - _wait_start < 30:',
                '        try:',
                '            import win32gui as _wg_w, win32process as _wp_w, win32con as _wc_w',
                '            # Süreç PID\'lerini bul',
                '            _pids_w = set()',
                '            try:',
                '                import psutil as _psu_w',
                '                _pids_w = {p.pid for p in _psu_w.process_iter(["name"]) if p.info["name"].lower() == _wait_exe}',
                '            except ImportError:',
                '                import subprocess as _sp_w',
                '                _r_w = _sp_w.run(["tasklist", "/FI", f"IMAGENAME eq {_wait_exe}", "/FO", "CSV", "/NH"], capture_output=True, text=True, timeout=5)',
                '                for _ln_w in _r_w.stdout.splitlines():',
                '                    _parts_w = _ln_w.strip(\'"\').split(\'","\')',
                '                    if len(_parts_w) > 1:',
                '                        try: _pids_w.add(int(_parts_w[1]))',
                '                        except: pass',
                '            if not _pids_w:',
                '                time.sleep(1); continue',
                '            _wins_w = []',
                '            def _cb_w(hwnd, _):',
                '                if _wg_w.IsWindowVisible(hwnd) and _wg_w.GetWindowText(hwnd):',
                '                    _, _pid_w = _wp_w.GetWindowThreadProcessId(hwnd)',
                '                    if _pid_w in _pids_w:',
                '                        _t_w = _wg_w.GetWindowText(hwnd)',
                '                        if not _wait_title or _wait_title.lower() in _t_w.lower():',
                '                            _wins_w.append(hwnd)',
                '                return True',
                '            _wg_w.EnumWindows(_cb_w, None)',
                '            if _wins_w:',
                '                _wg_w.ShowWindow(_wins_w[0], _wc_w.SW_RESTORE)',
                '                _wg_w.SetForegroundWindow(_wins_w[0])',
                '                _win_found = True',
                '                print(f"[OK] Pencere {time.time()-_wait_start:.1f}s sonra görüldü")',
                '                break',
                '        except Exception: pass',
                '        time.sleep(1)',
                '    if not _win_found: print("[WARN] Uygulama penceresi 30s içinde görünmedi, devam ediliyor")',
                '    time.sleep(2)',
                '',
            ]
        else:
            lines += [
                '# ── Uygulama başlat ─────────────────────────────────────────',
                f'print("Başlatılıyor: {app_path_esc}")',
                popen_line,
                'time.sleep(2)',
                '',
            ]

    # Pencereyi öne getir
    if window_title:
        lines += [
            '# ── Pencereyi öne getir ─────────────────────────────────────',
            'if _PW_OK:',
            '    try:',
            f'        _PwApp(backend="uia").connect(title_re=".*{window_title_esc}.*", timeout=10).top_window().set_focus()',
            f'        print("[OK] Pencere: {window_title_esc}")',
            '    except Exception as _we:',
            '        print(f"[WARN] Pencere bulunamadı: {{_we}}")',
            '',
        ]

    lines += ["try:"]
    for line in steps_code.splitlines():
        lines.append("    " + line if line.strip() else "")
    lines.append("    pass  # adım kodu boşsa (örn. oturum başlatma) try bloğu boş kalmasın")

    lines += [
        "",
        "except Exception as e:",
        '    print(f"HATA: {e}")',
        "    raise",
        "",
        "sys.exit(0)",
    ]

    return "\n".join(lines)


# ══════════════════════════════════════════════════════════════════════════════
# BÖLÜM 4: ANA DÖNÜŞÜM MOTORU
# ══════════════════════════════════════════════════════════════════════════════

def _image_click_code(idx: int, desc: str, b64: str, click_x: int, click_y: int, action: str, platform: str, browser: str = "") -> str:
    """Görsel tıklama adımı için Desktop+Web ortak kod üretir."""
    refocus_line = "    _refocus_window()\n" if platform == "desktop" else ""
    action_str = {"double_click": "double_click", "right_click": "right_click"}.get(action, "click")
    is_web_cdp = platform == "web" and browser in ("chrome", "edge")
    is_web_firefox = platform == "web" and browser not in ("chrome", "edge")

    if click_x >= 0 and click_y >= 0:
        if is_web_firefox:
            return (
                f"# [NOKTA TIKLAMA] {desc}\n"
                f"print('STEP_START:{idx}')\n"
                f"print('STEP_FAIL:{idx}')\n"
                f"raise Exception('Görsel adımlar şu anda Firefox\\'ta desteklenmiyor. Chrome veya Edge kullanın.')"
            )
        click_fn = "_find_b64_and_click_point_web" if is_web_cdp else "_find_b64_and_click_point"
        click_args = f"driver, _img_b64_{idx}, {click_x}, {click_y}, action='{action_str}')" if is_web_cdp else f"_img_b64_{idx}, {click_x}, {click_y}, action='{action_str}')"
        return (
            f"# [NOKTA TIKLAMA] {desc}\n"
            f"print('STEP_START:{idx}')\n"
            f"try:\n"
            f"{refocus_line}"
            f"    _img_b64_{idx} = '{b64}'\n"
            f"    _ok_{idx} = {click_fn}({click_args}\n"
            f"    if not _ok_{idx}: raise Exception('Nokta tiklanamiyor: {desc}')\n"
            f"    time.sleep(0.5)\n"
            f"    print('STEP_DONE:{idx}')\n"
            f"except Exception as _e_{idx}:\n"
            f"    print(f'[FAIL] {{_e_{idx}}}')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise"
        )

    if is_web_firefox:
        return (
            f"# [GÖRSEL TIKLAMA] {desc}\n"
            f"print('STEP_START:{idx}')\n"
            f"print('STEP_FAIL:{idx}')\n"
            f"raise Exception('Görsel adımlar şu anda Firefox\\'ta desteklenmiyor. Chrome veya Edge kullanın.')"
        )

    if is_web_cdp:
        find_fn = "_find_b64_in_page_web"
        find_args = f"driver, _img_b64_{idx}, timeout=10)"
        click_line = {
            "double_click": f"    _cdp_click(driver, _coords_{idx}[0], _coords_{idx}[1], click_count=2)",
            "right_click":  f"    _cdp_click(driver, _coords_{idx}[0], _coords_{idx}[1], button='right')",
        }.get(action, f"    _cdp_click(driver, _coords_{idx}[0], _coords_{idx}[1])")
    else:
        find_fn = "_find_b64_on_screen" if platform == "desktop" else "_find_b64_on_screen_web"
        find_args = f"_img_b64_{idx}, timeout=10)"
        click_line = {
            "double_click": f"    import pyautogui as _pag_i; _pag_i.doubleClick(_coords_{idx}[0], _coords_{idx}[1])",
            "right_click":  f"    import pyautogui as _pag_i; _pag_i.rightClick(_coords_{idx}[0], _coords_{idx}[1])",
        }.get(action, f"    _do_robust_click(_coords_{idx}[0], _coords_{idx}[1])")
    return (
        f"# [GÖRSEL TIKLAMA] {desc}\n"
        f"print('STEP_START:{idx}')\n"
        f"try:\n"
        f"{refocus_line}"
        f"    _img_b64_{idx} = '{b64}'\n"
        f"    _coords_{idx} = {find_fn}({find_args}\n"
        f"    if not _coords_{idx}: raise Exception('Gorsel bulunamadi: {desc}')\n"
        f"{click_line}\n"
        f"    print('[OK] Gorsel tiklandi: {desc}')\n"
        f"    time.sleep(0.5)\n"
        f"    print('STEP_DONE:{idx}')\n"
        f"except Exception as _e_{idx}:\n"
        f"    print(f'[FAIL] {{_e_{idx}}}')\n"
        f"    print('STEP_FAIL:{idx}')\n"
        f"    raise"
    )


def _image_step_code(item: "StepItem", platform: str, idx: int, browser: str = "") -> str:
    """Görsel adımı için Python kodu üretir. base64 scripte gömülür."""
    b64  = item.image_b64
    desc = item.content or f"Görsel adım {idx+1}"
    is_web_cdp = platform == "web" and browser in ("chrome", "edge")
    is_web_firefox = platform == "web" and browser not in ("chrome", "edge")

    if item.action == "drag":
        if is_web_firefox:
            return (
                f"# [GÖRSEL SÜRÜKLE] {desc}\n"
                f"print('STEP_START:{idx}')\n"
                f"print('STEP_FAIL:{idx}')\n"
                f"raise Exception('Görsel adımlar şu anda Firefox\\'ta desteklenmiyor. Chrome veya Edge kullanın.')"
            )
        refocus_line = "    _refocus_window()\n" if platform == "desktop" else ""
        drag_fn = "_find_b64_and_drag_web" if is_web_cdp else "_find_b64_and_drag"
        drag_args = (
            f"driver, _img_b64_{idx}, {item.click_x}, {item.click_y}, {item.drag_x}, {item.drag_y})"
            if is_web_cdp else
            f"_img_b64_{idx}, {item.click_x}, {item.click_y}, {item.drag_x}, {item.drag_y})"
        )
        return (
            f"# [GÖRSEL SÜRÜKLE] {desc}\n"
            f"print('STEP_START:{idx}')\n"
            f"try:\n"
            f"{refocus_line}"
            f"    _img_b64_{idx} = '{b64}'\n"
            f"    _ok_{idx} = {drag_fn}({drag_args}\n"
            f"    if not _ok_{idx}: raise Exception('Surukleme yapilamadi: {desc}')\n"
            f"    time.sleep(0.5)\n"
            f"    print('STEP_DONE:{idx}')\n"
            f"except Exception as _e_{idx}:\n"
            f"    print(f'[FAIL] {{_e_{idx}}}')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise"
        )

    if item.action == "assert":
        if is_web_firefox:
            return (
                f"# [GÖRSEL DOĞRULA] {desc}\n"
                f"print('STEP_START:{idx}')\n"
                f"print('STEP_FAIL:{idx}')\n"
                f"raise Exception('Görsel adımlar şu anda Firefox\\'ta desteklenmiyor. Chrome veya Edge kullanın.')"
            )
        refocus_line = "    _refocus_window()\n" if platform == "desktop" else ""
        if is_web_cdp:
            find_fn = "_find_b64_in_page_web"
            find_args = f"driver, _img_b64_{idx}, timeout=8)"
        else:
            find_fn = "_find_b64_on_screen" if platform == "desktop" else "_find_b64_on_screen_web"
            find_args = f"_img_b64_{idx}, timeout=8)"
        return (
            f"# [GÖRSEL DOĞRULA] {desc}\n"
            f"print('STEP_START:{idx}')\n"
            f"try:\n"
            f"{refocus_line}"
            f"    _img_b64_{idx} = '{b64}'\n"
            f"    _res_{idx} = {find_fn}({find_args}\n"
            f"    if not _res_{idx}: raise Exception('Gorsel bulunamadi: {desc}')\n"
            f"    print('[OK] Gorsel dogrulandi: {desc}')\n"
            f"    print('STEP_DONE:{idx}')\n"
            f"except Exception as _e_{idx}:\n"
            f"    print(f'[FAIL] {{_e_{idx}}}')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise"
        )

    return _image_click_code(idx, desc, b64, item.click_x, item.click_y, item.action, platform, browser)


# Yaygından en yeniye, script'in ÇALIŞTIĞI makinede kurulu ilk uyumlu sürücü seçilir
# (kod üretim anındaki makine ile çalışma zamanındaki makine farklı olabilir — IM002 hatasının nedeni buydu)
_SQLSERVER_ODBC_DRIVER_CANDIDATES = [
    "ODBC Driver 18 for SQL Server",
    "ODBC Driver 17 for SQL Server",
    "ODBC Driver 13 for SQL Server",
    "SQL Server Native Client 11.0",
    "SQL Server",
]


def _parse_jdbc_sqlserver(conn: str) -> Optional[dict]:
    """'jdbc:sqlserver://host;serverName=..;databaseName=..;user=..;password=..' formatını
    bileşenlerine ayırır. JDBC URL'leri bir ODBC driver adı belirtmez (JDBC'nin kendi
    sürücüsü vardır) — bu yüzden driver adı burada SEÇİLMEZ, çalışma zamanında hedef
    makinedeki kurulu sürücülere bakılarak seçilir (bkz. _sql_step_code)."""
    body = conn.strip()
    if not body.lower().startswith("jdbc:sqlserver:"):
        return None
    body = body[len("jdbc:sqlserver:"):]
    if body.startswith("//"):
        body = body[2:]
    parts = body.split(";")
    host_port = parts[0].strip()

    props = {}
    for part in parts[1:]:
        part = part.strip()
        if not part:
            continue
        sep = "=" if "=" in part else (":" if ":" in part else None)
        if not sep:
            continue
        key, val = part.split(sep, 1)
        props[key.strip().lower()] = val.strip()

    server = props.get("servername") or host_port.split(":")[0]
    if not server:
        return None
    port = props.get("portnumber") or (host_port.split(":", 1)[1] if ":" in host_port else "")
    return {
        "server": server,
        "port": port,
        "database": props.get("databasename", ""),
        "user": props.get("user", ""),
        "password": props.get("password", ""),
    }


def _sql_step_code(item: "StepItem", idx: int) -> str:
    """SQL adımı için Python kodu üretir."""
    desc  = item.content or f"SQL {idx+1}"
    conn  = item.db_connection or "sqlite:///test.db"
    query = item.sql_query or "SELECT 1"
    jdbc_info = _parse_jdbc_sqlserver(conn)
    is_print  = item.action != "execute"
    is_sqlite = conn.strip().lower().startswith("sqlite:///")
    # DRIVER= içeriyorsa doğrudan pyodbc — SQLAlchemy URL parsing sorununu önler
    is_odbc   = "DRIVER=" in conn.upper()

    conn_r  = repr(conn)
    query_r = repr(query)

    if jdbc_info:
        # JDBC bağlantı string'i: ODBC driver adı çalışma zamanında hedef makinede seçilir
        var_store = ""
        if item.var_name:
            var_store = (
                f"\n    _VARS[{repr(item.var_name)}] = str(_rows_{idx}[0][0] if _rows_{idx} else '')"
                f'\n    print(f"[SQL→_VARS] {item.var_name} = {{_VARS[{repr(item.var_name)}]}}")'
            )
        if is_print:
            result_code = (
                f"    _rows_{idx} = _cur_{idx}.fetchall()\n"
                f"    print(f'[SQL] {{len(_rows_{idx})}} kayit:')\n"
                f"    for _r in _rows_{idx}: print('  ', list(_r))"
                f"{var_store}"
            )
        else:
            result_code = (
                f"    _db_{idx}.commit()\n"
                f"    print(f'[SQL] {{_cur_{idx}.rowcount}} satir etkilendi')"
                f"{var_store}"
            )
        # 'tcp:' onekiyle protokol acikca TCP/IP'ye zorlanir — aksi halde surucu Named Pipes
        # deneyebilir ve ozel IP/uzak sunucularda "Named Pipes Provider" hatasi verir
        server_part = f"tcp:{jdbc_info['server']},{jdbc_info['port'] or '1433'}"
        auth_part = (
            f"UID={jdbc_info['user']};PWD={jdbc_info['password']}"
            if jdbc_info['user'] else "Trusted_Connection=yes"
        )
        conn_tail_r = repr(f"SERVER={server_part};DATABASE={jdbc_info['database']};{auth_part}")
        drv_candidates_r = repr(_SQLSERVER_ODBC_DRIVER_CANDIDATES)
        return (
            f"# [SQL] {desc}\n"
            f"print('STEP_START:{idx}')\n"
            f"try:\n"
            f"    import pyodbc as _pyodbc_{idx}\n"
            f"    _avail_{idx} = _pyodbc_{idx}.drivers()\n"
            f"    _drv_{idx} = next((_d for _d in {drv_candidates_r} if _d in _avail_{idx}), None)\n"
            f"    if not _drv_{idx}: raise Exception('Bu makinede SQL Server ODBC suruculeri bulunamadi. Kurulu suruculer: ' + ', '.join(_avail_{idx}))\n"
            f"    print(f'[SQL] ODBC surucu: {{_drv_{idx}}}')\n"
            f"    if _drv_{idx} == 'SQL Server': print('[SQL] UYARI: eski/deprecated \"SQL Server\" (DBNETLIB) suruculer kullaniliyor - bu makinede ODBC Driver 17/18 for SQL Server kurulu degil, TLS/encryption gerektiren sunucularda baglanti basarisiz olabilir')\n"
            f"    _connstr_{idx} = 'DRIVER=' + '{{' + _drv_{idx} + '}}' + ';' + {conn_tail_r}\n"
            f"    _db_{idx} = _pyodbc_{idx}.connect(_connstr_{idx})\n"
            f"    _cur_{idx} = _db_{idx}.cursor()\n"
            f"    _cur_{idx}.execute(_sub_vars({query_r}))\n"
            f"{result_code}\n"
            f"    _db_{idx}.close()\n"
            f"    print('STEP_DONE:{idx}')\n"
            f"except ImportError:\n"
            f"    print('[FAIL] pyodbc kurulu degil: pip install pyodbc (ayrica Microsoft ODBC Driver 17/18 for SQL Server gerekir)')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise\n"
            f"except Exception as _e_{idx}:\n"
            f"    print(f'[FAIL] {{_e_{idx}}}')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise"
        )
    elif is_sqlite:
        db_path = conn.strip()[len("sqlite:///"):]
        var_store = ""
        if item.var_name:
            var_store = (
                f"\n    _VARS[{repr(item.var_name)}] = str(_rows_{idx}[0][0] if _rows_{idx} else '')"
                f'\n    print(f"[SQL→_VARS] {item.var_name} = {{_VARS[{repr(item.var_name)}]}}")'
            )
        if is_print:
            result_code = (
                f"    _rows_{idx} = _cur_{idx}.fetchall()\n"
                f"    print(f'[SQL] {{len(_rows_{idx})}} kayit:')\n"
                f"    for _r in _rows_{idx}: print('  ', _r)"
                f"{var_store}"
            )
        else:
            result_code = (
                f"    _db_{idx}.commit()\n"
                f"    print(f'[SQL] {{_cur_{idx}.rowcount}} satir etkilendi')"
                f"{var_store}"
            )
        return (
            f"# [SQL] {desc}\n"
            f"print('STEP_START:{idx}')\n"
            f"try:\n"
            f"    import sqlite3 as _sq_{idx}\n"
            f"    _db_{idx} = _sq_{idx}.connect({repr(db_path)})\n"
            f"    _cur_{idx} = _db_{idx}.cursor()\n"
            f"    _cur_{idx}.execute(_sub_vars({query_r}))\n"
            f"{result_code}\n"
            f"    _db_{idx}.close()\n"
            f"    print('STEP_DONE:{idx}')\n"
            f"except Exception as _e_{idx}:\n"
            f"    print(f'[FAIL] {{_e_{idx}}}')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise"
        )
    elif is_odbc:
        # pyodbc doğrudan bağlantı (ODBC connection string — {DRIVER} parantezleri güvenli)
        var_store = ""
        if item.var_name:
            var_store = (
                f"\n    _VARS[{repr(item.var_name)}] = str(_rows_{idx}[0][0] if _rows_{idx} else '')"
                f'\n    print(f"[SQL→_VARS] {item.var_name} = {{_VARS[{repr(item.var_name)}]}}")'
            )
        if is_print:
            result_code = (
                f"    _rows_{idx} = _cur_{idx}.fetchall()\n"
                f"    print(f'[SQL] {{len(_rows_{idx})}} kayit:')\n"
                f"    for _r in _rows_{idx}: print('  ', list(_r))"
                f"{var_store}"
            )
        else:
            result_code = (
                f"    _db_{idx}.commit()\n"
                f"    print(f'[SQL] {{_cur_{idx}.rowcount}} satir etkilendi')"
                f"{var_store}"
            )
        return (
            f"# [SQL] {desc}\n"
            f"print('STEP_START:{idx}')\n"
            f"try:\n"
            f"    import pyodbc as _pyodbc_{idx}\n"
            f"    _db_{idx} = _pyodbc_{idx}.connect({conn_r})\n"
            f"    _cur_{idx} = _db_{idx}.cursor()\n"
            f"    _cur_{idx}.execute(_sub_vars({query_r}))\n"
            f"{result_code}\n"
            f"    _db_{idx}.close()\n"
            f"    print('STEP_DONE:{idx}')\n"
            f"except ImportError:\n"
            f"    print('[FAIL] pyodbc kurulu degil: pip install pyodbc (ayrica Microsoft ODBC Driver 17/18 for SQL Server gerekir)')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise\n"
            f"except Exception as _e_{idx}:\n"
            f"    print(f'[FAIL] {{_e_{idx}}}')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise"
        )
    else:
        var_store = ""
        if item.var_name:
            var_store = (
                f"\n        _VARS[{repr(item.var_name)}] = str(_rows_{idx}[0][0] if _rows_{idx} else '')"
                f'\n        print(f"[SQL→_VARS] {item.var_name} = {{_VARS[{repr(item.var_name)}]}}")'
            )
        if is_print:
            result_code = (
                f"        _rows_{idx} = list(_res_{idx})\n"
                f"        print(f'[SQL] {{len(_rows_{idx})}} kayit:')\n"
                f"        for _r in _rows_{idx}: print('  ', dict(_r._mapping))"
                f"{var_store}"
            )
        else:
            result_code = (
                f"        _conn_{idx}.commit()\n"
                f"        print(f'[SQL] {{_res_{idx}.rowcount}} satir etkilendi')"
                f"{var_store}"
            )
        return (
            f"# [SQL] {desc}\n"
            f"print('STEP_START:{idx}')\n"
            f"try:\n"
            f"    from sqlalchemy import create_engine as _ce_{idx}, text as _st_{idx}\n"
            f"    _eng_{idx} = _ce_{idx}({conn_r})\n"
            f"    with _eng_{idx}.connect() as _conn_{idx}:\n"
            f"        _res_{idx} = _conn_{idx}.execute(_st_{idx}(_sub_vars({query_r})))\n"
            f"{result_code}\n"
            f"    print('STEP_DONE:{idx}')\n"
            f"except ImportError:\n"
            f"    print('[FAIL] sqlalchemy kurulu degil: pip install sqlalchemy pymysql')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise\n"
            f"except Exception as _e_{idx}:\n"
            f"    print(f'[FAIL] {{_e_{idx}}}')\n"
            f"    print('STEP_FAIL:{idx}')\n"
            f"    raise"
        )


def _api_step_code(item: "StepItem", idx: int) -> str:
    """API adımı için Python kodu üretir."""
    desc    = item.content or f"API {idx+1}"
    url     = item.api_url or ""
    method  = (item.api_method or "GET").lower()
    headers = item.api_headers.strip() if item.api_headers else ""
    body    = item.api_body.strip()    if item.api_body    else ""
    action  = item.action or "print"

    headers_line = f"    _hdrs_{idx} = {headers}" if headers and headers.startswith("{") else f"    _hdrs_{idx} = {{}}"
    body_arg = ""
    if method in ("post", "put", "patch") and body and body.startswith("{"):
        body_line = f"    _body_{idx} = {body}\n"
        body_arg  = f", json=_body_{idx}"
    else:
        body_line = ""

    # json_field ve var_name için yardımcı kod
    json_field  = item.json_field.strip() if item.json_field else ""
    var_name    = item.var_name.strip()   if item.var_name   else ""

    extract_code = ""
    if json_field or var_name:
        extract_code = (
            f"    _api_json_{idx} = _resp_{idx}.json()\n"
        )
        if json_field:
            extract_code += (
                f"    def _jpath_{idx}(d, p):\n"
                f"        for k in p.split('.'):\n"
                f"            if isinstance(d, dict): d = d.get(k)\n"
                f"            elif isinstance(d, list):\n"
                f"                try: d = d[int(k)]\n"
                f"                except: return None\n"
                f"            else: return None\n"
                f"        return d\n"
                f"    _extracted_{idx} = _jpath_{idx}(_api_json_{idx}, {repr(json_field)})\n"
                f"    print(f'[API.{json_field}] {{_extracted_{idx}}}')\n"
            )
            if var_name:
                extract_code += (
                    f"    _VARS[{repr(var_name)}] = str(_extracted_{idx} if _extracted_{idx} is not None else '')\n"
                    f"    print(f'[API→_VARS] {var_name} = {{_VARS[{repr(var_name)}]}}')\n"
                )
        elif var_name:
            extract_code += (
                f"    _VARS[{repr(var_name)}] = str(_api_json_{idx})[:500]\n"
                f"    print(f'[API→_VARS] {var_name} = {{_VARS[{repr(var_name)}][:100]}}')\n"
            )

    if action == "print":
        resp_code = (
            f"    _resp_{idx}.raise_for_status()\n"
            f"    print(f'[API] Status: {{_resp_{idx}.status_code}}')\n"
            f"    print(f'[API] Yanit: {{_resp_{idx}.text[:500]}}')\n"
            f"{extract_code}"
        )
    elif action == "store":
        resp_code = (
            f"    _resp_{idx}.raise_for_status()\n"
            f"{extract_code}"
            f"    print(f'[API] Degisken: _api_result = {{str(_resp_{idx}.json())[:300]}}')"
        )
    else:
        resp_code = (
            f"    print(f'[API] Status: {{_resp_{idx}.status_code}}')  \n"
            f"{extract_code}"
        )

    return (
        f"# [API] {method.upper()} {url} — {desc}\n"
        f"print('STEP_START:{idx}')\n"
        f"try:\n"
        f"    import requests as _rm\n"
        f"{headers_line}\n"
        f"{body_line}"
        f"    _resp_{idx} = _rm.{method}({repr(url)}, headers=_hdrs_{idx}{body_arg}, timeout=30)\n"
        f"{resp_code}"
        f"    print('STEP_DONE:{idx}')\n"
        f"except Exception as _e_{idx}:\n"
        f"    print(f'[FAIL] {{_e_{idx}}}')\n"
        f"    print('STEP_FAIL:{idx}')\n"
        f"    raise"
    )


def _indent_code(code: str, prefix: str = "    ") -> str:
    """Çok satırlı bir kod bloğunun her satırının başına prefix ekler (boş satırlar hariç).
    _build_mobile/web/desktop_template'in steps_code'u indent'leme şeklinin aynısı."""
    return "\n".join((prefix + line if line.strip() else "") for line in code.splitlines())


def _ensure_block_body(code: str, indent: str) -> str:
    """if/else gövdesi olarak kullanılacak kodun GEÇERLİ bir Python bloğu olmasını garanti eder.
    Boşsa ya da (dal sadece 'yorum' tipi adımlardan oluştuğunda olduğu gibi) sadece '#' ile
    başlayan/boş satırlardan oluşuyorsa — yorumlar Python grameri için 'statement' sayılmaz ve
    tek başlarına bir blok gövdesi oluşturamaz — sona bir 'pass' eklenir."""
    has_executable_line = any(
        line.strip() and not line.strip().startswith("#")
        for line in code.splitlines()
    )
    if has_executable_line:
        return code
    pass_line = f"{indent}    pass"
    return f"{code}\n{pass_line}" if code else pass_line


def _condition_expr(item: "StepItem", platform: str, idx: int, browser: str = "") -> str:
    """Bir condition StepItem'ı, raise etmeyen bool döndüren tek satırlık bir Python ifadesine çevirir."""
    ctype = item.condition_type
    if ctype == "image_exists":
        if platform == "web" and browser in ("chrome", "edge"):
            return f"bool(_find_b64_in_page_web(driver, '{item.image_b64}', timeout=8))"
        if platform == "web":
            return "False  # Gorsel kosullar Firefox'ta desteklenmiyor"
        find_fn = "_find_b64_on_screen" if platform == "desktop" else "_find_b64_on_screen_web"
        return f"bool({find_fn}('{item.image_b64}', timeout=8))"
    if ctype == "element_exists":
        hint = item.content.replace("'", "\\'")
        if platform == "mobile":
            return f"_safe_assert(driver, '{hint}')"
        if platform == "web":
            return f"_safe_assert_web(driver, '{hint}', raise_on_fail=False)"
        return f"_safe_assert_desktop('{hint}', raise_on_fail=False)"
    if ctype == "variable_compare":
        var   = item.compare_var.replace("'", "\\'")
        value = item.compare_value.replace("'", "\\'")
        left  = f"str(_VARS.get('{var}', ''))"
        if item.compare_op == "neq":
            return f"{left} != '{value}'"
        if item.compare_op == "contains":
            return f"'{value}' in {left}"
        return f"{left} == '{value}'"   # "eq" varsayılan
    return "False"  # condition_type boş/bilinmeyen — güvenli varsayılan


def _wrap_text_step(idx: int, inner_lines: list, indent: str, platform: str, label: str,
                    comment: str = "") -> str:
    """Metin adımını STEP_START/DONE/FAIL ile sarar. Mobilde başarısız adımda
    _map_recover (uygulama haritasıyla kendini toparlama) çağrılır; ekran
    düzeltilebildiyse adım BİR KEZ daha denenir."""
    nl = chr(10)
    body = (nl + indent + "    ").join(inner_lines)
    head = f"{indent}_check_stop(){nl}{indent}print('STEP_START:{idx}'){nl}"
    if comment:
        head += f"{indent}# {comment}{nl}"
    if platform != "mobile":
        return head + (
            f"{indent}try:{nl}"
            f"{indent}    {body}{nl}"
            f"{indent}    print('STEP_DONE:{idx}'){nl}"
            f"{indent}except Exception as _e_{idx}:{nl}"
            f"{indent}    print(f'[FAIL] {{_e_{idx}}}'){nl}"
            f"{indent}    print('STEP_FAIL:{idx}'){nl}"
            f"{indent}    raise"
        )
    body2 = (nl + indent + "            ").join(inner_lines)
    return head + (
        f"{indent}try:{nl}"
        f"{indent}    {body}{nl}"
        f"{indent}    print('STEP_DONE:{idx}'){nl}"
        f"{indent}except Exception as _e_{idx}:{nl}"
        f"{indent}    print(f'[FAIL] {{_e_{idx}}}'){nl}"
        f"{indent}    _LAST_STEP_ERROR = str(_e_{idx}){nl}"
        f"{indent}    if _map_recover({idx}, {label!r}):{nl}"
        f"{indent}        try:{nl}"
        f"{indent}            {body2}{nl}"
        f"{indent}            print('STEP_DONE:{idx}'){nl}"
        f"{indent}        except Exception as _e2_{idx}:{nl}"
        f"{indent}            print(f'[FAIL] {{_e2_{idx}}}'){nl}"
        f"{indent}            print('STEP_FAIL:{idx}'){nl}"
        f"{indent}            raise{nl}"
        f"{indent}    else:{nl}"
        f"{indent}        print('STEP_FAIL:{idx}'){nl}"
        f"{indent}        raise"
    )


def _steps_to_code(items: list, platform: str, counter: list, rule_fn, indent: str = "", browser: str = "") -> tuple[str, list, list]:
    """items (üst seviye ya da bir then/else dalı) için Python kodu üretir.
    counter=[n] paylaşılan/mutable sayaç — dal içindeki adımlar da aynı artan idx'i kullanır.
    Returns: (indent'li kod string'i, code_lines rapor listesi, unrecognized metin listesi)
    code_lines elemanları: (kind, step_label, code_or_empty) — mevcut _generate_steps sözleşmesiyle aynı."""
    blocks = []
    code_lines = []
    unrecognized = []

    for item in items:
        idx = counter[0]; counter[0] += 1

        if item.disabled:
            label = item.content or f"Adım {idx+1}"
            code_lines.append(("comment", label, f"{indent}# [DEVRE DIŞI] {label}"))
            blocks.append(f"{indent}# [DEVRE DIŞI] {label}")
            continue

        if item.type == "condition":
            cond_expr = _condition_expr(item, platform, idx, browser)
            then_code, then_lines, then_unrec = _steps_to_code(
                item.then_steps, platform, counter, rule_fn, indent + "    ", browser)
            else_code, else_lines, else_unrec = _steps_to_code(
                item.else_steps, platform, counter, rule_fn, indent + "    ", browser)
            then_code = _ensure_block_body(then_code, indent)
            else_code = _ensure_block_body(else_code, indent)
            label = item.content or f"Koşul {idx+1}"
            block = (
                f"{indent}_check_stop()\n"
                f"{indent}print('STEP_START:{idx}')\n"
                f"{indent}_cond_{idx} = {cond_expr}\n"
                f"{indent}if _cond_{idx}:\n{then_code}\n"
                f"{indent}else:\n{else_code}\n"
                f"{indent}print('STEP_DONE:{idx}')"
            )
            blocks.append(block)
            code_lines.append(("condition", label, block))
            code_lines.extend(then_lines)
            code_lines.extend(else_lines)
            unrecognized.extend(then_unrec)
            unrecognized.extend(else_unrec)
            continue

        if item.type == "image":
            code = _indent_code("_check_stop()\n" + _image_step_code(item, platform, idx, browser), indent)
            blocks.append(code)
            code_lines.append(("image", item.content or f"Görsel {idx+1}", code))
        elif item.type == "sql":
            code = _indent_code("_check_stop()\n" + _sql_step_code(item, idx), indent)
            blocks.append(code)
            code_lines.append(("rule", item.content or f"SQL {idx+1}", code))
        elif item.type == "api":
            code = _indent_code("_check_stop()\n" + _api_step_code(item, idx), indent)
            blocks.append(code)
            code_lines.append(("rule", item.content or f"API {idx+1}", code))
        elif item.type == "comment":
            code = f"{indent}# [YORUM] {item.content}"
            blocks.append(code)
            code_lines.append(("comment", item.content or f"Yorum {idx+1}", code))
        else:
            step = item.content
            parsed = rule_fn(step)
            if parsed:
                wrapped = _wrap_text_step(idx, parsed, indent, platform, step)
                blocks.append(wrapped)
                code_lines.append(("rule", step, wrapped))
            else:
                blocks.append(None)   # LLM sonucu ikinci geçişte doldurulacak — yer tutucu
                code_lines.append(("llm", step, ""))
                unrecognized.append(step)

    return "\n\n".join(b for b in blocks if b is not None), code_lines, unrecognized


async def _generate_steps(req: GenerateRequest) -> tuple[str, list[dict]]:
    """
    Test case'i adım adım işler (metin + görsel + condition karma destekli).
    Döndürür: (kod_str, adım_raporu)
    """
    rule_fn = (
        _rule_mobile  if req.platform == "mobile"  else
        _rule_desktop if req.platform == "desktop" else
        _rule_web
    )
    lookup = _MOBILE_LOCATORS.get() if req.platform == "mobile" else None
    locator_hits: set = set()
    if lookup:
        def rule_fn(step: str, _base=rule_fn):
            lines, hit = _apply_locators_mobile(step, _base(step), lookup)
            if hit:
                locator_hits.add(step)
            return lines

    if req.steps_list:
        structured = req.steps_list
    else:
        structured = [
            StepItem(type="text", content=l.strip())
            for l in req.test_case.splitlines() if l.strip()
        ]

    counter = [0]
    steps_code, code_lines, unrecognized = _steps_to_code(structured, req.platform, counter, rule_fn, browser=req.browser or "")

    # LLM'e toplu sor (üst seviye + dal içi tüm tanınmayan metin adımları)
    llm_results = {}
    if unrecognized:
        llm_results = await _ask_llm_for_steps(unrecognized, req.platform)

    # LLM kodu üretilen ana string'e enjekte edilmeli — placeholder yaklaşımı yerine,
    # basitlik için LLM adımlarını üst seviyede yeniden üret (nested LLM adımı nadir durum,
    # 1-seviye sınırı ile dal içi LLM adımları da düz metinle aynı formatta üretilir).
    # report_out: bu geçişte üretilen HER adımın gerçek method/code'unu (rule/llm/todo dahil)
    # kaydeder — ilk geçişteki code_lines "llm" olarak damgaladığı ama burada asıl gerçek
    # sonucun (llm mi todo mu) belirlendiği adımlar için rapor artık bu listeden gelir.
    def _fill_llm(items: list, counter2: list, indent: str = "", report_out: list | None = None) -> str:
        if report_out is None:
            report_out = []
        blocks = []
        for item in items:
            idx = counter2[0]; counter2[0] += 1
            if item.disabled:
                label = item.content or f"Adım {idx+1}"
                block = f"{indent}# [DEVRE DIŞI] {label}"
                blocks.append(block)
                report_out.append({"step": label, "method": "comment", "code": block, "index": idx})
                continue
            if item.type == "condition":
                cond_expr = _condition_expr(item, req.platform, idx, req.browser or "")
                then_code = _ensure_block_body(_fill_llm(item.then_steps, counter2, indent + "    ", report_out), indent)
                else_code = _ensure_block_body(_fill_llm(item.else_steps, counter2, indent + "    ", report_out), indent)
                label = item.content or f"Koşul {idx+1}"
                block = (
                    f"{indent}_check_stop()\n"
                    f"{indent}print('STEP_START:{idx}')\n"
                    f"{indent}_cond_{idx} = {cond_expr}\n"
                    f"{indent}if _cond_{idx}:\n{then_code}\n"
                    f"{indent}else:\n{else_code}\n"
                    f"{indent}print('STEP_DONE:{idx}')"
                )
                blocks.append(block)
                report_out.append({"step": label, "method": "condition", "code": block, "index": idx})
                continue
            if item.type in ("image", "sql", "api", "comment"):
                # Bu tipler zaten _steps_to_code içinde tam üretildi (LLM'e hiç düşmezler),
                # burada sadece idx senkronizasyonu için tekrar üretiyoruz.
                if item.type == "image":
                    block  = _indent_code("_check_stop()\n" + _image_step_code(item, req.platform, idx, req.browser or ""), indent)
                    method = "image"
                    label  = item.content or f"Görsel {idx+1}"
                elif item.type == "sql":
                    block  = _indent_code("_check_stop()\n" + _sql_step_code(item, idx), indent)
                    method = "rule"
                    label  = item.content or f"SQL {idx+1}"
                elif item.type == "api":
                    block  = _indent_code("_check_stop()\n" + _api_step_code(item, idx), indent)
                    method = "rule"
                    label  = item.content or f"API {idx+1}"
                else:
                    block  = f"{indent}# [YORUM] {item.content}"
                    method = "comment"
                    label  = item.content or f"Yorum {idx+1}"
                blocks.append(block)
                report_out.append({"step": label, "method": method, "code": block, "index": idx})
                continue
            step = item.content
            parsed = rule_fn(step)
            if parsed:
                block = _wrap_text_step(idx, parsed, indent, req.platform, step)
                blocks.append(block)
                report_out.append({"step": step, "method": "rule", "code": block, "index": idx})
            else:
                llm_code = llm_results.get(step, "")
                llm_code = "\n".join(
                    l for l in llm_code.splitlines()
                    if l.strip().startswith("#") or re.match(r'^\s*[a-z_].*[\(\.]', l)
                ).strip()
                if llm_code:
                    llm_code = textwrap.dedent(llm_code)
                if not llm_code:
                    llm_code = f"# TODO (anlasilamadi): {step}\n{indent}time.sleep(1)"
                    _inner = llm_code
                    block = f"{indent}{_inner}"
                    blocks.append(block)
                    report_out.append({"step": step, "method": "todo", "code": block, "index": idx})
                else:
                    block = _wrap_text_step(idx, llm_code.splitlines(), indent, req.platform, step, comment=step)
                    blocks.append(block)
                    report_out.append({"step": step, "method": "llm", "code": block, "index": idx})
        return "\n\n".join(blocks)

    if unrecognized:
        # LLM sonucu varsa tüm kodu (idx sayacını sıfırlayıp) yeniden, bu sefer LLM
        # sonuçlarıyla üret — iki geçişli tasarım mevcut _generate_steps ile aynı prensip.
        # fill_report bu geçişte her adımın GERÇEK method/code'unu (rule/llm/todo dahil) taşır.
        fill_report: list = []
        steps_code = _fill_llm(structured, [0], report_out=fill_report)
        report = sorted(fill_report, key=lambda r: r["index"])
    else:
        # LLM'e hiç düşülmedi — ilk geçişteki code_lines zaten doğru rapor.
        report = [
            {"step": step, "method": kind, "code": code, "index": i}
            for i, (kind, step, code) in enumerate(code_lines)
        ]
    for r in report:
        if r.get("method") == "rule" and r.get("step") in locator_hits:
            r["locator"] = True

    return steps_code, report


async def _ask_llm_for_steps(steps: list[str], platform: str) -> dict[str, str]:
    """LLM'den tanınmayan adımlar için kod üret."""
    if platform == "mobile":
        api_note = (
            "Use ONLY these helpers — NEVER call .send_keys() directly:\n"
            "  _safe_click(driver, 'text')  — click by visible text\n"
            "  _safe_input(driver, '//android.widget.EditText[1]', 'value')  — type text\n"
            "  _safe_assert(driver, 'text')  — verify element visible"
        )
    elif platform == "desktop":
        api_note = (
            "Use ONLY these helpers (pyautogui-based desktop automation):\n"
            "  _safe_click_desktop('text', _WINDOW_TITLE)  — click element by visible text (pywinauto → OCR → vision)\n"
            "  _safe_type_desktop('value')                 — type text into active field\n"
            "  _safe_assert_desktop('text')                — verify text/element on screen\n"
            "  _pag.hotkey('ctrl', 's')                    — keyboard shortcut\n"
            "  _pag.press('enter')                         — press a key\n"
            "  _pag.scroll(-3)                             — scroll down\n"
            "  time.sleep(1)                               — wait"
        )
    else:
        api_note = (
            "Use ONLY these helpers:\n"
            "  _safe_click_web(driver, 'text')                        — click by visible text\n"
            "  _safe_input_web(driver, 'field_hint', 'value')         — type into input (hint = label/placeholder keyword)\n"
            "  _safe_input_web(driver, 'field_hint', 'value', is_password=True)  — for password fields\n"
            "  _safe_assert_web(driver, 'text')                       — verify text is visible\n"
            "  driver.get('url')                                       — navigate to URL\n"
            "  driver.execute_script('window.scrollBy(0,400);')        — scroll"
        )

    numbered = "\n".join(f"{i+1}. {s}" for i, s in enumerate(steps))
    auto_fill_hint = (
        "Auto-fill when no value specified: tutar/amount→'10000', telefon→'05551234567', "
        "şifre/password→'123456', tc/kimlik→'12345678901', generic→'121212'.\n"
        "For 'seçim yapılır' (dropdown): click first list item via find_elements(...)[0].click().\n"
        "For compound steps (e.g. 'seçim yapılır, tutar girilir'): handle each action separately."
    )
    prompt = f"""Write Python automation code for each step. Output only code + # comments. No explanations.
{api_note}
{auto_fill_hint}

Steps:
{numbered}

Code (one block per step, separated by blank lines):"""

    try:
        async with httpx.AsyncClient(timeout=60.0) as client:
            resp = await client.post(_ollama_url(), json={
                "model": _ollama_model(),
                "prompt": prompt,
                "stream": False,
                "options": {"temperature": 0.1, "num_predict": 512}
            })
            raw = resp.json().get("response", "")
    except Exception:
        raw = ""

    # Basit eşleştirme: bloğu adım sırasıyla eşleştir
    result = {}
    blocks = [b.strip() for b in re.split(r'\n\s*\n', raw) if b.strip()]
    for i, step in enumerate(steps):
        if i < len(blocks):
            # Prose filtrele — sadece kod satırları al
            code_only = "\n".join(
                l for l in blocks[i].splitlines()
                if l.startswith("#") or re.match(r'^[a-z_\s].*[\(\.]', l.strip())
            )
            result[step] = code_only if code_only else f"# TODO: {step}\ntime.sleep(1)"
        else:
            result[step] = f"# TODO: {step}\ntime.sleep(1)"
    return result


# ══════════════════════════════════════════════════════════════════════════════
# ENDPOINT'LER
# ══════════════════════════════════════════════════════════════════════════════

@router.post("/parse-file")
async def parse_file_upload(
    file: UploadFile = File(...),
    sheet_name: Optional[str] = Form(None),
    current_user: User = Depends(get_current_user)
):
    """
    Excel (.xlsx/.xls), CSV veya TXT dosyasını okur,
    test adımlarını düz metin olarak döndürür.
    sheet_name verilirse ve workbook'ta varsa o sheet okunur, yoksa aktif sheet
    (wb.active — önceki davranışla birebir uyumlu, sheet_name olmadan değişmez).
    """
    contents = await file.read()
    filename  = (file.filename or "").lower()

    try:
        if filename.endswith((".xlsx", ".xls")):
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            lines = []
            for row in ws.iter_rows(values_only=True):
                parts = [str(c).strip() for c in row if c is not None and str(c).strip() not in ("", "None")]
                if parts:
                    lines.append("  ".join(parts))
            wb.close()
            text = "\n".join(lines)

        elif filename.endswith(".csv"):
            import io, csv
            decoded = contents.decode("utf-8-sig", errors="replace")
            # Azure DevOps formatını önce dene
            ado_text = _parse_azure_devops_csv(decoded)
            if ado_text:
                text = ado_text
            else:
                reader = csv.reader(io.StringIO(decoded))
                lines  = ["  ".join(r) for r in reader if any(c.strip() for c in r)]
                text   = "\n".join(lines)

        else:
            # .txt / .md veya uzantısız
            for enc in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
                try:
                    text = contents.decode(enc)
                    break
                except Exception:
                    continue
            else:
                text = contents.decode("utf-8", errors="replace")

        # Null byte temizle
        text = text.replace("\x00", "").strip()

        return {"success": True, "text": text, "lines": len(text.splitlines())}

    except Exception as e:
        raise HTTPException(400, f"Dosya okunamadı: {str(e)}")


# ── ADO Excel Parser Helpers ──────────────────────────────────────────────────

_ADO_CMD_WORDS = [
    "görülür", "görüntülenir", "görünür", "gör",
    "tıklanır", "tıkla", "açılır", "kapanır",
    "girilir", "seçilir", "gelir", "kapat",
    "kaydedilir", "karşılaştırılır", "doğrulanır", "beklenir",
]

def _ado_cmd_word(text: str) -> str:
    stripped = text.strip().rstrip(".")
    lower = stripped.lower()
    for cw in _ADO_CMD_WORDS:
        if lower.endswith(cw):
            return cw
    return ""

def _ado_parse_expected(expected: str) -> list:
    text = (expected or "").strip()
    if not text:
        return []

    cmd = _ado_cmd_word(text)
    lines = [l.strip() for l in text.split("\n") if l.strip()]

    def _normalize_line(line: str) -> str:
        # Satır sonu ayraçları temizle: virgül, noktalı virgül, "ve"
        line = re.sub(r"[,;]\s*$", "", line).strip()
        line = re.sub(r"\s+ve\s*$", "", line, flags=re.IGNORECASE).strip()
        if cmd:
            # "geldiği {cmd}" veya sadece "{cmd}" varsa sona sıfırla
            line = re.sub(r"\s+geldi[gğ]i\s+" + re.escape(cmd) + r"\s*\.?\s*$", "", line, flags=re.IGNORECASE).strip()
            line = re.sub(r"\s+" + re.escape(cmd) + r"\s*\.?\s*$", "", line, flags=re.IGNORECASE).strip()
        return line

    # Tek satır → tek adım
    if len(lines) <= 1:
        single = _normalize_line(text.strip())
        return [single + (f" {cmd}." if cmd else ("." if not single.endswith(".") else ""))]

    # Çok satır → her satır ayrı adım
    result = []
    for line in lines:
        line = _normalize_line(line)
        if line:
            result.append(line + (f" {cmd}." if cmd else ("." if not line.endswith(".") else "")))
    return result


@router.post("/excel-sheets")
async def list_excel_sheets(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    """
    Excel dosyasındaki sheet adlarını döner — kullanıcı içe aktarmadan önce
    hangi sheet'i kullanacağını seçebilsin diye (workbook aktif olmayan diğer
    sheet'ler daha önce parse-file/parse-excel-ado tarafından sessizce
    atlanıyordu; wb.active sadece son kaydedilen aktif sekmeyi okuyordu).
    Döndürür: { sheets: [str], active: str }
    """
    contents = await file.read()
    fname = (file.filename or "").lower()
    if not (fname.endswith(".xlsx") or fname.endswith(".xls")):
        raise HTTPException(400, "Sadece .xlsx / .xls dosyaları desteklenir.")

    try:
        import io as _io
        import openpyxl as _xl

        wb = _xl.load_workbook(_io.BytesIO(contents), read_only=True, data_only=True)
        sheets = list(wb.sheetnames)
        active = wb.active.title if wb.active is not None else (sheets[0] if sheets else "")
        wb.close()
        return {"sheets": sheets, "active": active}
    except Exception as e:
        raise HTTPException(400, f"Excel okunamadı: {str(e)}")


@router.post("/parse-excel-ado")
async def parse_excel_ado(
    file: UploadFile = File(...),
    sheet_name: Optional[str] = Form(None),
    current_user: User = Depends(get_current_user),
):
    """
    Azure DevOps Excel export (.xlsx) parse eder.
    sheet_name verilirse ve workbook'ta varsa o sheet okunur, yoksa aktif sheet.
    Döndürür: { test_cases: [{ title, steps: [str] }] }
    """
    contents = await file.read()
    fname = (file.filename or "").lower()
    if not (fname.endswith(".xlsx") or fname.endswith(".xls")):
        raise HTTPException(400, "Sadece .xlsx / .xls dosyaları desteklenir.")

    try:
        import io as _io
        import openpyxl as _xl

        wb = _xl.load_workbook(_io.BytesIO(contents), data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))

        # Başlık satırını bul
        col_map = {}
        header_row_idx = None
        for ri, row in enumerate(rows):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if "Work Item Type" in cells and "Step Action" in cells:
                header_row_idx = ri
                col_map = {v: i for i, v in enumerate(cells) if v}
                break

        if header_row_idx is None:
            raise HTTPException(400, "ADO başlık satırı bulunamadı.")

        ci_type     = col_map.get("Work Item Type", 2)
        ci_title    = col_map.get("Title", 3)
        ci_step     = col_map.get("Test Step", 4)
        ci_action   = col_map.get("Step Action", 5)
        ci_expected = col_map.get("Step Expected", 6)

        def _cell(row, i):
            return str(row[i]).strip() if len(row) > i and row[i] is not None else ""

        test_cases = []
        current_tc = None

        for row in rows[header_row_idx + 1:]:
            wtype  = _cell(row, ci_type)
            title  = _cell(row, ci_title)
            action = _cell(row, ci_action)
            exp    = _cell(row, ci_expected)
            has_step = len(row) > ci_step and row[ci_step] is not None

            if wtype == "Test Case" and title:
                current_tc = {"title": title, "steps": []}
                test_cases.append(current_tc)
            elif has_step and current_tc is not None:
                if action:
                    a = action.rstrip(".")
                    current_tc["steps"].append(a + ".")
                if exp:
                    current_tc["steps"].extend(_ado_parse_expected(exp))

        wb.close()

        if not test_cases:
            raise HTTPException(400, "Excel'de test case bulunamadı.")

        return {"test_cases": test_cases}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Excel parse hatası: {str(e)}")


def _normalize_ado_cell(text: str) -> str:
    """
    Azure DevOps export'unda tırnaklar bazen ikileşmiş çıkar (standart CSV kaçışı
    ya da Excel apostrof davranışı): "''X''" / 'X'' gibi. Fazladan tekrar eden
    tırnak karakterlerini tekile indirger.
    """
    if not text:
        return text
    t = text.replace("\x00", "")
    t = re.sub(r"'{2,}", "'", t)
    t = re.sub(r'"{2,}', '"', t)
    return t.strip()


@router.post("/parse-csv-ado")
async def parse_csv_ado(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    """
    Azure DevOps CSV export parse eder (parse_excel_ado'nun CSV eşleniği).
    Bir CSV dosyası birden fazla Test Case içerebilir — her biri title+steps
    olarak gruplanır.
    Döndürür: { test_cases: [{ title, steps: [str] }] }
    """
    contents = await file.read()
    fname = (file.filename or "").lower()
    if not fname.endswith(".csv"):
        raise HTTPException(400, "Sadece .csv dosyaları desteklenir.")

    try:
        import io as _io
        import csv as _csv

        decoded = None
        for enc in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
            try:
                decoded = contents.decode(enc)
                break
            except Exception:
                continue
        if decoded is None:
            decoded = contents.decode("utf-8", errors="replace")

        rows = list(_csv.reader(_io.StringIO(decoded)))
        if not rows:
            raise HTTPException(400, "CSV dosyası boş.")

        header_row_idx = None
        col_map = {}
        for ri, row in enumerate(rows):
            cells = [c.strip() if c is not None else "" for c in row]
            if "Work Item Type" in cells and "Step Action" in cells:
                header_row_idx = ri
                col_map = {v: i for i, v in enumerate(cells) if v}
                break

        if header_row_idx is None:
            raise HTTPException(400, "ADO başlık satırı bulunamadı.")

        ci_type     = col_map.get("Work Item Type", 1)
        ci_title    = col_map.get("Title", 2)
        ci_step     = col_map.get("Test Step", 3)
        ci_action   = col_map.get("Step Action", 4)
        ci_expected = col_map.get("Step Expected", 5)

        def _cell(row, i):
            return _normalize_ado_cell(row[i]) if len(row) > i and row[i] is not None else ""

        test_cases = []
        current_tc = None

        for row in rows[header_row_idx + 1:]:
            wtype  = _cell(row, ci_type)
            title  = _cell(row, ci_title)
            action = _cell(row, ci_action)
            exp    = _cell(row, ci_expected)
            has_step = len(row) > ci_step and str(row[ci_step]).strip() != ""

            if wtype == "Test Case" and title:
                current_tc = {"title": title, "steps": []}
                test_cases.append(current_tc)
            elif has_step and current_tc is not None:
                if action:
                    a = action.rstrip(".")
                    current_tc["steps"].append(a + ".")
                if exp:
                    current_tc["steps"].extend(_ado_parse_expected(exp))

        if not test_cases:
            raise HTTPException(400, "CSV'de test case bulunamadı.")

        return {"test_cases": test_cases}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"CSV parse hatası: {str(e)}")


@router.get("/status")
async def ollama_status():
    cfg = _get_ollama_cfg()
    tags_url = cfg["url"].replace("/api/generate", "/api/tags")
    configured_model_base = cfg["model"].split(":")[0]
    try:
        async with httpx.AsyncClient(timeout=3.0) as client:
            resp = await client.get(tags_url)
            if resp.status_code == 200:
                models = [m["name"] for m in resp.json().get("models", [])]
                model_available = any(m.split(":")[0] == configured_model_base for m in models)
                return {"online": True, "model_available": model_available, "models": models, "configured_model": cfg["model"]}
    except Exception:
        pass
    return {"online": False, "model_available": False, "models": [], "configured_model": cfg["model"]}


def _find_adb() -> str:
    """PATH'e bağımlı kalmadan adb.exe'yi bulur — sunucu süreci PATH güncellenmeden önce
    başlamış olsa bile projenin kendi platform-tools'unu ve yaygın SDK konumlarını dener."""
    import shutil as _shutil
    found = _shutil.which("adb")
    if found:
        return found
    backend_dir = Path(__file__).resolve().parent.parent
    candidates = [
        backend_dir / "platform-tools" / "adb.exe",
        Path(os.path.expanduser("~")) / "AppData" / "Local" / "Android" / "Sdk" / "platform-tools" / "adb.exe",
        Path("C:/platform-tools/adb.exe"),
        Path("C:/Android/platform-tools/adb.exe"),
    ]
    for c in candidates:
        if c.exists():
            return str(c)
    return "adb"  # bulunamadı — subprocess FileNotFoundError fırlatacak, aşağıda yakalanıyor


@router.get("/adb/devices")
async def adb_devices(current_user: User = Depends(get_current_user)):
    # Tüm cihaz sunucuları (yerel + uzak); "devices" listesi eski biçimde (seri no) kalır
    found, _ = await asyncio.to_thread(dnodes.discover)
    ready = [d for d in found if d["adb_state"] == "device"]
    busy = {d: {"holder": h["display"], "mine": h["user_id"] == current_user.id,
                "purposes": leases.purposes_text(h)}
            for d, h in leases.all_holders().items()}
    return {"devices": [d["id"] for d in ready], "busy": busy,
            "info": {d["id"]: {"model": d["model"], "node": d["node_name"]} for d in ready}}


@router.get("/adb/activity")
async def adb_get_activity(device_id: str, app_package: str, current_user: User = Depends(get_current_user)):
    try:
        cmd = await asyncio.to_thread(dnodes.adb_cmd, device_id, "shell", "cmd", "package",
                                      "resolve-activity", "--brief", app_package)
        result = await asyncio.to_thread(subprocess.run, cmd, capture_output=True, timeout=10)
        output = result.stdout.decode("utf-8", errors="replace").strip()
        for line in output.splitlines():
            if "/" in line and app_package in line:
                activity = line.split("/", 1)[1].strip()
                if activity.startswith("."):
                    activity = app_package + activity
                return {"found": True, "activity": activity}
        return {"found": False, "activity": ""}
    except FileNotFoundError:
        raise HTTPException(500, "ADB bulunamadı")


@router.post("/generate")
async def generate_script(req: GenerateRequest, current_user: User = Depends(get_current_user),
                          db: Session = Depends(get_db)):
    """
    Hibrit dönüşüm:
    - Kural motoru → standart adımlar
    - LLM          → tanınmayanlar
    - Vision helper → runtime fallback (script içine gömülü)
    """
    tok = _MOBILE_LOCATORS.set(_locator_lookup(current_user.id, db, req.locator_file)
                               if req.platform == "mobile" else None)
    try:
        steps_code, report = await _generate_steps(req)
    finally:
        _MOBILE_LOCATORS.reset(tok)

    vars_seed = _get_step_vars(current_user.id) if req.single_step_test else None

    if req.platform == "mobile":
        final_script = _build_mobile_template(req, steps_code)
    elif req.platform == "web":
        final_script = _build_web_template(req, steps_code)
    else:
        final_script = _build_desktop_template(req, steps_code, vars_seed=vars_seed)

    todo_steps = [r for r in report if r["method"] == "todo"]
    warnings   = [f"Anlaşılamayan adım: {r['step']}" for r in todo_steps]

    return {
        "success":      True,
        "script":       final_script,
        "platform":     req.platform,
        "report":       report,
        "warnings":     warnings,
        "script_valid": True
    }


def _run_log_dir(user_id: int, db) -> Path:
    """Kullanıcının run log klasörünü döndürür (workspace > local fallback)."""
    try:
        ws = _get_user_workspace(user_id, db)
        if ws:
            d = Path(ws.local_path) / "ai_run_logs"
            d.mkdir(parents=True, exist_ok=True)
            return d
    except Exception:
        pass
    d = Path(__file__).parent.parent / "local_run_logs" / str(user_id)
    d.mkdir(parents=True, exist_ok=True)
    return d


@router.post("/execute")
async def execute_script(
    req: ExecuteRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if not req.script or len(req.script.strip()) < 10:
        raise HTTPException(400, "Geçersiz script")

    is_desktop = req.platform == "desktop"
    if is_desktop:
        lock = _get_desktop_lock()
        if lock.locked():
            raise HTTPException(423, "Şu an başka bir desktop testi çalışıyor. Lütfen bekleyin.")

    run_id     = str(uuid.uuid4())[:8]

    # Mobil koşum telefonu koşum boyunca kiralar (başkasının oturumu/koşumu varken reddedilir).
    # Telefon, üretilen script'teki udid'den okunur; yoksa tek bağlı telefon.
    run_device, run_lease = "", None
    if req.platform == "mobile":
        _um = re.search(r'options\.udid\s*=\s*["\']([^"\']+)["\']', req.script)
        run_device = _um.group(1) if _um else await asyncio.to_thread(_resolve_device, "")
        try:
            run_lease = leases.acquire(run_device, current_user, "run",
                                       req.scenario_name or "Adsız Senaryo", ref=run_id)
        except leases.DeviceBusyError as e:
            raise _busy_http(e)

    def _release_run_device():
        leases.release(run_device, run_lease)   # iki kez çağrılması zararsız

    log_dir    = _run_log_dir(current_user.id, db)
    started_at = _datetime.datetime.utcnow().isoformat()
    step_outcomes: dict[int, str] = {}   # idx -> "pass" | "fail"
    terminal_lines: list[str] = []        # tam terminal çıktısı
    _stop_fd, stop_file = tempfile.mkstemp(suffix=".stop", prefix=f"ai_stop_{run_id}_")
    os.close(_stop_fd)
    os.unlink(stop_file)  # Dosya yokluğu = devam; varlığı = dur sinyali
    _stop_files[run_id] = stop_file

    # Uygulama haritası: mobil koşumda ekran izi tutulur, koşum bitince arka planda
    # haritaya işlenir (ai_config.json → "app_map_auto": false ile kapatılabilir)
    map_trace_dir = ""
    map_recover = "1"   # kendini toparlama ("app_map_recovery": false ile kapatılır)
    if req.platform == "mobile":
        try:
            from services.ai_service import get_ai_config as _gac
            _mcfg = _gac()
            if _mcfg.get("app_map_auto", True):
                map_trace_dir = tempfile.mkdtemp(prefix=f"ai_maptrace_{run_id}_")
            map_recover = "1" if _mcfg.get("app_map_recovery", True) else "0"
        except Exception:
            map_trace_dir = ""

    async def stream():
        nonlocal step_outcomes
        exit_code = -1
        ended_at  = None

        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".py", delete=False, encoding="utf-8", prefix="ai_test_"
        ) as f:
            f.write(req.script)
            tmp = f.name

        yield _sse("run_id", run_id)
        yield _sse("info", f"Script: {os.path.basename(tmp)}")
        yield _sse("info", f"Platform: {req.platform}")
        yield _sse("info", "Kosum baslatiliyor...\n")

        desktop_lock_held = False
        if is_desktop:
            await _get_desktop_lock().acquire()
            desktop_lock_held = True

        proc = None
        try:
            _popen_kw: dict = dict(
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                env={**os.environ, "PYTHONUNBUFFERED": "1", "PYTHONIOENCODING": "utf-8",
                     "_AI_STOP_FILE": stop_file, "_AI_MAP_TRACE": map_trace_dir,
                     "_AI_MAP_RECOVER": map_recover},
            )
            if sys.platform == "win32":
                _popen_kw["creationflags"] = subprocess.CREATE_NEW_PROCESS_GROUP
            proc = subprocess.Popen([sys.executable, tmp], **_popen_kw)
            _running_processes[run_id] = proc
            loop = asyncio.get_event_loop()
            while True:
                line = await loop.run_in_executor(None, proc.stdout.readline)
                if not line:
                    break
                text = line.decode("utf-8", errors="replace").rstrip()
                if text.startswith("STEP_START:"):
                    idx_s = text.split(":", 1)[1].strip()
                    n_label = req.step_labels[int(idx_s)] if idx_s.isdigit() and int(idx_s) < len(req.step_labels) else f"Adım {idx_s}"
                    terminal_lines.append(f"▶ [{idx_s}] {n_label}")
                    yield _sse("step_start", idx_s)
                elif text.startswith("STEP_DONE:"):
                    idx_s = text.split(":", 1)[1].strip()
                    try: step_outcomes[int(idx_s)] = "pass"
                    except ValueError: pass
                    terminal_lines.append(f"✓ [{idx_s}] PASS")
                    yield _sse("step_done", idx_s)
                elif text.startswith("STEP_FAIL:"):
                    idx_s = text.split(":", 1)[1].strip()
                    try: step_outcomes[int(idx_s)] = "fail"
                    except ValueError: pass
                    terminal_lines.append(f"✗ [{idx_s}] FAIL")
                    yield _sse("step_fail", idx_s)
                else:
                    terminal_lines.append(text)
                    if "→_VARS]" in text:
                        _ingest_step_vars(current_user.id, text)
                    yield _sse("log", text)

            await loop.run_in_executor(None, proc.wait)
            exit_code = proc.returncode
            ended_at  = _datetime.datetime.utcnow().isoformat()
            if exit_code == 0:
                yield _sse("success", f"\nKosum tamamlandi. (exit: {exit_code})")
            else:
                yield _sse("error", f"\nKosum hata ile bitti. (exit: {exit_code})")
        except Exception as e:
            yield _sse("error", str(e))
        finally:
            _running_processes.pop(run_id, None)
            if proc and proc.poll() is None:
                proc.terminate()
            try:
                os.unlink(tmp)
            except Exception:
                pass

            # Stop dosyasını temizle
            _stop_files.pop(run_id, None)
            try:
                if os.path.exists(stop_file):
                    os.unlink(stop_file)
            except Exception:
                pass

            # ── Run log kaydet ───────────────────────────────────────────────
            try:
                n = len(req.step_labels)
                log_entry = {
                    "run_id":          run_id,
                    "scenario_name":   req.scenario_name,
                    "scenario_folder": req.scenario_folder,
                    "platform":        req.platform,
                    "started_at":      started_at,
                    "ended_at":        ended_at or _datetime.datetime.utcnow().isoformat(),
                    "exit_code":       exit_code,
                    "total_steps":     n,
                    "passed":          sum(1 for v in step_outcomes.values() if v == "pass"),
                    "failed":          sum(1 for v in step_outcomes.values() if v == "fail"),
                    "steps": [
                        {
                            "index":  i,
                            "label":  req.step_labels[i] if i < n else f"Adım {i+1}",
                            "status": step_outcomes.get(i, "unknown"),
                        }
                        for i in range(n)
                    ],
                    "terminal_log": "\n".join(terminal_lines),
                }
                (log_dir / f"{run_id}.json").write_text(
                    _json.dumps(log_entry, ensure_ascii=False, indent=2), encoding="utf-8"
                )
            except Exception:
                pass

            # ── Uygulama haritası: izi arka planda işle (OCR koşumu yavaşlatmasın) ──
            if map_trace_dir:
                import threading as _thr
                from services.app_map_service import process_run_trace as _prt
                _labels, _outs, _user = list(req.step_labels), dict(step_outcomes), current_user.username

                def _map_job():
                    try:
                        _sm = _prt(map_trace_dir, _labels, _outs, _user)
                        print(f"[APP-MAP] run {run_id}: {_sm}")
                    except Exception as _me:
                        print(f"[APP-MAP] run {run_id} işlenemedi: {_me}")
                _thr.Thread(target=_map_job, daemon=True).start()

            # Yalnizca BU istek gercekten kilidi aldiysa serbest birakilir (bkz.
            # ai_job_service._run_job_cells'teki ayni duzeltme / final review notu).
            if desktop_lock_held:
                try:
                    _get_desktop_lock().release()
                except RuntimeError:
                    pass
            _release_run_device()

            yield _sse("done", "")

    # background: istemci akışı hiç okumadan koparsa (generator başlamazsa) da kira bırakılır
    from starlette.background import BackgroundTask
    return StreamingResponse(stream(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
                             background=BackgroundTask(_release_run_device))


class SessionStartRequest(BaseModel):
    platform:     str
    browser:      Optional[str] = "chrome"
    target_url:   Optional[str] = ""
    device_id:    Optional[str] = ""
    app_package:  Optional[str] = ""
    app_activity: Optional[str] = ""
    appium_url:   Optional[str] = "http://localhost:4723"


@router.post("/session/start")
async def session_start(req: SessionStartRequest, current_user: User = Depends(get_current_user)):
    """Web/Mobil için görünür, kalıcı bir tarayıcı/uygulama oturumu açar (driver.quit() etmez)."""
    if req.platform not in ("web", "mobile"):
        raise HTTPException(400, "Oturum sadece web/mobile platformlarında desteklenir.")
    if _get_session(current_user.id, req.platform):
        raise HTTPException(409, "Bu platform için zaten aktif bir oturum var.")

    device_id, lease_token = req.device_id or "", None
    if req.platform == "mobile":
        device_id = await asyncio.to_thread(_resolve_device, req.device_id)
        try:
            lease_token = leases.acquire(device_id, current_user, "session", req.app_package or "")
        except leases.DeviceBusyError as e:
            raise _busy_http(e)
    try:
        return await _session_start(req, current_user, device_id, lease_token)
    except BaseException:
        leases.release(device_id, lease_token)   # oturum açılamadı → telefon boşa düşer
        raise


async def _session_start(req: "SessionStartRequest", current_user: User, device_id: str,
                         lease_token: Optional[str]):
    _reset_step_vars(current_user.id)

    gen_req = GenerateRequest(
        platform=req.platform, browser=req.browser, target_url=req.target_url,
        device_id=device_id, app_package=req.app_package, app_activity=req.app_activity,
        appium_url=req.appium_url,
    )
    mjpeg_port = 0
    if req.platform == "mobile":
        # Canlı MJPEG akışı yalnız yerel telefonda: Appium akış portunu telefonun takılı olduğu
        # makinenin localhost'una yönlendirir, uzak telefonda VM'den erişilemez → adb ekran görüntüsü
        if not dnodes.is_remote(device_id):
            import socket as _sock
            with _sock.socket() as _s:          # boş yerel port (birden fazla telefon/oturum çakışmasın)
                _s.bind(("127.0.0.1", 0))
                mjpeg_port = _s.getsockname()[1]
        script = _build_mobile_template(gen_req, "", session_mode="start", mjpeg_port=mjpeg_port)
    else:
        script = _build_web_template(gen_req, "", session_mode="start")

    with tempfile.NamedTemporaryFile(mode="w", suffix=".py", delete=False, encoding="utf-8", prefix="ai_session_start_") as f:
        f.write(script)
        tmp = f.name

    # stdout/stderr bir DOSYAYA yönlendirilir (PIPE değil): script os._exit(0) ile bitse bile
    # chromedriver/appium'un henüz yaşayan çocuk süreçleri PIPE'ın yazma ucunu miras alıp EOF'u
    # engelleyebilir - subprocess.run(capture_output=True) bu yüzden timeout'a kadar asılı kalırdı.
    out_fd, out_path = tempfile.mkstemp(suffix=".log", prefix="ai_session_start_out_")
    os.close(out_fd)

    try:
        with open(out_path, "w", encoding="utf-8") as out_f:
            proc = subprocess.Popen(
                [sys.executable, tmp], stdout=out_f, stderr=out_f,
                env={**os.environ, "PYTHONUNBUFFERED": "1", "PYTHONIOENCODING": "utf-8"},
            )
            # iPhone: WebDriverAgent ilk açılışta derlenip yüklenir (birkaç dakika sürebilir)
            is_ios = req.platform == "mobile" and dnodes.platform_for(device_id) == "ios"
            try:
                await asyncio.to_thread(proc.wait, 300 if is_ios else 60)
            except subprocess.TimeoutExpired:
                proc.kill()
                raise HTTPException(500, "Oturum başlatma zaman aşımına uğradı."
                                    + (" (iPhone: WebDriverAgent imzalama/kurulumunu kontrol edin)" if is_ios else ""))

        out = Path(out_path).read_text(encoding="utf-8", errors="replace")
        sid_m = re.search(r"SESSION_ID:\s*(\S+)", out)
        url_m = re.search(r"EXECUTOR_URL:\s*(\S+)", out)
        if not sid_m or not url_m:
            raise HTTPException(500, f"Oturum başlatılamadı: {out}")
        _active_sessions[(current_user.id, req.platform)] = {
            "session_id": sid_m.group(1),
            "command_executor_url": url_m.group(1),
            "app_package": req.app_package or "",   # uygulama haritası kayıt modu için
            "device_id": device_id,
            "device_os": "ios" if is_ios else ("android" if req.platform == "mobile" else ""),
            "lease_token": lease_token,             # _remove_session telefon kirasını bırakır
            "mjpeg_port": mjpeg_port,
            "started_at": _time_mod.time(),
            "last_used_at": _time_mod.time(),
        }
        return {"success": True, "platform": req.platform}
    finally:
        try:
            os.unlink(tmp)
        except Exception:
            pass
        try:
            os.unlink(out_path)
        except Exception:
            pass


@router.post("/vars/reset")
async def reset_step_vars(current_user: User = Depends(get_current_user)):
    """Tek-adım test havuzundaki {{değişken}}leri temizler (Yeni Senaryo / Sıfırla)."""
    _reset_step_vars(current_user.id)
    return {"success": True}


class SessionRunStepRequest(BaseModel):
    platform:    str
    step:        StepItem
    step_label:  Optional[str] = "Adım"
    locator_file: Optional[str] = ""


@router.post("/session/run-step")
async def session_run_step(req: SessionRunStepRequest, current_user: User = Depends(get_current_user),
                           db: Session = Depends(get_db)):
    """Aktif oturuma bağlanıp TEK bir adımı, önceki adımların bıraktığı durumdan devam ederek çalıştırır."""
    sess = _get_session(current_user.id, req.platform)
    if not sess:
        raise HTTPException(404, "Aktif oturum yok. Önce oturumu başlatın.")

    gen_req = GenerateRequest(platform=req.platform, steps_list=[req.step],
                              device_id=sess.get("device_id", ""))   # uzak telefonda adb ortamı için
    tok = _MOBILE_LOCATORS.set(_locator_lookup(current_user.id, db, req.locator_file)
                               if req.platform == "mobile" else None)
    try:
        steps_code, report = await _generate_steps(gen_req)
    finally:
        _MOBILE_LOCATORS.reset(tok)
    build_fn = _build_mobile_template if req.platform == "mobile" else _build_web_template
    script = build_fn(
        gen_req, steps_code, session_mode="attach",
        session_id=sess["session_id"], command_executor_url=sess["command_executor_url"],
        vars_seed=_get_step_vars(current_user.id),
    )

    with tempfile.NamedTemporaryFile(mode="w", suffix=".py", delete=False, encoding="utf-8", prefix="ai_session_step_") as f:
        f.write(script)
        tmp = f.name

    try:
        proc = subprocess.run(
            [sys.executable, tmp], capture_output=True, text=True, encoding="utf-8", timeout=60,
            env={**os.environ, "PYTHONUNBUFFERED": "1", "PYTHONIOENCODING": "utf-8"},
        )
        _touch_session(current_user.id, req.platform)
        _ingest_step_vars(current_user.id, (proc.stdout or "") + (proc.stderr or ""))
        passed = proc.returncode == 0
        return {
            "success": passed,
            "output": (proc.stdout or "") + (proc.stderr or ""),
            "report": report,
        }
    except subprocess.TimeoutExpired:
        _touch_session(current_user.id, req.platform)
        return {"success": False, "output": "Adım zaman aşımına uğradı.", "report": report}
    finally:
        try:
            os.unlink(tmp)
        except Exception:
            pass


@router.post("/session/end")
async def session_end(platform: str, current_user: User = Depends(get_current_user)):
    """Aktif oturuma son kez bağlanıp kapatır (driver.quit())."""
    sess = _get_session(current_user.id, platform)
    if not sess:
        raise HTTPException(404, "Aktif oturum yok.")

    gen_req = GenerateRequest(platform=platform, device_id=sess.get("device_id", ""))
    build_fn = _build_mobile_template if platform == "mobile" else _build_web_template
    script = build_fn(
        gen_req, "", session_mode="end",
        session_id=sess["session_id"], command_executor_url=sess["command_executor_url"],
    )
    with tempfile.NamedTemporaryFile(mode="w", suffix=".py", delete=False, encoding="utf-8", prefix="ai_session_end_") as f:
        f.write(script)
        tmp = f.name
    try:
        subprocess.run(
            [sys.executable, tmp], capture_output=True, text=True, encoding="utf-8", timeout=30,
            env={**os.environ, "PYTHONUNBUFFERED": "1", "PYTHONIOENCODING": "utf-8"},
        )
    except subprocess.TimeoutExpired:
        pass
    finally:
        if platform == "web":
            # attach modundaki driver'in Service nesnesi yok - driver.quit() chromedriver.exe'yi
            # kapatmaz, /shutdown endpoint'ine dogrudan istek atmak gerekir (orphan surec onlenir)
            try:
                async with httpx.AsyncClient(timeout=5) as client:
                    await client.get(sess["command_executor_url"].rstrip("/") + "/shutdown")
            except Exception:
                pass
        try:
            os.unlink(tmp)
        except Exception:
            pass
        _remove_session(current_user.id, platform)
    return {"success": True}


@router.get("/session/status")
async def session_status(platform: str, current_user: User = Depends(get_current_user)):
    sess = _get_session(current_user.id, platform)
    return {"active": sess is not None, "started_at": sess["started_at"] if sess else None,
            "device_id": (sess or {}).get("device_id", ""), "device_os": (sess or {}).get("device_os", "")}


@router.get("/run-logs")
async def list_run_logs(
    platform: str | None = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Kullanıcının tüm koşum loglarını listeler (platform filtresi opsiyonel)."""
    log_dir = _run_log_dir(current_user.id, db)
    logs = []
    if log_dir.exists():
        for f in sorted(log_dir.glob("*.json"), reverse=True):
            try:
                data = _json.loads(f.read_text(encoding="utf-8"))
                if platform and data.get("platform") != platform:
                    continue
                # Hesaplanan süre
                try:
                    s = _datetime.datetime.fromisoformat(data.get("started_at", ""))
                    e = _datetime.datetime.fromisoformat(data.get("ended_at", ""))
                    dur = round((e - s).total_seconds())
                except Exception:
                    dur = None
                logs.append({
                    "run_id":        data.get("run_id", f.stem),
                    "scenario_name": data.get("scenario_name", "?"),
                    "platform":      data.get("platform", ""),
                    "started_at":    data.get("started_at", ""),
                    "duration_sec":  dur,
                    "total_steps":   data.get("total_steps", 0),
                    "passed":        data.get("passed", 0),
                    "failed":        data.get("failed", 0),
                    "exit_code":     data.get("exit_code", -1),
                })
            except Exception:
                pass
    return {"logs": logs}


@router.get("/run-logs/{run_id}")
async def get_run_log(
    run_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Tek bir koşumun detaylı logunu döndürür."""
    log_dir  = _run_log_dir(current_user.id, db)
    log_file = log_dir / f"{run_id}.json"
    if not log_file.exists():
        raise HTTPException(404, "Log bulunamadı")
    return _json.loads(log_file.read_text(encoding="utf-8"))


@router.post("/run-logs/push")
async def push_run_logs(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """ai_run_logs/ klasörünü Git'e commit edip push eder."""
    ws = _get_user_workspace(current_user.id, db)
    if not ws:
        raise HTTPException(400, "Git workspace bulunamadı. Önce bir repo clone edin.")
    repo = Path(ws.local_path)
    subprocess.run(["git", "add", "ai_run_logs/"], cwd=repo, capture_output=True)
    commit = subprocess.run(
        ["git", "commit", "-m", "chore: AI test run logları güncellendi"],
        cwd=repo, capture_output=True, text=True, encoding="utf-8"
    )
    if "nothing to commit" in commit.stdout or "nothing to commit" in commit.stderr:
        return {"ok": True, "message": "Yeni log yok, push yapılmadı"}
    push = subprocess.run(
        ["git", "push"], cwd=repo, capture_output=True, text=True, encoding="utf-8"
    )
    if push.returncode != 0:
        raise HTTPException(500, f"Push hatası: {push.stderr.strip()}")
    return {"ok": True, "message": "Loglar başarıyla push edildi ✓"}


@router.post("/stop/{run_id}")
async def stop_execution(run_id: str, current_user: User = Depends(get_current_user)):
    import signal as _sig
    proc = _running_processes.get(run_id)
    if not proc:
        raise HTTPException(404, "Koşum bulunamadı veya zaten bitti")

    # 1. Stop flag dosyası yaz — adımlar arası _check_stop() yakalar, hızlı çıkar
    sf = _stop_files.get(run_id)
    if sf:
        try:
            open(sf, "w").close()
        except Exception:
            pass

    # 2. Kısa bekle — adımlar arasındaysa graceful çıkış yeterli
    for _ in range(6):          # 3 saniye, 0.5s aralıklarla
        if proc.poll() is not None:
            return {"ok": True, "run_id": run_id}
        await asyncio.sleep(0.5)

    # 3. Hâlâ çalışıyor → CTRL_C_EVENT / SIGINT ile KeyboardInterrupt gönder
    #    Bu, script içindeki `finally: driver.quit()` bloğunu tetikler
    if proc.poll() is None:
        try:
            if sys.platform == "win32":
                proc.send_signal(_sig.CTRL_C_EVENT)
            else:
                proc.send_signal(_sig.SIGINT)
        except Exception:
            pass

    # 4. finally bloğunun çalışması için 3 saniye daha bekle
    for _ in range(6):
        if proc.poll() is not None:
            return {"ok": True, "run_id": run_id}
        await asyncio.sleep(0.5)

    # 5. Son çare: zorla öldür
    if proc.poll() is None:
        proc.terminate()

    return {"ok": True, "run_id": run_id}


def _sse(event: str, data: str) -> str:
    result = f"event: {event}\n"
    for line in data.split("\n"):
        result += f"data: {line}\n"
    return result + "\n"


# ─── Senaryo Kaydetme (Git Workspace) ─────────────────────────────────────────

class SaveScenarioRequest(BaseModel):
    name:       str
    steps:      list[dict]
    script:     str
    folder:     str | None  = None   # dolu → mevcut üzerine yaz
    platform:   str | None  = None   # "web" | "mobile" | "desktop"
    params:     dict | None = None   # platform-spesifik ayarlar


def _sanitize_folder_name(name: str) -> str:
    """Senaryo adını güvenli klasör adına çevirir."""
    clean = re.sub(r'[^\w\s-]', '', name, flags=re.UNICODE)
    clean = re.sub(r'[\s]+', '_', clean.strip())
    return clean[:64] or "senaryo"


def _get_user_workspace(user_id: int, db: Session):
    """Kullanıcının en son clone'lanmış workspace'ini döndürür."""
    ws = (
        db.query(UserWorkspace)
        .filter(UserWorkspace.user_id == user_id, UserWorkspace.is_cloned == True)
        .order_by(UserWorkspace.id.desc())
        .first()
    )
    return ws


@router.post("/scenarios")
async def save_scenario(
    req: SaveScenarioRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ws = _get_user_workspace(current_user.id, db)
    if not ws:
        raise HTTPException(400, "Git workspace bulunamadı. Önce Workspace sayfasından bir repo clone edin.")

    folder_name = req.folder if req.folder else _sanitize_folder_name(req.name)
    scenario_dir = Path(ws.local_path) / "ai_scenarios" / folder_name
    try:
        scenario_dir.mkdir(parents=True, exist_ok=True)
    except Exception as e:
        raise HTTPException(500, f"Klasör oluşturulamadı: {e}")

    # steps.json — b64 alanını koru (büyük veri)
    try:
        (scenario_dir / "steps.json").write_text(
            _json.dumps({
                "name":     req.name,
                "steps":    req.steps,
                "platform": req.platform,
                "params":   req.params or {},
            }, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (scenario_dir / "script.py").write_text(req.script, encoding="utf-8")
    except Exception as e:
        raise HTTPException(500, f"Dosya yazılamadı: {e}")

    # Git commit + push
    author_email = getattr(current_user, "email", None) or f"{current_user.username}@localhost"
    git_author = f"{current_user.username} <{author_email}>"
    committed = False
    try:
        committed = git_service.commit_changes(
            Path(ws.local_path),
            f"ai-gen: senaryo kaydet '{req.name}'",
            author=git_author,
        )
    except Exception as e:
        # Git hatası olsa bile local kayıt yapıldı — uyarı ver ama başarısız sayma
        print(f"[WARN] Git commit başarısız ({folder_name}): {e}")

    push_ok = False
    if committed:
        repo = ws.repository
        try:
            git_service.push_branch(
                Path(ws.local_path),
                ws.branch_name,
                username=repo.git_username if repo else None,
                token=repo.git_token if repo else None,
            )
            push_ok = True
        except Exception:
            pass  # push opsiyonel — local kayıt yeterli

    return {"ok": True, "folder": folder_name, "committed": committed, "pushed": push_ok}


@router.get("/scenarios")
async def list_scenarios(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ws = _get_user_workspace(current_user.id, db)
    if not ws:
        return {"scenarios": []}

    base = Path(ws.local_path) / "ai_scenarios"
    if not base.exists():
        return {"scenarios": []}

    result = []
    for folder in sorted(base.iterdir()):
        if not folder.is_dir():
            continue
        meta_file = folder / "steps.json"
        name      = folder.name
        platform  = ""
        if meta_file.exists():
            try:
                meta     = _json.loads(meta_file.read_text(encoding="utf-8"))
                name     = meta.get("name",     folder.name)
                platform = meta.get("platform", "")
            except Exception:
                pass
        result.append({"folder": folder.name, "name": name, "platform": platform})

    return {"scenarios": result}


@router.get("/scenarios/{folder_name}")
async def load_scenario(
    folder_name: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ws = _get_user_workspace(current_user.id, db)
    if not ws:
        raise HTTPException(400, "Git workspace bulunamadı.")

    scenario_dir = Path(ws.local_path) / "ai_scenarios" / folder_name
    if not scenario_dir.exists():
        raise HTTPException(404, "Senaryo bulunamadı.")

    steps_file  = scenario_dir / "steps.json"
    script_file = scenario_dir / "script.py"

    steps_data = _json.loads(steps_file.read_text(encoding="utf-8")) if steps_file.exists() else {}
    script     = script_file.read_text(encoding="utf-8") if script_file.exists() else ""

    return {
        "name":     steps_data.get("name", folder_name),
        "steps":    steps_data.get("steps", []),
        "script":   script,
        "platform": steps_data.get("platform"),
        "params":   steps_data.get("params", {}),
    }


@router.delete("/scenarios/{folder_name}")
async def delete_scenario(
    folder_name: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ws = _get_user_workspace(current_user.id, db)
    if not ws:
        raise HTTPException(400, "Git workspace bulunamadı.")

    scenario_dir = Path(ws.local_path) / "ai_scenarios" / folder_name
    if not scenario_dir.exists():
        raise HTTPException(404, "Senaryo bulunamadı.")

    import shutil
    shutil.rmtree(scenario_dir)

    _del_author_email = getattr(current_user, "email", None) or f"{current_user.username}@localhost"
    committed = git_service.commit_changes(
        Path(ws.local_path),
        f"ai-gen: senaryo sil '{folder_name}'",
        author=f"{current_user.username} <{_del_author_email}>",
    )
    push_ok = False
    if committed:
        repo = ws.repository
        try:
            git_service.push_branch(
                Path(ws.local_path),
                ws.branch_name,
                username=repo.git_username if repo else None,
                token=repo.git_token if repo else None,
            )
            push_ok = True
        except Exception:
            pass

    return {"ok": True, "committed": committed, "pushed": push_ok}


# ── Git Pull / Push ────────────────────────────────────────────────────────────

@router.post("/git-pull")
async def git_pull(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Remote'tan en son degisiklikleri cek (fetch + merge origin/main)."""
    ws = _get_user_workspace(current_user.id, db)
    if not ws:
        raise HTTPException(400, "Git workspace bulunamadi. Once workspace klonlayin.")

    local_path = Path(ws.local_path)
    if not local_path.exists():
        raise HTTPException(400, "Workspace klasoru bulunamadi.")

    # Once remote credential guncelle
    repo = ws.repository
    if repo and repo.git_username and repo.git_token:
        try:
            git_service.push_branch.__wrapped__ if False else None  # dummy
            from services import git_service as _gs
            _gs._run_git_command(
                ["git", "remote", "set-url", "origin",
                 repo.repo_url.replace("https://", f"https://{repo.git_username}:{repo.git_token}@")],
                cwd=local_path
            )
        except Exception:
            pass

    result = git_service.pull_from_main(local_path)

    from datetime import datetime, timezone
    ws.last_pulled = datetime.now(timezone.utc)
    db.commit()

    return {
        "ok":            result["success"],
        "message":       result["message"],
        "has_conflicts": result.get("has_conflicts", False),
    }


@router.post("/git-push")
async def git_push(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """ai_scenarios klasorunu commit edip remote branch'e push et."""
    ws = _get_user_workspace(current_user.id, db)
    if not ws:
        raise HTTPException(400, "Git workspace bulunamadi. Once workspace klonlayin.")

    local_path = Path(ws.local_path)
    if not local_path.exists():
        raise HTTPException(400, "Workspace klasoru bulunamadi.")

    repo = ws.repository

    commit_msg = f"ai-gen: senaryo guncelleme {_datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M')}"
    try:
        committed = git_service.commit_changes(
            local_path,
            commit_msg,
            author=f"{current_user.full_name or current_user.username} <{current_user.email}>",
        )
    except Exception as e:
        raise HTTPException(500, f"Commit hatasi: {e}")

    push_ok = False
    push_msg = "Commit basarili, push atlanamadi"
    if committed:
        try:
            git_service.push_branch(
                local_path,
                ws.branch_name,
                username=repo.git_username if repo else None,
                token=repo.git_token if repo else None,
            )
            push_ok = True
            push_msg = f"'{ws.branch_name}' branch'ine push edildi"
            from datetime import datetime, timezone
            ws.last_pushed = datetime.now(timezone.utc)
            ws.has_changes = False
            db.commit()
        except Exception as e:
            push_msg = f"Push hatasi: {e}"
    else:
        push_msg = "Commit edilecek degisiklik yok"

    return {
        "ok":       push_ok or not committed,
        "committed": committed,
        "pushed":   push_ok,
        "message":  push_msg,
    }
