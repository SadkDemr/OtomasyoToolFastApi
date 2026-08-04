"""
TestCaseGen Service
--------------------
Ekran görüntülerinden test case başlığı çıkarma, kural tabanlı
önceliklendirme, adım doğrulama ve ADO formatlı Excel export.
Veritabanı YOK — state dosya tabanlı (bkz. session_* fonksiyonları).
"""

import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font


# ──────────────────────────────────────────────────────
# JSON temizleme / parse
# ──────────────────────────────────────────────────────

def clean_json_text(text: str) -> str:
    """Markdown code fence'leri (```json ... ``` veya ``` ... ```) temizler."""
    cleaned = text.strip()
    cleaned = re.sub(r'^```(?:json)?\s*', '', cleaned)
    cleaned = re.sub(r'\s*```$', '', cleaned)
    return cleaned.strip()


def extract_json_object(text: str) -> Optional[dict]:
    """LLM çıktısından ilk JSON nesnesini çıkarır. Parse edilemezse None döner."""
    if not text:
        return None
    cleaned = clean_json_text(text)
    start = cleaned.find('{')
    end = cleaned.rfind('}')
    if start == -1 or end == -1 or end <= start:
        return None
    candidate = cleaned[start:end + 1]
    try:
        return json.loads(candidate)
    except json.JSONDecodeError:
        return None


# ──────────────────────────────────────────────────────
# Oturum/Ekran Depolama (Dosya Tabanlı)
# ──────────────────────────────────────────────────────

DATA_DIR = Path(__file__).parent.parent / "testcasegen_data"


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _session_dir(user_id: int, session_id: str) -> Path:
    return DATA_DIR / str(user_id) / session_id


def _session_json_path(user_id: int, session_id: str) -> Path:
    return _session_dir(user_id, session_id) / "session.json"


def _screens_dir(user_id: int, session_id: str) -> Path:
    return _session_dir(user_id, session_id) / "screenshots"


def create_session(user_id: int) -> dict:
    session_id = str(uuid.uuid4())
    session = {
        "id": session_id,
        "user_id": user_id,
        "created_at": _now_iso(),
        "updated_at": _now_iso(),
        "current_step": 1,
        "screens": [],
        "titles": [],
        "testcases": [],
    }
    _session_dir(user_id, session_id).mkdir(parents=True, exist_ok=True)
    _screens_dir(user_id, session_id).mkdir(parents=True, exist_ok=True)
    save_session(user_id, session_id, session)
    return session


def list_sessions(user_id: int) -> list:
    base = DATA_DIR / str(user_id)
    if not base.exists():
        return []
    sessions = []
    for entry in sorted(base.iterdir()):
        json_path = entry / "session.json"
        if json_path.exists():
            try:
                sessions.append(load_session(user_id, entry.name))
            except (FileNotFoundError, json.JSONDecodeError):
                continue
    sessions.sort(key=lambda s: s.get("created_at", ""), reverse=True)
    return sessions


def load_session(user_id: int, session_id: str) -> dict:
    json_path = _session_json_path(user_id, session_id)
    if not json_path.exists():
        raise FileNotFoundError(f"Oturum bulunamadı: {session_id}")
    with open(json_path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_session(user_id: int, session_id: str, session: dict) -> None:
    session["updated_at"] = _now_iso()
    json_path = _session_json_path(user_id, session_id)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(session, f, ensure_ascii=False, indent=2)


def delete_session(user_id: int, session_id: str) -> None:
    import shutil
    d = _session_dir(user_id, session_id)
    if d.exists():
        shutil.rmtree(d)


def screen_image_path(user_id: int, session_id: str, index: int, ext: str = ".png") -> Path:
    return _screens_dir(user_id, session_id) / f"{index}{ext}"


def add_screen(user_id: int, session_id: str, image_bytes: bytes, ext: str = ".png", note: str = "") -> dict:
    session = load_session(user_id, session_id)
    index = len(session["screens"])
    _screens_dir(user_id, session_id).mkdir(parents=True, exist_ok=True)
    screen_image_path(user_id, session_id, index, ext).write_bytes(image_bytes)
    session["screens"].append({
        "index": index,
        "image_ext": ext,
        "user_note": note,
        "candidate_titles": [],
        "screen_summary": "",
        "parse_error": False,
    })
    save_session(user_id, session_id, session)
    return session


def _reindex_and_rename_screens(user_id: int, session_id: str, session: dict) -> None:
    """Ekranlar yer değiştirdikten sonra dosya adlarını da yeni index'e göre yeniden adlandırır."""
    tmp_paths = []
    for new_idx, screen in enumerate(session["screens"]):
        old_path = screen_image_path(user_id, session_id, screen["index"], screen["image_ext"])
        tmp_path = old_path.with_suffix(old_path.suffix + f".tmp{new_idx}")
        old_path.rename(tmp_path)
        tmp_paths.append((tmp_path, new_idx))
    for tmp_path, new_idx in tmp_paths:
        screen = session["screens"][new_idx]
        final_path = screen_image_path(user_id, session_id, new_idx, screen["image_ext"])
        tmp_path.rename(final_path)
        screen["index"] = new_idx


def reorder_screens(user_id: int, session_id: str, order: list) -> dict:
    session = load_session(user_id, session_id)
    if sorted(order) != list(range(len(session["screens"]))):
        raise ValueError("Geçersiz sıralama: her index tam olarak bir kez geçmeli")
    session["screens"] = [session["screens"][i] for i in order]
    _reindex_and_rename_screens(user_id, session_id, session)
    save_session(user_id, session_id, session)
    return session


def update_screen_note(user_id: int, session_id: str, index: int, note: str) -> dict:
    session = load_session(user_id, session_id)
    session["screens"][index]["user_note"] = note
    save_session(user_id, session_id, session)
    return session


def delete_screen(user_id: int, session_id: str, index: int) -> dict:
    session = load_session(user_id, session_id)
    screen = session["screens"][index]
    screen_image_path(user_id, session_id, index, screen["image_ext"]).unlink(missing_ok=True)
    del session["screens"][index]
    _reindex_and_rename_screens(user_id, session_id, session)
    save_session(user_id, session_id, session)
    return session


# ──────────────────────────────────────────────────────
# Kural tabanlı önceliklendirme (LLM karar vermez, sadece skor üzerinden gerekçe alır)
# ──────────────────────────────────────────────────────

WEIGHTS = {
    "transaction": 3.0,        # form/tutar girişi/onay butonu içeren ekran
    "info_only": -1.0,         # salt bilgi gösterimi
    "negative_potential": 1.5, # validasyon gerektiren input alanı → hata senaryosu riski
    "depth_penalty": 0.3,      # ana akıştan her adım uzaklaştıkça düşen ağırlık
}
THRESHOLD_HIGH = 3.0
THRESHOLD_MEDIUM = 1.0

TRANSACTION_KEYWORDS = [
    "onayla", "öde", "ödeme", "gönder", "kaydet", "tutar", "miktar",
    "transfer", "satın al", "işlem yap", "onaylanır", "gönderilir",
]
INFO_KEYWORDS = [
    "hakkımızda", "sıkça sorulan", "sss", "bilgilendirme", "yardım",
    "iletişim", "kullanım koşulları", "gizlilik",
]
NEGATIVE_KEYWORDS = [
    "doğrulama", "zorunlu alan", "hatalı", "geçersiz", "input", "şifre",
    "gir", "girilir", "validasyon",
]


def score_title(text: str, screen_note: str, screen_index: int) -> tuple[float, str]:
    haystack = f"{text} {screen_note}".lower()
    score = 0.0
    if any(k in haystack for k in TRANSACTION_KEYWORDS):
        score += WEIGHTS["transaction"]
    if any(k in haystack for k in INFO_KEYWORDS):
        score += WEIGHTS["info_only"]
    if any(k in haystack for k in NEGATIVE_KEYWORDS):
        score += WEIGHTS["negative_potential"]
    score -= WEIGHTS["depth_penalty"] * screen_index

    if score >= THRESHOLD_HIGH:
        level = "high"
    elif score >= THRESHOLD_MEDIUM:
        level = "medium"
    else:
        level = "low"
    return round(score, 2), level


# ──────────────────────────────────────────────────────
# Test case adım validator'ü (4 kesin kural)
# ──────────────────────────────────────────────────────

_IMPERATIVE_PATTERNS = [
    r'\btıkla\b', r'\btıklayınız\b', r'\baç\b', r'\byaz\b', r'\bgir\b',
    r'\bgiriniz\b', r'\bseç\b', r'\bseçiniz\b', r'\bonayla\b', r'\bonaylayınız\b',
    r'\bbas\b', r'\bbasınız\b', r'\bkaydet\b', r'\bdoldur\b', r'\bdolduruniz\b',
    r'\bkapat\b', r'\bsil\b', r'\bgönder\b', r'\bgönderiniz\b',
]

_VARIABLE_KEYWORDS = ["kur", "tutar", "bakiye", "hesaplanan", "miktar", "oran", "komisyon", "vade"]

_ACTION_VERB_PATTERN = re.compile(
    r'\b(tıklanır|açılır|yazılır|girilir|seçilir|onaylanır|basılır|kaydedilir|gönderilir|doldurulur)\b'
)


def _check_imperative(step_action: str) -> Optional[str]:
    unquoted = re.sub(r'"[^"]*"', '', step_action)
    lowered = unquoted.lower()
    for pattern in _IMPERATIVE_PATTERNS:
        if re.search(pattern, lowered):
            return "Emir kipi tespit edildi — pasif çatı kullanılmalı (örn. 'tıklanır', 'açılır')"
    return None


def _check_variable_notation(step_action: str) -> Optional[str]:
    lowered = step_action.lower()
    at_tokens = re.findall(r'@\w+', lowered)
    for keyword in _VARIABLE_KEYWORDS:
        match = re.search(rf'\b{keyword}\w*\b', lowered)
        if not match:
            continue
        word = match.group(0)
        if any(keyword in tok for tok in at_tokens):
            continue
        return f"'{keyword}' değişken gibi görünüyor ama @ notasyonu kullanılmamış (örn. @{keyword})"
    return None


def _check_multiple_actions(step_action: str) -> Optional[str]:
    lowered = step_action.lower()
    if ' ve ' not in lowered:
        return None
    if len(_ACTION_VERB_PATTERN.findall(lowered)) >= 2:
        return 'Step Action birden fazla aksiyon içeriyor gibi görünüyor (" ve " ile birleştirilmiş)'
    return None


def _check_empty_expected(step_expected: str) -> Optional[str]:
    if not step_expected or not step_expected.strip():
        return "Step Expected boş bırakılamaz"
    return None


def validate_steps(steps: list) -> list:
    checks = [
        ("imperative_mood", lambda s: _check_imperative(s["step_action"])),
        ("missing_variable_notation", lambda s: _check_variable_notation(s["step_action"])),
        ("multiple_actions", lambda s: _check_multiple_actions(s["step_action"])),
        ("empty_expected", lambda s: _check_empty_expected(s["step_expected"])),
    ]
    result = []
    for step in steps:
        violations = []
        for rule, check in checks:
            message = check(step)
            if message:
                violations.append({"rule": rule, "message": message})
        result.append({**step, "violations": violations})
    return result


# ──────────────────────────────────────────────────────
# Ollama prompt üreticileri
# ──────────────────────────────────────────────────────

def build_title_extraction_prompt(screen_note: str, prev_note: str, prev_summary: str) -> str:
    context = ""
    if prev_note or prev_summary:
        context = (
            f"\nBir önceki ekranın notu: \"{prev_note}\"\n"
            f"Bir önceki ekranın özeti: \"{prev_summary}\"\n"
            "Bu bağlamı akış bütünlüğünü korumak için kullan.\n"
        )
    note_line = f"\nKullanıcının bu ekran için notu: \"{screen_note}\"\n" if screen_note else ""
    return (
        "Sen bir mobil uygulama test analistisin. Sana verilen ekran görüntüsünü incele.\n"
        f"{context}{note_line}\n"
        "Bu ekrandan test edilebilecek 2-5 aday test case başlığı çıkar ve ekranın kısa bir özetini yaz.\n\n"
        "SADECE aşağıdaki şemaya uyan JSON döndür, başka hiçbir metin ekleme, markdown code fence kullanma:\n"
        '{"screen_titles": [{"screen_index": 0, "candidate_titles": ["...", "..."], "screen_summary": "..."}]}'
    )


def build_rationale_prompt(title_text: str, level: str, score: float) -> str:
    return (
        "Aşağıdaki test case başlığına, kural tabanlı bir motor tarafından zaten "
        f"'{level}' önceliği ve {score} skoru atandı. Karar değiştirme, sadece bu "
        "skoru 1-2 cümlelik kısa bir gerekçeye çevir (Türkçe, test analistine hitaben):\n\n"
        f"Başlık: \"{title_text}\"\n"
        f"Öncelik: {level}\n"
        f"Skor: {score}\n\n"
        "Yalnızca gerekçe cümlesini döndür, başka açıklama ekleme."
    )


def build_testcase_prompt(title_text: str, screen_summary: str, screen_note: str) -> str:
    return (
        "Sen bir mobil uygulama test otomasyon uzmanısın. Aşağıdaki test case başlığı için "
        "adım adım test senaryosu üret.\n\n"
        f"Test case başlığı: \"{title_text}\"\n"
        f"Ekran özeti: \"{screen_summary}\"\n"
        f"Kullanıcı notu: \"{screen_note}\"\n\n"
        "KESİN KURALLAR:\n"
        "1. Her Step Action pasif fiil çatısıyla yazılmalı (\"tıklanır\", \"açılır\", \"görülür\" — "
        "\"tıkla\", \"aç\" gibi EMİR KİPİ YASAK).\n"
        "2. Her Step Action'da TEK bir aksiyon olmalı (birden fazla tıklama/yazma tek satırda birleştirilmemeli).\n"
        "3. Dinamik/değişken değerler (güncel kur, hesaplanan tutar gibi) @ParamAdı notasyonuyla yazılmalı "
        "(örn. @güncel_kur, @hesaplanan_tutar).\n"
        "4. Her step'in boş olmayan bir Step Expected (beklenen sonuç) alanı olmalı.\n\n"
        "SADECE aşağıdaki şemaya uyan JSON döndür, başka metin/markdown fence ekleme:\n"
        '{"steps": [{"step_action": "...", "step_expected": "..."}]}'
    )


# ──────────────────────────────────────────────────────
# Excel export builder (ADO format)
# ──────────────────────────────────────────────────────

def build_export_workbook(testcases: list, titles_by_id: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    headers = ["ID", "Work Item Type", "Title", "Test Step", "Step Action",
               "Step Expected", "Area Path", "Assigned To", "State"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)

    widths = [8, 16, 40, 10, 55, 45, 20, 20, 12]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    for tc in testcases:
        title_text = titles_by_id.get(tc["title_id"], {}).get("text", "")
        ws.append([None, "Test Case", title_text, None, None, None, None, None, None])
        for i, step in enumerate(tc["steps"], start=1):
            ws.append([None, None, None, i, step.get("step_action", ""),
                       step.get("step_expected", ""), None, None, None])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.font.name != "Arial" or not cell.font.bold:
                cell.font = Font(name="Arial", bold=False)

    return wb
